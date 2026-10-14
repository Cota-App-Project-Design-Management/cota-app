import streamlit as st
from streamlit_gsheets import GSheetsConnection
import pandas as pd
import base64
import altair as alt
import unicodedata
from datetime import datetime, timedelta
import calendar
import uuid
import io
import json
import re


# ==========================================
# 1. CONFIGURACIÓN Y ESTÉTICA (COTA APP)
# ==========================================
raw_title = st.secrets.get("APP_TITLE", "COTA APP")
APP_TITLE = "COTA APP" if raw_title in ["Binda Constructores", "Found Operating Systems", "Found OS"] else raw_title

st.set_page_config(page_title="Cota App", layout="wide", page_icon=st.secrets.get("PAGE_ICON", "icono.png"), initial_sidebar_state="expanded")

lista_socios = ["Socio 1", "Socio 2"]
SOCIOS = lista_socios

def safe_float(val, default=0.0):
    """Convierte a float de forma segura, soportando comas decimales, formatos latinos y valores vacíos."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return default
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if s == '' or s.lower() in ['nan', 'none', 'null']:
        return default
    try:
        return float(s)
    except ValueError:
        try:
            if ',' in s and '.' in s:
                if s.rfind(',') > s.rfind('.'):
                    s_clean = s.replace('.', '').replace(',', '.')
                else:
                    s_clean = s.replace(',', '')
                return float(s_clean)
            elif ',' in s:
                return float(s.replace(',', '.'))
            elif '.' in s:
                return float(s)
        except ValueError:
            return default
    return default

def limpiar_especialidad(texto):
    if not texto or pd.isna(texto): return ""
    return ''.join(c for c in unicodedata.normalize('NFD', str(texto).strip()) if unicodedata.category(c) != 'Mn').capitalize()

def get_image_base64(path):
    if not path:
        return None
    if path.startswith("http://") or path.startswith("https://"):
        return path
    try:
        with open(path, "rb") as f: return f"data:image/png;base64,{base64.b64encode(f.read()).decode()}"
    except: return None

# CSS Estilo Cota App (Minimalista, Técnico, Estilo Palantir - Premium Edition)
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600;800&family=Inter:wght@300;400;600;700&display=swap');
    
    :root {
        --custom-bg: #070B14;
        --card-bg: #0d1424;
        --card-bg-hover: #111b30;
        --border-color: #1e293b;
        --text-muted: #cbd5e1;
        --text-secondary: #94a3b8;
        --text-color: #ffffff;
        --btn-bg-start: #070B14;
        --btn-bg-end: #0d1424;
        --card-shadow: 0 4px 6px rgba(0, 0, 0, 0.3);
        --login-bg: rgba(10, 15, 28, 0.85);
        --login-shadow: 0 10px 30px rgba(0,0,0,0.5);
    }
    
    @media (prefers-color-scheme: light) {
        :root {
            --custom-bg: #ffffff;
            --card-bg: #f8fafc;
            --card-bg-hover: #f1f5f9;
            --border-color: #cbd5e1;
            --text-muted: #475569;
            --text-secondary: #64748b;
            --text-color: #0f172a;
            --btn-bg-start: #ffffff;
            --btn-bg-end: #f1f5f9;
            --card-shadow: 0 4px 6px rgba(0, 0, 0, 0.05);
            --login-bg: rgba(248, 250, 252, 0.95);
            --login-shadow: 0 10px 30px rgba(0,0,0,0.08);
        }
    }
    
    html, body { 
        font-size: 13px !important; 
        font-family: 'Outfit', 'Inter', 'Segoe UI', sans-serif !important; 
        background-color: var(--custom-bg) !important; 
    }
    
    /* Evitar que las opciones del selectbox se vean delgadas o borrosas */
    div[role="listbox"] ul li, 
    div[role="listbox"] [role="option"], 
    div[role="listbox"] div,
    ul[role="listbox"] li {
        font-family: 'Inter', 'Segoe UI', sans-serif !important;
        font-weight: 600 !important;
        color: var(--text-color) !important;
        font-size: 13px !important;
        -webkit-font-smoothing: antialiased !important;
        -moz-osx-font-smoothing: grayscale !important;
    }
    
    div[data-baseweb="select"] div {
        font-family: 'Inter', 'Segoe UI', sans-serif !important;
        font-weight: 600 !important;
        -webkit-font-smoothing: antialiased !important;
        -moz-osx-font-smoothing: grayscale !important;
    }
    
    /* BRANDING HEADER - ESTILO PALANTIR */
    .header-container { display: flex; align-items: center; padding: 25px 0; gap: 20px; border-bottom: 1px solid var(--border-color); margin-bottom: 25px; }
    
    /* Animación del logotipo */
    @keyframes pulse-branding {
        0% { filter: drop-shadow(0 0 2px rgba(0, 229, 255, 0.15)); }
        50% { filter: drop-shadow(0 0 10px rgba(0, 229, 255, 0.45)); }
        100% { filter: drop-shadow(0 0 2px rgba(0, 229, 255, 0.15)); }
    }
    .header-logo { 
        height: 60px; 
        width: auto; 
        object-fit: contain; 
        animation: pulse-branding 4s infinite ease-in-out;
    }
    
    .ice-card-logo {
        filter: drop-shadow(0 0 5px rgba(255, 255, 255, 0.5));
        transition: all 0.3s ease;
    }
    .ice-card-logo:hover {
        filter: drop-shadow(0 0 8px rgba(0, 229, 255, 0.6));
    }
    
    .header-divider { width: 2px; height: 50px; background-color: var(--border-color); margin: 0 10px; }
    .header-title-box { display: flex; flex-direction: column; justify-content: center;}
    .header-title { color: var(--text-color) !important; font-size: 34px !important; font-weight: 800; letter-spacing: -1px; margin:0; line-height: 1.1; }
    .header-subtitle { color: #00E5FF; font-family: 'Courier New', monospace; font-size: 13px; letter-spacing: 4px; font-weight: bold; margin-top: 4px;}
    
    /* Tarjetas Métricas con Glassmorphism y Micro-animaciones */
    .metric-card { 
        background-color: var(--card-bg) !important; 
        padding: 18px; 
        border-radius: 6px; 
        text-align: center; 
        border-top: 3px solid #00E5FF !important; 
        border: 1px solid var(--border-color) !important; 
        transition: all 0.3s cubic-bezier(0.25, 0.8, 0.25, 1);
        box-shadow: var(--card-shadow) !important;
    }
    .metric-card:hover {
        transform: translateY(-5px);
        box-shadow: 0 12px 20px rgba(0, 229, 255, 0.15);
        border-color: #00E5FF;
        border-top-width: 4px;
    }
    
    .metric-card h4 { margin-bottom: 5px; color: var(--text-muted) !important; font-weight: 600; font-size: 14px; }
    .metric-card h2 { margin-top: 0; margin-bottom: 5px; font-size: 28px; color: var(--text-color) !important; }
    
    /* Cajas VDC interactivas */
    .vdc-framework-box { 
        border: 1px solid var(--border-color) !important; 
        background-color: var(--card-bg) !important; 
        padding: 15px; 
        text-align: center; 
        border-radius: 6px; 
        margin-bottom: 5px; 
        transition: all 0.2s ease-in-out;
        color: var(--text-color) !important;
    }
    .vdc-framework-box:hover {
        border-color: #00E5FF !important;
        background-color: var(--card-bg-hover) !important;
    }
    
    .vdc-framework-title { font-weight: bold; font-family: monospace; margin-bottom: 8px; font-size: 14px; color: var(--text-color) !important; }
    .vdc-arrow { text-align: center; font-size: 20px; color: #94a3b8; margin-bottom: 5px; }
    
    /* Estilo premium de expander estilo consola */
    div[data-testid="stExpander"] {
        border: none !important;
        background-color: transparent !important;
        box-shadow: none !important;
        margin-bottom: 8px !important;
    }
    div[data-testid="stExpander"] > details {
        border: none !important;
        background-color: transparent !important;
    }
    div[data-testid="stExpander"] > details > summary, button[data-baseweb="accordion"] {
        border: 1px solid var(--border-color) !important;
        background-color: var(--card-bg) !important;
        border-radius: 4px !important;
        padding: 10px 15px !important;
        transition: all 0.25s ease-in-out !important;
    }
    div[data-testid="stExpander"] > details > summary:hover, button[data-baseweb="accordion"]:hover {
        border-color: #00E5FF !important;
        box-shadow: 0 0 10px rgba(0, 229, 255, 0.1) !important;
    }
    div[data-testid="stExpander"] summary div, button[data-baseweb="accordion"] div {
        font-family: 'Courier New', monospace !important;
        font-size: 13px !important;
        color: var(--text-muted) !important;
        font-weight: 600 !important;
        white-space: pre !important;
    }
    
    .proj-label { color: #00E5FF; font-family: monospace; font-weight: bold; font-size: 11px; text-transform: uppercase; margin-top: 15px; margin-bottom: 8px; border-bottom: 1px solid var(--border-color); padding-bottom: 4px; letter-spacing: 1px; }
    .tag-user-a { background-color: #1e3a8a; color: #bfdbfe; padding: 2px 8px; border-radius: 2px; font-size: 10px; font-weight: bold; font-family: monospace; }
    .tag-user-r { background-color: #14532d; color: #bbf7d0; padding: 2px 8px; border-radius: 2px; font-size: 10px; font-weight: bold; font-family: monospace; }
    
    .vdc-hub-title { color: #00E5FF; font-family: 'Courier New', monospace; font-weight: bold; border-left: 3px solid #00E5FF; padding-left: 15px; text-transform: uppercase; }
    
    /* CALENDARIO TÉCNICO */
    .calendar-container { display: grid; grid-template-columns: repeat(7, 1fr); gap: 5px; margin-top: 20px; }
    .calendar-header-day { text-align: center; font-weight: bold; color: #00E5FF; padding: 10px 0; border-bottom: 1px solid var(--border-color); font-family: monospace; text-transform: uppercase; }
    .calendar-cell { min-height: 120px; border: 1px solid var(--border-color); border-radius: 6px; padding: 8px; background-color: var(--card-bg); transition: border-color 0.2s ease; }
    .calendar-cell:hover { border-color: #00E5FF; }
    .calendar-cell.today { border-color: #00E5FF; background-color: rgba(0, 229, 255, 0.05); }
    .cal-task-item { font-size: 10px; padding: 4px 6px; margin-bottom: 4px; border-radius: 2px; font-weight: 600; position: relative !important; overflow: visible !important; cursor: pointer !important; }
    .cal-task-item .tooltip-text {
        visibility: hidden !important;
        width: 220px !important;
        background-color: #070b14 !important;
        color: #cbd5e1 !important;
        border: 1px solid var(--border-color) !important;
        border-left: 3px solid #00E5FF !important;
        border-radius: 4px !important;
        padding: 8px 10px !important;
        position: absolute !important;
        z-index: 9999 !important;
        bottom: 125% !important;
        left: 50% !important;
        transform: translateX(-50%) !important;
        opacity: 0 !important;
        transition: opacity 0.15s ease-in-out !important;
        font-family: monospace !important;
        font-size: 10px !important;
        line-height: 1.4 !important;
        box-shadow: 0 8px 20px rgba(0,0,0,0.6) !important;
        pointer-events: none !important;
        white-space: normal !important;
        text-align: left !important;
    }
    .cal-task-item:hover .tooltip-text {
        visibility: visible !important;
        opacity: 1 !important;
    }
    .cal-task-item .tooltip-text::after {
        content: "" !important;
        position: absolute !important;
        top: 100% !important;
        left: 50% !important;
        margin-left: -5px !important;
        border-width: 5px !important;
        border-style: solid !important;
        border-color: var(--border-color) transparent transparent transparent !important;
    }
    
    .calendar-legend {
        display: flex !important;
        flex-wrap: wrap !important;
        justify-content: center !important;
        gap: 15px !important;
        margin-top: 25px !important;
        padding: 15px !important;
        background: var(--card-bg) !important;
        border: 1px solid var(--border-color) !important;
        border-radius: 8px !important;
    }
    .calendar-legend-item {
        display: flex !important;
        align-items: center !important;
        gap: 8px !important;
        font-family: monospace !important;
        font-size: 11px !important;
        color: var(--text-secondary) !important;
    }
    .calendar-legend-color {
        width: 12px !important;
        height: 12px !important;
        border-radius: 3px !important;
    }
    
    /* Estilo premium para botones secundarios/estándar de Streamlit */
    .stButton>button[kind="secondary"], .stButton>button:not([kind="primary"]) {
        background: linear-gradient(135deg, var(--btn-bg-start) 0%, var(--btn-bg-end) 100%) !important;
        color: var(--text-muted) !important;
        border: 1px solid var(--border-color) !important;
        border-radius: 4px !important;
        padding: 6px 16px !important;
        font-family: 'Outfit', sans-serif !important;
        font-weight: 600 !important;
        letter-spacing: 0.5px !important;
        transition: all 0.25s cubic-bezier(0.25, 0.8, 0.25, 1) !important;
    }
    .stButton>button[kind="secondary"]:hover, .stButton>button:not([kind="primary"]):hover {
        background: linear-gradient(135deg, #00E5FF 0%, #008b9a 100%) !important;
        color: #070B14 !important;
        border-color: #00E5FF !important;
        box-shadow: 0 0 15px rgba(0, 229, 255, 0.4) !important;
        transform: scale(1.02);
    }
    .stButton>button[kind="secondary"]:active, .stButton>button:not([kind="primary"]):active {
        transform: scale(0.98);
    }
    
    /* Estilo premium destacado para el botón primario de Sincronizar (Rojo original) */
    .stButton>button[kind="primary"] {
        background: linear-gradient(135deg, #FF1744 0%, #D50000 100%) !important;
        color: #ffffff !important;
        border: 1px solid #FF1744 !important;
        border-radius: 4px !important;
        padding: 8px 16px !important;
        font-family: 'Outfit', sans-serif !important;
        font-weight: 800 !important;
        letter-spacing: 1px !important;
        transition: all 0.25s cubic-bezier(0.25, 0.8, 0.25, 1) !important;
        box-shadow: 0 4px 6px rgba(213, 0, 0, 0.2) !important;
    }
    .stButton>button[kind="primary"]:hover {
        background: linear-gradient(135deg, #FF5252 0%, #FF1744 100%) !important;
        color: #ffffff !important;
        border-color: #FF5252 !important;
        box-shadow: 0 0 18px rgba(255, 23, 68, 0.55) !important;
        transform: scale(1.02);
    }
    .stButton>button[kind="primary"]:active {
        transform: scale(0.98);
    }

    /* LOGIN CARD PREMIUM STYLE */
    .login-card {
        background: var(--login-bg) !important;
        border: 1px solid var(--border-color) !important;
        border-radius: 12px;
        padding: 35px 30px;
        box-shadow: var(--login-shadow) !important;
        text-align: center;
        font-family: 'Outfit', sans-serif;
    }
    .login-title {
        color: var(--text-color) !important;
        font-weight: 800;
        font-family: 'Outfit', sans-serif;
        letter-spacing: 1.5px;
        margin-bottom: 5px;
        font-size: 24px;
        margin-top: 0;
    }
    .login-subtitle {
        color: #00E5FF;
        font-size: 13px;
        font-family: 'Inter', sans-serif;
        font-weight: 600;
        letter-spacing: 1.5px;
        margin-bottom: 15px;
        text-transform: uppercase;
    }
    .login-access {
        color: var(--text-secondary) !important;
        font-size: 11px;
        font-family: 'Inter', sans-serif;
        letter-spacing: 2px;
        margin-bottom: 5px;
    }
    
    /* Xian Style Additions for Matriz Operaciones */
    .xian-section-header {
        color: #00E5FF !important;
        font-family: 'Courier New', monospace !important;
        font-weight: bold !important;
        font-size: 13px !important;
        text-transform: uppercase !important;
        margin-top: 25px !important;
        margin-bottom: 12px !important;
        border-bottom: 1px solid var(--border-color) !important;
        padding-bottom: 6px !important;
        letter-spacing: 1px !important;
        display: flex !important;
        align-items: center !important;
    }
    .xian-section-header-icon {
        color: #00E5FF !important;
        margin-right: 8px !important;
        font-size: 10px !important;
    }
    .xian-badge-container {
        display: flex !important;
        gap: 8px !important;
        font-family: 'Courier New', monospace !important;
        font-size: 11px !important;
        align-items: center !important;
        justify-content: flex-end !important;
        height: 100% !important;
    }
    .xian-badge {
        padding: 6px 12px !important;
        border-radius: 2px !important;
        font-weight: bold !important;
        border: 1px solid transparent !important;
        white-space: nowrap !important;
    }
    .xian-badge-pactado {
        background-color: var(--card-bg) !important;
        color: var(--text-muted) !important;
        border-color: var(--border-color) !important;
    }
    .xian-badge-pagado {
        background-color: rgba(0, 229, 255, 0.05) !important;
        color: #00E5FF !important;
        border-color: rgba(0, 229, 255, 0.3) !important;
    }
    .xian-badge-saldo-ok {
        background-color: rgba(16, 185, 129, 0.05) !important;
        color: #10b981 !important;
        border-color: rgba(16, 185, 129, 0.3) !important;
    }
    .xian-badge-saldo-error {
        background-color: rgba(239, 68, 68, 0.05) !important;
        color: #ef4444 !important;
        border-color: rgba(239, 68, 68, 0.3) !important;
    }
    .xian-status-header {
        font-family: 'Courier New', monospace !important;
        font-weight: bold !important;
        letter-spacing: 1px !important;
        margin-top: 25px !important;
        margin-bottom: 15px !important;
        border-bottom: 1px solid var(--border-color) !important;
        padding-bottom: 6px !important;
        display: flex !important;
        align-items: center !important;
    }
    .xian-status-header-bar {
        width: 4px !important;
        height: 14px !important;
        display: inline-block !important;
        margin-right: 8px !important;
    }
    </style>
    """, unsafe_allow_html=True)

# Listas de Estados 
status_proyectos = ["🟡 Planificación", "🔵 En Proceso", "🟢 Completado"]
status_tasks = ["🟡 Planificado", "🔵 En Proceso", "⚪ En espera", "🟢 Completado a tiempo", "🔴 Completado con retraso"]

# ==========================================
# 1.5. SISTEMA DE AUTENTICACIÓN (MARCA BLANCA SEGURA)
# ==========================================
password_configurado = st.secrets.get("APP_PASSWORD", "")
if not password_configurado:
    # Soporte por si se pegó al final del TOML bajo la sección [connections.gsheets]
    try:
        password_configurado = st.secrets.get("connections", {}).get("gsheets", {}).get("APP_PASSWORD", "")
    except:
        pass

if password_configurado:
    if 'authenticated' not in st.session_state:
        st.session_state.authenticated = False
        
    if not st.session_state.authenticated:
        # Renderizar formulario de login elegante estilo Cota App
        c_space1, c_login, c_space2 = st.columns([1.5, 2, 1.5])
        with c_login:
            st.markdown("<br><br><br>", unsafe_allow_html=True)
            logo_src = get_image_base64(st.secrets.get("LOGO_IMAGE", "Logo2.png"))
            logo_html = f'<img src="{logo_src}" style="max-height: 80px; display: block; margin: 0 auto 15px auto;" class="header-logo">' if logo_src else ""
            
            st.markdown(f"""
                <div class="login-card">
                    {logo_html}
                    <h2 class="login-title">{APP_TITLE} <span style='font-size: 12px; color: #FF9800; background-color: rgba(255, 152, 0, 0.15); border: 1px solid #FF9800; border-radius: 4px; padding: 1px 6px; vertical-align: middle; font-family: monospace;'>BETA</span></h2>
                    <p class="login-subtitle">PROJECT & DESIGN MANAGEMENT</p>
                    <p class="login-access">ACCESO AUTORIZADO</p>
                </div>
            """, unsafe_allow_html=True)
            
            with st.form("form_login"):
                p_in = st.text_input("Contraseña de Acceso:", type="password", placeholder="Ingrese su contraseña")
                submit_btn = st.form_submit_button("INGRESAR AL SISTEMA", use_container_width=True)
                
                if submit_btn:
                    if p_in == password_configurado:
                        st.session_state.authenticated = True
                        st.success("Acceso concedido.")
                        st.rerun()
                    else:
                        st.error("Contraseña incorrecta. Por favor intente de nuevo.")
                        
        st.stop()

# ==========================================
# 2. MOTOR DE BASE DE DATOS
# ==========================================
try:
    conn = st.connection("gsheets", type=GSheetsConnection)
    sheet_url = st.secrets.get("GSHEETS_URL", "")
    nube_ok = True if sheet_url else False
except: nube_ok = False

cols_proy = ['Nombre', 'Encargado', 'Ubicación', 'Estado', 'Inicio', 'Fin', 'Notas', 'Color', 'Cliente', 'Presupuesto', 'Utilidades', 'GG_Presupuesto', 'Archivado', 'Metodologia', 'Division', 'Tipo_Proyecto', 'Num_Pisos', 'Num_Sotanos']
cols_tasks = ['Proyecto', 'Partida', 'Orden', 'Tarea', 'Especialidad', 'Contratista', 'Estado', 'Inicio', 'Fin', 'Causa_Cumplimiento', 'Culpa_Contratista', 'Permitir_Cruce', 'Dependencia', 'Notas']
cols_maestro = ['ID_Contratista', 'Especialidad', 'Ubicación_Base', 'Notas']
cols_hist = ['ID_Contratista', 'Estado', 'Responsable']
cols_gg = ['Proyecto', 'Fecha', 'Responsable', 'Concepto', 'Monto']
cols_todo = ['ID', 'Proyecto', 'Categoria', 'Tarea', 'Responsable', 'Check']
cols_pagos = ['ID', 'Proyecto', 'Partida', 'Contratista', 'Fecha', 'Monto', 'Tipo', 'Notas'] 
cols_p_info = ['Proyecto', 'Nombre_Partida', 'Monto_Pactado', 'Nivel_Actual', 'Contratista', 'Inicio', 'Fin', 'Estado', 'Etapa', 'Nivel_Inicio', 'Nivel_Fin']
cols_ingresos = ['ID', 'Fecha', 'Proyecto', 'Socio', 'Categoria', 'Monto', 'Notas']
cols_mat = ['ID', 'Proyecto', 'Material', 'Cantidad', 'Unidad', 'Costo_Estimado', 'Costo_Real', 'Proveedor', 'Estado']


# SCHEMA ICE SESSIONS
cols_ice_sesiones = ['ID_Sesion', 'Proyecto', 'Nombre_Sesion', 'Etapa', 'Fecha_Prog', 'Fecha_Real', 'Hora', 'Organizador', 'Lider', 'Coordinador_VDC', 'Supervisor_VDC']
cols_ice_invitados = ['ID_Sesion', 'Nombre', 'Empresa', 'Cargo', 'Asistencia']
cols_ice_acuerdos = ['ID_Sesion', 'Descripcion', 'Responsable', 'Acuerdo']

# SCHEMA MULTIFAMILIAR LIFECYCLE (DESIGN & HANDOVER)
cols_entregables = ['ID', 'Proyecto', 'Especialidad', 'Nombre_Entregable', 'Revision', 'Estado', 'Fecha_Limite', 'Responsable']
cols_observaciones = ['ID', 'Proyecto', 'Nro_Observacion', 'Especialidad', 'Descripcion', 'Estado', 'Fecha_Ingreso', 'Fecha_Subsanacion', 'Responsable', 'Notas']
cols_dossier = ['ID', 'Proyecto', 'Especialidad', 'Documento', 'Estado', 'Responsable', 'Notas']
cols_unidades = ['ID', 'Proyecto', 'Unidad', 'Propietario', 'Estado', 'Fecha_Inspeccion', 'Fecha_Entrega', 'Detalle_Observaciones']


def cargar_tabla(pestaña, columnas):
    if 'original_tables' not in st.session_state:
        st.session_state.original_tables = {}
    if nube_ok and sheet_url:
        try:
            df = conn.read(spreadsheet=sheet_url, worksheet=pestaña, ttl=300)
            if df is None or df.empty:
                # Si una de las tablas clave viene vacía, levantamos la bandera de alerta para evitar sobreescritura accidental
                if pestaña in ['proyectos', 'tasks', 'maestro']:
                    st.session_state.hubo_error_conexion = True
                df_empty = pd.DataFrame(columns=columnas)
                st.session_state.original_tables[pestaña] = df_empty.copy()
                return df_empty
            
            # Migración de columnas antiguas de materiales y mapeo de estados
            if pestaña == 'materiales':
                if 'Ppto_Estimado' in df.columns and 'Costo_Estimado' not in df.columns:
                    df['Costo_Estimado'] = df['Ppto_Estimado']
                if 'Ppto_Real' in df.columns and 'Costo_Real' not in df.columns:
                    df['Costo_Real'] = df['Ppto_Real']
                if 'Estado' in df.columns:
                    df['Estado'] = df['Estado'].replace({
                        'No Comprado': '🔴 No Comprado',
                        'Comprado': '🟢 Comprado'
                    })
                    
            for col in columnas:
                if col not in df.columns: 
                    if col in ['Archivado', 'Check']: df[col] = False
                    elif col == 'Division': df[col] = "General"
                    elif col in ['Meta', 'Tolerancia', 'Valor']: df[col] = 0.0
                    elif col == 'Tipo_Proyecto': df[col] = "Unifamiliar"
                    elif col in ['Num_Pisos', 'Num_Sotanos']: df[col] = 0
                    elif col == 'Estado': df[col] = "🔴 No Comprado" if pestaña == 'materiales' else "No iniciado"
                    elif col == 'Nivel_Actual': df[col] = "No iniciado"
                    elif col == 'Etapa': df[col] = "Subestructura"
                    elif col in ['Nivel_Inicio', 'Nivel_Fin']: df[col] = ""
                    else: df[col] = ""
            
            # Asegurar alineación y selección de columnas para materiales
            if pestaña == 'materiales':
                df = df[columnas]
                
            st.session_state.original_tables[pestaña] = df.copy()
            return df
        except Exception as e:
            # Solo bloqueamos la conexión si falla una tabla crítica (evita bloquear por hojas secundarias faltantes)
            if pestaña in ['proyectos', 'tasks', 'maestro']:
                st.session_state.hubo_error_conexion = True
            df_empty = pd.DataFrame(columns=columnas)
            st.session_state.original_tables[pestaña] = df_empty.copy()
            return df_empty
    df_empty = pd.DataFrame(columns=columnas)
    st.session_state.original_tables[pestaña] = df_empty.copy()
    return df_empty

def guardar_tabla(df, pestaña):
    # 1. Si hubo error de conexión al cargar, bloqueamos cualquier escritura para proteger los datos en la nube
    if st.session_state.get('hubo_error_conexion', False):
        st.error(f"⚠️ Guardado bloqueado para '{pestaña}': Se detectó un fallo de conexión al iniciar la aplicación. Por favor, reintenta la conexión.")
        return False
        
    # 2. Si es una tabla crítica y viene vacía (0 filas), bloqueamos la escritura para evitar pérdidas accidentales
    if pestaña in ['proyectos', 'tasks', 'maestro'] and (df is None or len(df) == 0):
        st.error(f"⚠️ Guardado cancelado para '{pestaña}': Intento de guardar una tabla vacía. Esto previno la pérdida de tus datos.")
        return False
        
    # 3. Evitar peticiones innecesarias si la tabla no ha cambiado desde la última carga o guardado exitoso
    original_df = st.session_state.get('original_tables', {}).get(pestaña)
    if original_df is not None and original_df.equals(df):
        return True
        
    if nube_ok and sheet_url:
        try:
            # Optimización de cuota de API de Google Sheets: intentamos actualizar directamente.
            # Solo si falla la actualización realizamos la validación proactiva y auto-sanación con gspread.
            try:
                conn.update(spreadsheet=sheet_url, worksheet=pestaña, data=df)
                if 'original_tables' in st.session_state:
                    st.session_state.original_tables[pestaña] = df.copy()
                st.cache_data.clear()
                return True
            except Exception as e_update:
                try:
                    import gspread
                    client = conn.client._client
                    spreadsheet = client.open_by_url(sheet_url)
                    try:
                        spreadsheet.worksheet(pestaña)
                    except gspread.exceptions.WorksheetNotFound:
                        # Crear hoja de cálculo vacía
                        spreadsheet.add_worksheet(title=pestaña, rows="100", cols="20")
                    # Reintentar actualización
                    conn.update(spreadsheet=sheet_url, worksheet=pestaña, data=df)
                    if 'original_tables' in st.session_state:
                        st.session_state.original_tables[pestaña] = df.copy()
                    st.cache_data.clear()
                    return True
                except Exception as e_heal:
                    # Si falla la auto-sanación, lanzar el error de actualización original
                    raise e_update
        except Exception as e:
            st.error(f"❌ Error al guardar '{pestaña}': {str(e)}")
            return False
    return False

def forzar_importacion():
    st.session_state.hubo_error_conexion = False  # Reseteamos la bandera al reintentar la conexión
    st.session_state.proyectos = cargar_tabla('proyectos', cols_proy)
    if 'Archivado' in st.session_state.proyectos.columns:
        st.session_state.proyectos['Archivado'] = st.session_state.proyectos['Archivado'].apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y']).astype(bool)
    st.session_state.tasks = cargar_tabla('tasks', cols_tasks)
    st.session_state.maestro = cargar_tabla('maestro', cols_maestro)
    st.session_state.df_historial = cargar_tabla('historial', cols_hist)
    st.session_state.gg_log = cargar_tabla('gg_log', cols_gg)
    df_todo = cargar_tabla('logistica', cols_todo)
    if 'Check' in df_todo.columns:
        df_todo['Check'] = df_todo['Check'].apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y']).astype(bool)
    st.session_state.todo_list = df_todo
    st.session_state.pagos = cargar_tabla('pagos_maestros', cols_pagos)
    st.session_state.p_info = cargar_tabla('partidas_info', cols_p_info)
    st.session_state.ingresos_socios = cargar_tabla('ingresos_socios', cols_ingresos)
    st.session_state.materiales = cargar_tabla('materiales', cols_mat)
    st.session_state.ice_sesiones = cargar_tabla('ice_sesiones', cols_ice_sesiones)
    st.session_state.ice_invitados = cargar_tabla('ice_invitados', cols_ice_invitados)
    st.session_state.ice_acuerdos = cargar_tabla('ice_acuerdos', cols_ice_acuerdos)
    st.session_state.entregables_planos = cargar_tabla('entregables_planos', cols_entregables)
    st.session_state.observaciones_municipales = cargar_tabla('observaciones_municipales', cols_observaciones)
    st.session_state.entrega_dossier = cargar_tabla('entrega_dossier', cols_dossier)
    st.session_state.entrega_unidades = cargar_tabla('entrega_unidades', cols_unidades)

    partidas_list = []
    if not st.session_state.tasks.empty and 'Partida' in st.session_state.tasks.columns:
        df_tasks_p = st.session_state.tasks[['Proyecto', 'Partida']].rename(columns={'Partida': 'Nombre_Partida'})
        partidas_list.append(df_tasks_p)
    if not st.session_state.p_info.empty and 'Nombre_Partida' in st.session_state.p_info.columns:
        df_pinfo_p = st.session_state.p_info[['Proyecto', 'Nombre_Partida']]
        partidas_list.append(df_pinfo_p)
    if partidas_list:
        st.session_state.partidas = pd.concat(partidas_list, ignore_index=True).drop_duplicates().reset_index(drop=True)
    else:
        st.session_state.partidas = pd.DataFrame(columns=['Proyecto', 'Nombre_Partida'])



# CARGA INDIVIDUAL ESTRICTA
if 'hubo_error_conexion' not in st.session_state:
    st.session_state.hubo_error_conexion = False

if 'proyectos' not in st.session_state: 
    df_p = cargar_tabla('proyectos', cols_proy)
    if 'Archivado' in df_p.columns:
        df_p['Archivado'] = df_p['Archivado'].apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y']).astype(bool)
    st.session_state.proyectos = df_p

if 'tasks' not in st.session_state: st.session_state.tasks = cargar_tabla('tasks', cols_tasks)
if 'maestro' not in st.session_state: st.session_state.maestro = cargar_tabla('maestro', cols_maestro)
if 'df_historial' not in st.session_state: st.session_state.df_historial = cargar_tabla('historial', cols_hist)
if 'gg_log' not in st.session_state: st.session_state.gg_log = cargar_tabla('gg_log', cols_gg)
if 'todo_list' not in st.session_state:
    df_todo = cargar_tabla('logistica', cols_todo)
    if 'Check' in df_todo.columns:
        df_todo['Check'] = df_todo['Check'].apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y']).astype(bool)
    st.session_state.todo_list = df_todo
if 'pagos' not in st.session_state: st.session_state.pagos = cargar_tabla('pagos_maestros', cols_pagos)
if 'p_info' not in st.session_state: st.session_state.p_info = cargar_tabla('partidas_info', cols_p_info)
if 'ingresos_socios' not in st.session_state: st.session_state.ingresos_socios = cargar_tabla('ingresos_socios', cols_ingresos)
if 'materiales' not in st.session_state: st.session_state.materiales = cargar_tabla('materiales', cols_mat)
if 'ice_sesiones' not in st.session_state: st.session_state.ice_sesiones = cargar_tabla('ice_sesiones', cols_ice_sesiones)
if 'ice_invitados' not in st.session_state: st.session_state.ice_invitados = cargar_tabla('ice_invitados', cols_ice_invitados)
if 'ice_acuerdos' not in st.session_state: st.session_state.ice_acuerdos = cargar_tabla('ice_acuerdos', cols_ice_acuerdos)
if 'entregables_planos' not in st.session_state: st.session_state.entregables_planos = cargar_tabla('entregables_planos', cols_entregables)
if 'observaciones_municipales' not in st.session_state: st.session_state.observaciones_municipales = cargar_tabla('observaciones_municipales', cols_observaciones)
if 'entrega_dossier' not in st.session_state: st.session_state.entrega_dossier = cargar_tabla('entrega_dossier', cols_dossier)
if 'entrega_unidades' not in st.session_state: st.session_state.entrega_unidades = cargar_tabla('entrega_unidades', cols_unidades)


# Formateadores incondicionales de Caché
if 'Archivado' in st.session_state.proyectos.columns:
    st.session_state.proyectos['Archivado'] = st.session_state.proyectos['Archivado'].apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y']).astype(bool)
if 'Check' in st.session_state.todo_list.columns:
    st.session_state.todo_list['Check'] = st.session_state.todo_list['Check'].apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y']).astype(bool)

if 'partidas' not in st.session_state:
    partidas_list = []
    if not st.session_state.tasks.empty and 'Partida' in st.session_state.tasks.columns:
        df_tasks_p = st.session_state.tasks[['Proyecto', 'Partida']].rename(columns={'Partida': 'Nombre_Partida'})
        partidas_list.append(df_tasks_p)
    if not st.session_state.p_info.empty and 'Nombre_Partida' in st.session_state.p_info.columns:
        df_pinfo_p = st.session_state.p_info[['Proyecto', 'Nombre_Partida']]
        partidas_list.append(df_pinfo_p)
    if partidas_list:
        st.session_state.partidas = pd.concat(partidas_list, ignore_index=True).drop_duplicates().reset_index(drop=True)
    else:
        st.session_state.partidas = pd.DataFrame(columns=['Proyecto', 'Nombre_Partida'])

if 'vista_actual' not in st.session_state: st.session_state.vista_actual = 'main'
if 'proy_activo' not in st.session_state: st.session_state.proy_activo = None
if 'vdc_proy_activo' not in st.session_state: st.session_state.vdc_proy_activo = None
if 'cal_month_offset' not in st.session_state: st.session_state.cal_month_offset = 0

hora_peru = datetime.utcnow() - timedelta(hours=5)
fecha_hoy = pd.to_datetime(hora_peru.date())
df_hoy = pd.DataFrame({'Hoy': [fecha_hoy], 'Etiqueta': ['DIA ACTUAL']})

# ==========================================
# 3. LÓGICAS AI Y CÁLCULOS TÉCNICOS
# ==========================================
client_ia = None
ai_conectada = False

def calcular_rating(nombre):
    base = 5
    penal_hist = len(st.session_state.df_historial[(st.session_state.df_historial['ID_Contratista'] == nombre) & (st.session_state.df_historial['Estado'].astype(str).str.contains('retraso', case=False, na=False))])
    penal_app = len(st.session_state.tasks[(st.session_state.tasks['Contratista'] == nombre) & (st.session_state.tasks['Estado'].astype(str).str.contains('retraso', case=False, na=False)) & (st.session_state.tasks['Culpa_Contratista'] == True)])
    notas = st.session_state.maestro.loc[st.session_state.maestro['ID_Contratista'] == nombre, 'Notas'].values
    notas_str = notas[0] if len(notas) > 0 and pd.notna(notas[0]) else ""
    rating_final = base - (penal_hist + penal_app) + (notas_str.count("[+]") * 0.5) - (notas_str.count("[-]") * 1.0)
    return max(1, min(5, int(rating_final)))

def extraer_esp():
    raw = st.session_state.maestro['Especialidad'].dropna().astype(str).tolist()
    todas = set([limpiar_especialidad(sub) for e in raw for sub in e.split(',') if limpiar_especialidad(sub)])
    return ["Ninguna"] + [e.replace("Albanileria", "Albañilería") for e in sorted(list(todas))]

def sugerir_contratista(esp_req, ubi_proy, req_i, req_f):
    if not esp_req or esp_req == "Ninguna": return None
    maestro = st.session_state.maestro.copy()
    candidatos = maestro[maestro['Especialidad'].apply(lambda x: limpiar_especialidad(x)).str.contains(limpiar_especialidad(esp_req), case=False, na=False)].copy()
    if candidatos.empty: return f"Alerta: Sin especialistas en '{esp_req}'."
    
    candidatos['Puntaje'] = 0
    if pd.notna(req_i) and pd.notna(req_f):
        for idx, r in candidatos.iterrows():
            c_n = r['ID_Contratista']
            t_c = st.session_state.tasks[(st.session_state.tasks['Contratista'] == c_n) & (st.session_state.tasks['Permitir_Cruce'] == False)]
            for _, t in t_c.iterrows():
                try:
                    t_i = pd.to_datetime(t['Inicio']).date() if pd.notna(t['Inicio']) else None
                    t_f = pd.to_datetime(t['Fin']).date() if pd.notna(t['Fin']) else None
                    if t_i and t_f and req_i <= t_f and t_i <= req_f:
                        candidatos.at[idx, 'Puntaje'] -= 1000 
                        break
                except: pass
                    
    candidatos = candidatos[candidatos['Puntaje'] >= 0]
    if candidatos.empty: return f"Alerta: Especialistas saturados en fecha seleccionada."
    
    candidatos.loc[candidatos['Ubicación_Base'] == ubi_proy, 'Puntaje'] += 2
    candidatos['Estrellas'] = [calcular_rating(n) for n in candidatos['ID_Contratista']]
    candidatos['Puntaje'] += candidatos['Estrellas']
    
    mejor = candidatos.sort_values(by='Puntaje', ascending=False).iloc[0]
    return f"Sugerencia {st.secrets.get('APP_TITLE', 'COTA APP')}: {mejor['ID_Contratista']} (Rating: {mejor['Estrellas']}/5)."

def generar_tabla_cife(df):
    if df.empty:
        return ""
    
    html = """
    <div style="overflow-x: auto; margin: 15px 0 25px 0;">
        <table style="width: 100%; border-collapse: collapse; text-align: left; font-family: 'Outfit', 'Inter', sans-serif; font-size: 13px; color: var(--text-muted); background-color: var(--custom-bg); box-shadow: var(--card-shadow); border-radius: 4px; overflow: hidden;">
            <thead>
                <tr style="border-top: 3px solid #00E5FF; border-bottom: 2px solid var(--border-color); background-color: var(--card-bg);">
                    <th style="padding: 14px 12px; color: #00E5FF; font-weight: 800; text-transform: uppercase; width: 10%; text-align: center;">Pilar</th>
                    <th style="padding: 14px 12px; color: #00E5FF; font-weight: 800; text-transform: uppercase; width: 8%; text-align: center;">Clasif.</th>
                    <th style="padding: 14px 12px; color: #00E5FF; font-weight: 800; text-transform: uppercase; width: 34%;">Description</th>
                    <th style="padding: 14px 12px; color: #00E5FF; font-weight: 800; text-transform: uppercase; width: 24%;">Metric</th>
                    <th style="padding: 14px 12px; color: #00E5FF; font-weight: 800; text-transform: uppercase; width: 12%;">Target Value</th>
                    <th style="padding: 14px 12px; color: #00E5FF; font-weight: 800; text-transform: uppercase; width: 12%;">Source</th>
                </tr>
            </thead>
            <tbody>
    """
    
    pilares_order = ["PPM", "ICE", "BIM"]
    rows_added = 0
    
    for pilar in pilares_order:
        df_p = df[df['Pilar'] == pilar]
        if df_p.empty:
            continue
            
        pilar_rowspan = len(df_p)
        first_row = True
        
        # Sort so PO is first, then CF
        df_p = df_p.copy()
        df_p['sort_idx'] = df_p['Clasificacion'].apply(lambda x: 0 if 'PO' in str(x) else 1)
        df_p = df_p.sort_values('sort_idx')
        
        for idx, row in df_p.iterrows():
            lbl_c = str(row['Clasificacion'])[:2]
            
            # Format target value beautifully
            cond = row['Condicion']
            
            try:
                meta = float(row['Meta']) if pd.notna(row['Meta']) and row['Meta'] is not None else 0.0
            except:
                meta = 0.0
                
            try:
                tol = float(row.get('Tolerancia', 0.0)) if pd.notna(row.get('Tolerancia', 0.0)) and row.get('Tolerancia', 0.0) is not None else 0.0
            except:
                tol = 0.0
                
            unidad = row.get('Unidad', '')
            
            # Formatting meta and tol values nicely (remove trailing .0 if integer-like)
            meta_str = f"{int(meta)}" if isinstance(meta, (int, float)) and hasattr(meta, 'is_integer') and meta.is_integer() else f"{meta:.2f}"
            tol_str = f"{int(tol)}" if isinstance(tol, (int, float)) and hasattr(tol, 'is_integer') and tol.is_integer() else f"{tol:.2f}"
            
            if cond == '±':
                target_str = f"{meta_str} {unidad} ± {tol_str}"
            else:
                target_str = f"{cond}{meta_str} {unidad}"
                
            # Row style
            row_border = "border-bottom: 1px solid var(--border-color);"
            bg_color = "var(--custom-bg)" if rows_added % 2 == 0 else "var(--card-bg)"
            
            html += f'<tr style="background-color: {bg_color}; {row_border}">'
            
            if first_row:
                html += f'''
                    <td rowspan="{pilar_rowspan}" style="padding: 14px 12px; font-weight: 800; color: var(--text-color); text-align: center; vertical-align: middle; border-right: 1px solid var(--border-color); background-color: var(--card-bg); font-family: 'Outfit', sans-serif; font-size: 16px; letter-spacing: 1px; border-bottom: 1px solid var(--border-color);">
                        {pilar}
                    </td>
                '''
                first_row = False
                
            html += f'''
                <td style="padding: 14px 12px; font-weight: bold; color: #00E5FF; font-family: monospace; font-size: 13px; text-align: center; vertical-align: middle;">{lbl_c}</td>
                <td style="padding: 14px 12px; color: var(--text-muted); vertical-align: middle; line-height: 1.4; font-weight: 400;">{row['Descripcion']}</td>
                <td style="padding: 14px 12px; color: var(--text-secondary); vertical-align: middle; line-height: 1.4; font-weight: 400;">{row['Nombre']}</td>
                <td style="padding: 14px 12px; color: #00E5FF; font-weight: 700; font-family: 'Outfit', sans-serif; vertical-align: middle;">{target_str}</td>
                <td style="padding: 14px 12px; color: var(--text-muted); font-family: monospace; vertical-align: middle;">{row['Fuente']}</td>
            </tr>
            '''
            rows_added += 1
            
    html += """
            </tbody>
        </table>
    </div>
    <style>
    /* Styling to match CIFE academic style but dynamically */
    table tbody tr:last-child {
        border-bottom: 3px solid #00E5FF !important;
    }
    </style>
    """
    cleaned_html = "\n".join([line.strip() for line in html.split("\n")])
    return cleaned_html



def generar_pdf_sesion(r_ses, invitados_df, acuerdos_df):
    from fpdf import FPDF
    from fpdf.fonts import FontFace
    
    def safe_str(text):
        if not text or pd.isna(text):
            return ""
        return str(text).encode('latin-1', 'replace').decode('latin-1')
        
    class PDF(FPDF):
        def header(self):
            # Header color block (#1e293b)
            self.set_fill_color(30, 41, 59)
            self.rect(0, 0, 210, 28, 'F')
            
            # Title text (#00E5FF)
            self.set_text_color(0, 229, 255)
            self.set_font('Helvetica', 'B', 15)
            self.set_xy(10, 10)
            self.cell(190, 8, safe_str(self.pdf_title).upper(), border=0, align='C')
            
            # Logo image
            import os
            if os.path.exists("Logo2.png"):
                try:
                    self.image("Logo2.png", x=175, y=4, w=25)
                except:
                    pass
            self.ln(18)
            
        def footer(self):
            self.set_y(-15)
            self.set_font('Helvetica', 'I', 8)
            self.set_text_color(128, 128, 128)
            self.cell(0, 10, f'Página {self.page_no()}/{{nb}}', align='C')
            
    pdf = PDF()
    pdf.pdf_title = r_ses['Nombre_Sesion']
    pdf.alias_nb_pages()
    pdf.add_page()
    pdf.set_margins(15, 32, 15)
    pdf.set_y(32)
    
    # 1. Datos del proyecto
    pdf.set_fill_color(17, 27, 48)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(0, 7, '  DATOS DEL PROYECTO', fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)
    
    pdf.set_text_color(0, 0, 0)
    pdf.set_font('Helvetica', '', 9)
    pdf.cell(45, 6, 'Proyecto:', border='B')
    pdf.set_font('Helvetica', 'B', 9)
    pdf.cell(0, 6, safe_str(r_ses['Proyecto']), border='B', new_x="LMARGIN", new_y="NEXT")
    
    pdf.set_font('Helvetica', '', 9)
    pdf.cell(45, 6, 'Etapa del proyecto:', border='B')
    pdf.set_font('Helvetica', 'B', 9)
    pdf.cell(0, 6, safe_str(r_ses['Etapa']), border='B', new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    
    # 2. Datos de la sesión
    pdf.set_fill_color(17, 27, 48)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(0, 7, '  DATOS DE LA SESIÓN', fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)
    
    pdf.set_text_color(0, 0, 0)
    
    session_fields = [
        ('Fecha de programación:', r_ses['Fecha_Prog']),
        ('Fecha de sesión:', r_ses['Fecha_Real']),
        ('Hora:', r_ses['Hora']),
        ('Organizador:', r_ses['Organizador'])
    ]
    for label, val in session_fields:
        pdf.set_font('Helvetica', '', 9)
        pdf.cell(45, 6, label, border='B')
        pdf.set_font('Helvetica', 'B', 9)
        pdf.cell(0, 6, safe_str(val), border='B', new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    
    # 3. Encargado
    pdf.set_fill_color(17, 27, 48)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(0, 7, '  ENCARGADO', fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)
    
    pdf.set_text_color(0, 0, 0)
    
    encargados_fields = [
        ('Líder:', r_ses['Lider'])
    ]
    for label, val in encargados_fields:
        pdf.set_font('Helvetica', '', 9)
        pdf.cell(45, 6, label, border='B')
        pdf.set_font('Helvetica', 'B', 9)
        pdf.cell(0, 6, safe_str(val), border='B', new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    
    # 4. Invitados
    pdf.set_fill_color(17, 27, 48)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(0, 7, '  INVITADOS Y ASISTENCIA', fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    
    hs = FontFace(emphasis="B", fill_color=(255, 255, 255), color=(0, 0, 0))
    
    if invitados_df.empty:
        pdf.set_text_color(0, 0, 0)
        pdf.set_font('Helvetica', 'I', 9)
        pdf.cell(0, 6, 'No hay invitados registrados.', border=1, new_x="LMARGIN", new_y="NEXT")
    else:
        pdf.set_text_color(0, 0, 0)
        pdf.set_fill_color(255, 255, 255)
        with pdf.table(col_widths=(65, 45, 45, 25), line_height=5, headings_style=hs, borders_layout="NO_HORIZONTAL_LINES", cell_fill_mode="NONE") as table:
            headers = table.row()
            pdf.set_font('Helvetica', 'B', 9)
            headers.cell(safe_str("Nombre"))
            headers.cell(safe_str("Empresa"))
            headers.cell(safe_str("Cargo / Especialidad"))
            headers.cell(safe_str("Asistencia"))
            
            pdf.set_font('Helvetica', '', 9)
            for _, row in invitados_df.iterrows():
                r = table.row()
                r.cell(safe_str(row.get('Nombre', '')))
                r.cell(safe_str(row.get('Empresa', '')))
                r.cell(safe_str(row.get('Cargo', '')))
                r.cell(safe_str(row.get('Asistencia', 'Asistió')))
    pdf.ln(6)
    
    # 5. Acuerdos
    pdf.set_fill_color(17, 27, 48)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(0, 7, '  ACTIVIDADES Y ACUERDOS', fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    
    if acuerdos_df.empty:
        pdf.set_text_color(0, 0, 0)
        pdf.set_font('Helvetica', 'I', 9)
        pdf.cell(0, 6, 'No hay acuerdos registrados.', border=1, new_x="LMARGIN", new_y="NEXT")
    else:
        pdf.set_text_color(0, 0, 0)
        pdf.set_fill_color(255, 255, 255)
        with pdf.table(col_widths=(75, 40, 65), line_height=5, headings_style=hs, borders_layout="NO_HORIZONTAL_LINES", cell_fill_mode="NONE") as table:
            headers = table.row()
            pdf.set_font('Helvetica', 'B', 9)
            headers.cell(safe_str("Descripción / Actividad"))
            headers.cell(safe_str("Responsable"))
            headers.cell(safe_str("Acuerdo"))
            
            pdf.set_font('Helvetica', '', 9)
            for _, row in acuerdos_df.iterrows():
                r = table.row()
                r.cell(safe_str(row.get('Descripcion', '')))
                r.cell(safe_str(row.get('Responsable', '')))
                r.cell(safe_str(row.get('Acuerdo', '')))
                
    return bytes(pdf.output())


# Helper to execute lightweight reruns inside fragments if supported
if hasattr(st, "fragment"):
    fragment_decorator = st.fragment
elif hasattr(st, "experimental_fragment"):
    fragment_decorator = st.experimental_fragment
else:
    def fragment_decorator(func):
        return func

def clear_base_dfs():
    for k in list(st.session_state.keys()):
        if k.startswith("base_df_"):
            del st.session_state[k]

def fragment_rerun():
    clear_base_dfs()
    try:
        st.rerun(scope="fragment")
    except:
        st.rerun()

def on_table_edit(base_key, widget_key, master_state_key, group_cols=None, group_vals=None, date_cols=None, text_cols=None, bool_cols=None, num_cols=None, id_col=None):
    if widget_key not in st.session_state:
        return
    edits = st.session_state[widget_key]
    df = st.session_state[base_key]
    
    # 1. Apply edits to the base dataframe
    df_new = df.copy()
    for idx, cols in edits.get("edited_rows", {}).items():
        row_idx = int(idx)
        for col, val in cols.items():
            df_new.at[row_idx, col] = val
            
    # 2. Apply deletions
    deleted_indices = [int(i) for i in edits.get("deleted_rows", [])]
    if deleted_indices:
        df_new = df_new.drop(index=deleted_indices).reset_index(drop=True)
        
    # 3. Apply additions
    added_rows = edits.get("added_rows", [])
    if added_rows:
        df_added = pd.DataFrame(added_rows)
        df_new = pd.concat([df_new, df_added], ignore_index=True)
        
    # 4. Generate unique IDs if necessary
    if id_col:
        if id_col not in df_new.columns:
            df_new[id_col] = None
        for i, row in df_new.iterrows():
            val = row.get(id_col)
            if pd.isna(val) or val == "" or str(val).strip() == "" or str(val).lower() == "none":
                if master_state_key == 'entregables_planos':
                    import uuid
                    df_new.at[i, id_col] = f"ENT_{int(datetime.now().timestamp())}_{uuid.uuid4().hex[:6]}"
                elif master_state_key == 'observaciones_municipales':
                    import uuid
                    df_new.at[i, id_col] = f"OBS_{int(datetime.now().timestamp())}_{uuid.uuid4().hex[:6]}"
                elif id_col in ['ID', 'ID_Log']:
                    import time
                    df_new.at[i, id_col] = time.time() + i * 0.001
                else:
                    import uuid
                    prefix = id_col.split("_")[0] if "_" in id_col else id_col
                    df_new.at[i, id_col] = f"{prefix}_{int(datetime.now().timestamp())}_{uuid.uuid4().hex[:6]}"
                    
    # 5. Sanitization of dates/types to prevent ArrowTypeError and DateColumn crashes
    if date_cols:
        for c in date_cols:
            if c in df_new.columns:
                df_new[c] = pd.to_datetime(df_new[c], errors='coerce').dt.date
    if text_cols:
        for c in text_cols:
            if c in df_new.columns:
                df_new[c] = df_new[c].fillna("").astype(str).replace('nan', '').replace('None', '')
    if bool_cols:
        for c in bool_cols:
            if c in df_new.columns:
                df_new[c] = df_new[c].fillna(False).apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y', 'yes', 'si']).astype(bool)
    if num_cols:
        for c in num_cols:
            if c in df_new.columns:
                df_new[c] = pd.to_numeric(df_new[c], errors='coerce').fillna(0.0)
                
    # Update base dataframe in session state
    st.session_state[base_key] = df_new
    
    # 6. Update master dataframe in session state
    if master_state_key in st.session_state:
        master_df = st.session_state[master_state_key]
        if group_cols and group_vals:
            clean_mask = pd.Series(True, index=master_df.index)
            for g_col, g_val in zip(group_cols, group_vals):
                clean_mask = clean_mask & (master_df[g_col] == g_val)
            master_df = master_df[~clean_mask]
        else:
            master_df = pd.DataFrame(columns=master_df.columns)
            
        n_t = df_new.copy()
        if group_cols and group_vals:
            for g_col, g_val in zip(group_cols, group_vals):
                n_t[g_col] = g_val
                
        cols_master = master_df.columns.tolist()
        for c in cols_master:
            if c not in n_t.columns:
                n_t[c] = None
        st.session_state[master_state_key] = pd.concat([master_df, n_t[cols_master]], ignore_index=True)
        
    # 7. Increment counter to force new widget key and avoid StreamlitValueAssignmentNotAllowedError
    c_key = f"c_{base_key}"
    st.session_state[c_key] = st.session_state.get(c_key, 0) + 1
        


def extraer_texto_documento(uploaded_file):
    """Extrae texto de archivos .txt, .md, .docx, .pdf"""
    if uploaded_file is None:
        return ""
    name = uploaded_file.name.lower()
    try:
        if name.endswith(('.txt', '.md')):
            return uploaded_file.getvalue().decode('utf-8', errors='replace')
        elif name.endswith('.docx'):
            try:
                import docx
                doc = docx.Document(io.BytesIO(uploaded_file.getvalue()))
                full_text = []
                for p in doc.paragraphs:
                    if p.text.strip():
                        full_text.append(p.text)
                for t in doc.tables:
                    for row in t.rows:
                        row_text = [c.text.strip() for c in row.cells if c.text.strip()]
                        if row_text:
                            full_text.append(" | ".join(row_text))
                return "\n".join(full_text)
            except Exception:
                return uploaded_file.getvalue().decode('latin-1', errors='replace')
        elif name.endswith('.pdf'):
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(uploaded_file.getvalue()))
                text_pages = []
                for page in reader.pages:
                    t = page.extract_text()
                    if t:
                        text_pages.append(t)
                return "\n\n".join(text_pages)
            except Exception:
                return ""
        else:
            return uploaded_file.getvalue().decode('utf-8', errors='replace')
    except Exception:
        return ""

def _parse_resumen_meet_heuristico(texto, proyecto, count=0):
    lines = [l.strip() for l in str(texto).splitlines() if l.strip()]
    num_s = count + 1
    num_str = f"0{num_s}" if num_s < 10 else str(num_s)
    
    fecha_actual_str = datetime.now().strftime("%d/%m/%Y")
    
    m_fecha = re.search(r'(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})', texto)
    if m_fecha:
        d, m, y = m_fecha.groups()
        if len(y) == 2: y = "20" + y
        fecha_str = f"{int(d):02d}/{int(m):02d}/{y}"
    else:
        fecha_str = fecha_actual_str
        
    m_hora = re.search(r'(\d{1,2}:\d{2}\s*(?:a\.?\s*m\.?|p\.?\s*m\.?|AM|PM)?)', texto, re.IGNORECASE)
    hora_str = m_hora.group(1) if m_hora else "3:00 p. m."
    
    invitados = []
    acuerdos = []
    
    modo = None
    for line in lines:
        l_lower = line.lower()
        if any(h in l_lower for h in ['asistentes', 'participantes', 'attendees']):
            modo = 'invitados'
            continue
        elif any(h in l_lower for h in ['próximos pasos', 'proximos pasos', 'action items', 'tareas', 'acuerdos', 'decisiones', 'compromisos']):
            modo = 'acuerdos'
            continue
        elif any(h in l_lower for h in ['resumen', 'overview', 'temas tratados']):
            modo = 'resumen'
            continue
            
        if modo == 'invitados':
            clean_l = re.sub(r'^[•\-\*\d\.\)]\s*', '', line).strip()
            if clean_l and not clean_l.endswith(':'):
                partes = clean_l.split('(')
                nom = partes[0].strip()
                rol = partes[1].replace(')', '').strip() if len(partes) > 1 else "Participante"
                if nom:
                    invitados.append({
                        "nombre": nom,
                        "empresa": "",
                        "cargo": rol,
                        "asistencia": "Asistió"
                    })
        elif modo == 'acuerdos':
            clean_l = re.sub(r'^[•\-\*\d\.\)]\s*', '', line).strip()
            if clean_l and not clean_l.endswith(':'):
                resp = "Equipo del Proyecto"
                act = clean_l
                if ':' in clean_l:
                    p = clean_l.split(':', 1)
                    if len(p[0].split()) <= 4:
                        resp = p[0].strip()
                        act = p[1].strip()
                elif '-' in clean_l and len(clean_l.split('-')[0].split()) <= 4:
                    p = clean_l.split('-', 1)
                    resp = p[0].strip()
                    act = p[1].strip()
                acuerdos.append({
                    "descripcion": "Seguimiento y acuerdos de reunión",
                    "responsable": resp,
                    "acuerdo": act
                })
                
    if not invitados:
        invitados.append({"nombre": "Equipo del Proyecto", "empresa": "", "cargo": "Participante", "asistencia": "Asistió"})
    if not acuerdos:
        acuerdos.append({"descripcion": "Revisión general de avances", "responsable": "Equipo del Proyecto", "acuerdo": "Revisar acuerdos en la siguiente sesión"})
        
    return {
        "nombre_sesion": f"Agenda Sesión ICE N-{num_str} - {proyecto}",
        "etapa": "Diseño",
        "fecha_prog": fecha_str,
        "fecha_real": fecha_str,
        "hora": hora_str,
        "organizador": invitados[0]["nombre"] if invitados else "",
        "lider": invitados[0]["nombre"] if invitados else "",
        "invitados": invitados,
        "acuerdos": acuerdos,
        "resumen_ejecutivo": "Sesión registrada a partir del resumen de reunión de Google Meet."
    }

def procesar_resumen_meet_gemini(texto_resumen, nombre_proyecto, sesiones_count=0):
    """
    Analiza un resumen de Google Meet (generado por Gemini / 'Tomar notas por mí')
    y extrae estructuradamente todos los datos para la sesión ICE.
    """
    if not texto_resumen or not str(texto_resumen).strip():
        return None

    if ai_conectada and client_ia:
        num_s = sesiones_count + 1
        num_str = f"0{num_s}" if num_s < 10 else str(num_s)
        prompt_ia = f"""Eres un asistente experto en metodología VDC (Virtual Design and Construction) y Sesiones ICE (Integrated Concurrent Engineering) en construcción e ingeniería.
Analiza el siguiente texto correspondiente a una minuta o resumen de reunión de Google Meet (generada por Google Gemini / 'Tomar notas por mí' o notas de reunión) del proyecto "{nombre_proyecto}".
El proyecto tiene actualmente {sesiones_count} sesiones registradas.

Tu tarea es estructurar la información en un objeto JSON VÁLIDO con los datos necesarios para registrar la Sesión ICE en el sistema.

FORMATO JSON OBLIGATORIO (responde ÚNICAMENTE con el objeto JSON, sin formato markdown ```json, sin comentarios):
{{
  "nombre_sesion": "Agenda Sesión ICE N-{num_str} - [Tema principal de la reunión]",
  "etapa": "Diseño",
  "fecha_prog": "DD/MM/YYYY",
  "fecha_real": "DD/MM/YYYY",
  "hora": "3:00 p. m.",
  "organizador": "Nombre del moderador u organizador",
  "lider": "Nombre del líder del proyecto o coordinador VDC",
  "invitados": [
    {{
      "nombre": "Nombre completo",
      "empresa": "Empresa o especialidad (ej. Arquitectura, Estructuras, Venti, Cliente, etc.)",
      "cargo": "Cargo o rol (ej. Modelador BIM, Coordinador VDC, Ing. Sanitario, o Participante)",
      "asistencia": "Asistió"
    }}
  ],
  "acuerdos": [
    {{
      "descripcion": "Tema, entregable o actividad tratada (ej. Trazado de Sanitarias en Sótano 1, Entrega de RFI-02)",
      "responsable": "Nombre del responsable o especialidad",
      "acuerdo": "Tarea o acuerdo acordado"
    }}
  ],
  "resumen_ejecutivo": "Síntesis breve de 2 oraciones con los puntos clave tratados."
}}

REGLAS DE EXTRACCIÓN:
1. Extrae a todos los participantes mencionados en 'invitados'. Si alguien no pudo asistir o se excusó, pon 'Faltó' en 'asistencia', de lo contrario 'Asistió'.
2. Convierte TODOS los 'Action items', 'Próximos pasos', 'Tareas', 'Decisiones' y 'Acuerdos' en filas dentro de 'acuerdos', identificando con precisión a la persona asignada en 'responsable' y el compromiso en 'acuerdo'.
3. 'etapa' debe ser estrictamente una de: "Diseño", "Construcción", "Planificación", "Cierre" (inferir del contexto o por defecto "Diseño").
4. Formato de fechas: DD/MM/YYYY.

TEXTO DEL RESUMEN DE GOOGLE MEET:
\"\"\"
{texto_resumen}
\"\"\"
"""
        for mod_name in ['gemini-flash-latest', 'gemini-2.5-flash', 'gemini-1.5-flash']:
            try:
                response = client_ia.models.generate_content(
                    model=mod_name,
                    contents=prompt_ia,
                )
                raw_text = response.text.strip()
                if "```" in raw_text:
                    raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text, flags=re.MULTILINE)
                    raw_text = re.sub(r"```\s*$", "", raw_text, flags=re.MULTILINE).strip()
                
                start = raw_text.find("{")
                end = raw_text.rfind("}")
                if start != -1 and end != -1:
                    json_clean = raw_text[start:end+1]
                    data = json.loads(json_clean)
                    return data
            except Exception:
                continue

    return _parse_resumen_meet_heuristico(texto_resumen, nombre_proyecto, sesiones_count)



@fragment_decorator
def render_gestor_ice_sesiones(v_proy):
    st.markdown("### 📅 AGENDA Y REGISTRO DE SESIONES ICE")
    logo_src = get_image_base64("Logo2.png")
    
    # Load sheets if not initialized (fallback verification)
    if 'ice_sesiones' not in st.session_state:
        st.session_state.ice_sesiones = pd.DataFrame(columns=cols_ice_sesiones)
    if 'ice_invitados' not in st.session_state:
        st.session_state.ice_invitados = pd.DataFrame(columns=cols_ice_invitados)
    if 'ice_acuerdos' not in st.session_state:
        st.session_state.ice_acuerdos = pd.DataFrame(columns=cols_ice_acuerdos)
        
    proj_sessions = st.session_state.ice_sesiones[st.session_state.ice_sesiones['Proyecto'] == v_proy]
    
    c_act1, c_act2 = st.columns([2.2, 4.8])
    
    with c_act1:
        # Popover to register new session
        with st.popover("➕ REGISTRAR NUEVA SESIÓN", use_container_width=True):
            tab_gemini_ice, tab_manual_ice = st.tabs(["✨ GOOGLE MEET (GEMINI)", "✍️ REGISTRO MANUAL"])
            
            with tab_gemini_ice:
                st.markdown("<p style='font-size: 13px; color: #cbd5e1; margin-bottom: 8px;'>Sube el resumen/minuta generado por <b>Gemini en Google Meet</b> (o notas de reunión) para estructurar y registrar automáticamente la sesión con todos sus asistentes y acuerdos.</p>", unsafe_allow_html=True)
                
                f_meet = st.file_uploader(
                    "Subir archivo de resumen (.txt, .docx, .pdf, .md)",
                    type=['txt', 'docx', 'pdf', 'md'],
                    key=f"f_meet_{v_proy}",
                    help="Documento de notas o transcripción de Google Meet"
                )
                
                txt_meet = st.text_area(
                    "O pega el texto del resumen de Gemini / Meet:",
                    placeholder="Resumen de la reunión: Sesión ICE...\nAsistentes: Juan Perez, Maria Gomez...\nPróximos pasos / Tareas:\n- Juan Perez: Revisar trazado de sanitarias...",
                    height=130,
                    key=f"txt_meet_{v_proy}"
                )
                
                if st.button("🤖 PROCESAR Y AUTO-REGISTRAR SESIÓN", type="primary", use_container_width=True, key=f"btn_proc_meet_{v_proy}"):
                    texto_a_procesar = ""
                    if f_meet is not None:
                        texto_a_procesar = extraer_texto_documento(f_meet)
                    elif txt_meet and txt_meet.strip():
                        texto_a_procesar = txt_meet.strip()
                        
                    if not texto_a_procesar:
                        st.error("Por favor, sube un archivo o pega el texto del resumen de la reunión.")
                    else:
                        with st.spinner("⚡ Spicy Gala & Gemini están analizando el resumen de Google Meet..."):
                            data_ext = procesar_resumen_meet_gemini(texto_a_procesar, v_proy, len(proj_sessions))
                            
                        if not data_ext:
                            st.error("No se pudo estructurar la información del documento. Revisa el contenido e inténtalo nuevamente.")
                        else:
                            new_id = f"SES_{int(datetime.now().timestamp())}_{uuid.uuid4().hex[:6]}"
                            n_nom = data_ext.get('nombre_sesion', f"Agenda Sesión ICE N-0{len(proj_sessions) + 1}")
                            n_etp = data_ext.get('etapa', 'Diseño')
                            if n_etp not in ["Diseño", "Construcción", "Planificación", "Cierre"]:
                                n_etp = "Diseño"
                                
                            new_session = pd.DataFrame([{
                                'ID_Sesion': new_id,
                                'Proyecto': v_proy,
                                'Nombre_Sesion': n_nom,
                                'Etapa': n_etp,
                                'Fecha_Prog': data_ext.get('fecha_prog', datetime.today().strftime("%d/%m/%Y")),
                                'Fecha_Real': data_ext.get('fecha_real', datetime.today().strftime("%d/%m/%Y")),
                                'Hora': data_ext.get('hora', '3:00 p. m.'),
                                'Organizador': data_ext.get('organizador', ''),
                                'Lider': data_ext.get('lider', ''),
                                'Coordinador_VDC': "",
                                'Supervisor_VDC': ""
                            }])
                            st.session_state.ice_sesiones = pd.concat([st.session_state.ice_sesiones, new_session], ignore_index=True)
                            
                            # Invitados extraídos
                            invitados_list = data_ext.get('invitados', [])
                            if invitados_list:
                                new_guests_rows = []
                                for inv in invitados_list:
                                    nom_inv = inv.get('nombre', '').strip()
                                    if nom_inv:
                                        new_guests_rows.append({
                                            'ID_Sesion': new_id,
                                            'Nombre': nom_inv,
                                            'Empresa': inv.get('empresa', ''),
                                            'Cargo': inv.get('cargo', ''),
                                            'Asistencia': inv.get('asistencia', 'Asistió')
                                        })
                                if new_guests_rows:
                                    st.session_state.ice_invitados = pd.concat([st.session_state.ice_invitados, pd.DataFrame(new_guests_rows)], ignore_index=True)
                                else:
                                    st.session_state.ice_invitados = pd.concat([st.session_state.ice_invitados, pd.DataFrame([{'ID_Sesion': new_id, 'Nombre': '', 'Empresa': '', 'Cargo': '', 'Asistencia': 'Asistió'}])], ignore_index=True)
                            else:
                                st.session_state.ice_invitados = pd.concat([st.session_state.ice_invitados, pd.DataFrame([{'ID_Sesion': new_id, 'Nombre': '', 'Empresa': '', 'Cargo': '', 'Asistencia': 'Asistió'}])], ignore_index=True)
                                
                            # Acuerdos y tareas extraídos
                            acuerdos_list = data_ext.get('acuerdos', [])
                            if acuerdos_list:
                                new_acu_rows = []
                                for acu in acuerdos_list:
                                    desc_acu = acu.get('descripcion', '').strip()
                                    resp_acu = acu.get('responsable', '').strip()
                                    texto_acu = acu.get('acuerdo', '').strip()
                                    if desc_acu or resp_acu or texto_acu:
                                        new_acu_rows.append({
                                            'ID_Sesion': new_id,
                                            'Descripcion': desc_acu,
                                            'Responsable': resp_acu,
                                            'Acuerdo': texto_acu
                                        })
                                if new_acu_rows:
                                    st.session_state.ice_acuerdos = pd.concat([st.session_state.ice_acuerdos, pd.DataFrame(new_acu_rows)], ignore_index=True)
                                else:
                                    st.session_state.ice_acuerdos = pd.concat([st.session_state.ice_acuerdos, pd.DataFrame([{'ID_Sesion': new_id, 'Descripcion': '', 'Responsable': '', 'Acuerdo': ''}])], ignore_index=True)
                            else:
                                st.session_state.ice_acuerdos = pd.concat([st.session_state.ice_acuerdos, pd.DataFrame([{'ID_Sesion': new_id, 'Descripcion': '', 'Responsable': '', 'Acuerdo': ''}])], ignore_index=True)
                                
                            st.session_state[f"sel_sesion_{v_proy}"] = new_id
                            st.success(f"🎉 Sesión '{n_nom}' registrada automáticamente.")
                            fragment_rerun()
            
            with tab_manual_ice:
                with st.form(f"form_nueva_sesion_ice_{v_proy}", clear_on_submit=True):
                    st.markdown("##### NUEVA SESIÓN ICE MANUAL")
                    n_nombre = st.text_input("Nombre de la Sesión", value=f"Agenda Sesión ICE N-0{len(proj_sessions) + 1}" if len(proj_sessions) < 9 else f"Agenda Sesión ICE N-{len(proj_sessions) + 1}", key=f"ice_n_nombre_{v_proy}")
                    n_etapa = st.selectbox("Etapa del Proyecto", ["Diseño", "Construcción", "Planificación", "Cierre"], index=0, key=f"ice_n_etapa_{v_proy}")
                    n_fecha_prog = st.date_input("Fecha de Programación", value=datetime.today(), key=f"ice_n_fecha_prog_{v_proy}")
                    n_fecha_real = st.date_input("Fecha de Sesión", value=datetime.today(), key=f"ice_n_fecha_real_{v_proy}")
                    n_hora = st.text_input("Hora", value="3:00 p. m.", key=f"ice_n_hora_{v_proy}")
                    n_organizador = st.text_input("Organizador", value="", key=f"ice_n_organizador_{v_proy}")
                    n_lider = st.text_input("Líder", value="", key=f"ice_n_lider_{v_proy}")
                    
                    if st.form_submit_button("CREAR SESIÓN"):
                        if not n_nombre.strip():
                            st.error("El nombre de la sesión es requerido.")
                        else:
                            new_id = f"SES_{int(datetime.now().timestamp())}_{uuid.uuid4().hex[:6]}"
                            new_session = pd.DataFrame([{
                                'ID_Sesion': new_id,
                                'Proyecto': v_proy,
                                'Nombre_Sesion': n_nombre,
                                'Etapa': n_etapa,
                                'Fecha_Prog': n_fecha_prog.strftime("%d/%m/%Y"),
                                'Fecha_Real': n_fecha_real.strftime("%d/%m/%Y"),
                                'Hora': n_hora,
                                'Organizador': n_organizador,
                                'Lider': n_lider,
                                'Coordinador_VDC': "",
                                'Supervisor_VDC': ""
                            }])

                            st.session_state.ice_sesiones = pd.concat([st.session_state.ice_sesiones, new_session], ignore_index=True)
                            
                            # Keep guests empty initially with a single empty row
                            new_guests_rows = [{
                                'ID_Sesion': new_id,
                                'Nombre': '',
                                'Empresa': '',
                                'Cargo': '',
                                'Asistencia': 'Asistió'
                            }]
                            st.session_state.ice_invitados = pd.concat([st.session_state.ice_invitados, pd.DataFrame(new_guests_rows)], ignore_index=True)

                            # Pre-populate an empty agreements row so the user has an initial row to write into
                            new_acuerdo_row = pd.DataFrame([{
                                'ID_Sesion': new_id,
                                'Descripcion': '',
                                'Responsable': '',
                                'Acuerdo': ''
                            }])
                            st.session_state.ice_acuerdos = pd.concat([st.session_state.ice_acuerdos, new_acuerdo_row], ignore_index=True)
                            
                            st.success(f"Sesión '{n_nombre}' creada exitosamente.")
                            st.session_state[f"sel_sesion_{v_proy}"] = new_id
                            fragment_rerun()
        
        # Select session
        if proj_sessions.empty:
            st.info("No hay sesiones ICE registradas para este proyecto.")
            sel_sesion_id = None
        else:
            session_options = proj_sessions.set_index('ID_Sesion')['Nombre_Sesion'].to_dict()
            default_sel_key = f"sel_sesion_{v_proy}"
            selected_idx = 0
            if default_sel_key in st.session_state and st.session_state[default_sel_key] in session_options:
                selected_idx = list(session_options.keys()).index(st.session_state[default_sel_key])
            
            sel_sesion_id = st.selectbox(
                "Seleccionar Sesión ICE",
                options=list(session_options.keys()),
                format_func=lambda x: session_options[x],
                index=selected_idx,
                key=f"sb_sesion_{v_proy}"
            )
            st.session_state[default_sel_key] = sel_sesion_id
            
    with c_act2:
        if sel_sesion_id:
            r_ses = proj_sessions[proj_sessions['ID_Sesion'] == sel_sesion_id].iloc[0]
            
            # Excel-like layout summary card
            logo_html_card = f'<img src="{logo_src}" class="ice-card-logo" style="max-height: 40px; object-fit: contain;" />' if logo_src else '<span style="font-weight: 800; color: #002060; font-size: 24px; letter-spacing: 3px; font-family: Outfit, sans-serif;">VENTI</span>'
            html_card = f"""
            <div style="font-family: 'Inter', sans-serif; border: 2px solid #1e293b; width: 100%; margin-bottom: 20px; background-color: #0d1424; color: #ffffff; border-radius: 8px; overflow: hidden; box-shadow: 0 4px 6px rgba(0, 0, 0, 0.3);">
              <!-- Title -->
              <div style="background-color: #1e293b; text-align: center; font-weight: bold; font-size: 16px; padding: 10px; border-bottom: 2px solid #1e293b; color: #00E5FF; letter-spacing: 1px;">
                {r_ses['Nombre_Sesion'].upper()}
              </div>
              
              <!-- Datos del proyecto -->
              <div style="background-color: #111b30; font-weight: bold; font-size: 11px; padding: 6px 12px; border-bottom: 1px solid #1e293b; color: #94a3b8; text-transform: uppercase; letter-spacing: 0.5px;">
                Datos del proyecto
              </div>
              <div style="display: flex; border-bottom: 1px solid #1e293b; flex-wrap: wrap;">
                <div style="flex: 2; min-width: 250px; padding: 12px; border-right: 1px solid #1e293b;">
                  <div style="display: flex; margin-bottom: 6px; font-size: 13px;">
                    <span style="font-weight: 600; width: 150px; color: #cbd5e1;">Proyecto:</span>
                    <span style="color: #ffffff;">{r_ses['Proyecto']}</span>
                  </div>
                  <div style="display: flex; font-size: 13px;">
                    <span style="font-weight: 600; width: 150px; color: #cbd5e1;">Etapa del proyecto:</span>
                    <span style="color: #ffffff;">{r_ses['Etapa']}</span>
                  </div>
                </div>
                <div style="flex: 1; min-width: 120px; display: flex; align-items: center; justify-content: center; padding: 12px;">
                  {logo_html_card}
                </div>

              </div>

              
              <!-- Datos de la Sesión -->
              <div style="background-color: #111b30; font-weight: bold; font-size: 11px; padding: 6px 12px; border-bottom: 1px solid #1e293b; color: #94a3b8; text-transform: uppercase; letter-spacing: 0.5px;">
                Datos de la Sesión
              </div>
              <div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); border-bottom: 1px solid #1e293b;">
                <div style="padding: 12px; border-right: 1px solid #1e293b; border-bottom: 1px solid #1e293b; font-size: 13px;">
                  <div style="font-weight: 600; color: #cbd5e1; margin-bottom: 4px;">Fecha de programación</div>
                  <div style="color: #ffffff;">{r_ses['Fecha_Prog']}</div>
                </div>
                <div style="padding: 12px; border-right: 1px solid #1e293b; border-bottom: 1px solid #1e293b; font-size: 13px;">
                  <div style="font-weight: 600; color: #cbd5e1; margin-bottom: 4px;">Fecha de sesión</div>
                  <div style="color: #ffffff;">{r_ses['Fecha_Real']}</div>
                </div>
                <div style="padding: 12px; border-right: 1px solid #1e293b; border-bottom: 1px solid #1e293b; font-size: 13px;">
                  <div style="font-weight: 600; color: #cbd5e1; margin-bottom: 4px;">Hora</div>
                  <div style="color: #ffffff;">{r_ses['Hora']}</div>
                </div>
                <div style="padding: 12px; border-bottom: 1px solid #1e293b; font-size: 13px;">
                  <div style="font-weight: 600; color: #cbd5e1; margin-bottom: 4px;">Organizador</div>
                  <div style="color: #ffffff;">{r_ses['Organizador']}</div>
                </div>
              </div>
              
              <!-- Encargado -->
              <div style="background-color: #111b30; font-weight: bold; font-size: 11px; padding: 6px 12px; border-bottom: 1px solid #1e293b; color: #94a3b8; text-transform: uppercase; letter-spacing: 0.5px;">
                Encargado
              </div>
              <div style="display: grid; grid-template-columns: 1fr;">
                <div style="padding: 12px; font-size: 13px;">
                  <div style="font-weight: 600; color: #cbd5e1; margin-bottom: 4px;">Líder</div>
                  <div style="color: #ffffff;">{r_ses['Lider']}</div>
                </div>
              </div>

            </div>
            """
            st.markdown(html_card, unsafe_allow_html=True)
            
            with st.expander("⚙️ EDITAR DATOS GENERALES DE LA SESIÓN", expanded=False, key=f"exp_edit_sesion_{v_proy}_{sel_sesion_id}"):
                with st.form(f"form_edit_sesion_{v_proy}_{sel_sesion_id}"):
                    c_e1, c_e2 = st.columns(2)
                    e_nombre = c_e1.text_input("Nombre de la Sesión", value=r_ses['Nombre_Sesion'], key=f"ice_e_nombre_{v_proy}_{sel_sesion_id}")
                    e_etapa = c_e2.selectbox("Etapa del Proyecto", ["Diseño", "Construcción", "Planificación", "Cierre"], index=["Diseño", "Construcción", "Planificación", "Cierre"].index(r_ses['Etapa']) if r_ses['Etapa'] in ["Diseño", "Construcción", "Planificación", "Cierre"] else 0, key=f"ice_e_etapa_{v_proy}_{sel_sesion_id}")
                    
                    c_e3, c_e4, c_e5 = st.columns(3)
                    try:
                        e_fecha_prog_val = datetime.strptime(r_ses['Fecha_Prog'], "%d/%m/%Y")
                    except:
                        e_fecha_prog_val = datetime.today()
                    try:
                        e_fecha_real_val = datetime.strptime(r_ses['Fecha_Real'], "%d/%m/%Y")
                    except:
                        e_fecha_real_val = datetime.today()
                        
                    e_fecha_prog = c_e3.date_input("Fecha de Programación", value=e_fecha_prog_val, key=f"ice_e_fecha_prog_{v_proy}_{sel_sesion_id}")
                    e_fecha_real = c_e4.date_input("Fecha de Sesión", value=e_fecha_real_val, key=f"ice_e_fecha_real_{v_proy}_{sel_sesion_id}")
                    e_hora = c_e5.text_input("Hora", value=r_ses['Hora'], key=f"ice_e_hora_{v_proy}_{sel_sesion_id}")
                    
                    c_e6, c_e7 = st.columns(2)
                    e_organizador = c_e6.text_input("Organizador", value=r_ses['Organizador'], key=f"ice_e_organizador_{v_proy}_{sel_sesion_id}")
                    e_lider = c_e7.text_input("Líder", value=r_ses['Lider'], key=f"ice_e_lider_{v_proy}_{sel_sesion_id}")
                    
                    if st.form_submit_button("ACTUALIZAR DATOS GENERALES"):
                        idx_s = st.session_state.ice_sesiones[st.session_state.ice_sesiones['ID_Sesion'] == sel_sesion_id].index[0]
                        st.session_state.ice_sesiones.at[idx_s, 'Nombre_Sesion'] = e_nombre
                        st.session_state.ice_sesiones.at[idx_s, 'Etapa'] = e_etapa
                        st.session_state.ice_sesiones.at[idx_s, 'Fecha_Prog'] = e_fecha_prog.strftime("%d/%m/%Y")
                        st.session_state.ice_sesiones.at[idx_s, 'Fecha_Real'] = e_fecha_real.strftime("%d/%m/%Y")
                        st.session_state.ice_sesiones.at[idx_s, 'Hora'] = e_hora
                        st.session_state.ice_sesiones.at[idx_s, 'Organizador'] = e_organizador
                        st.session_state.ice_sesiones.at[idx_s, 'Lider'] = e_lider
                        st.session_state.ice_sesiones.at[idx_s, 'Coordinador_VDC'] = ""
                        st.session_state.ice_sesiones.at[idx_s, 'Supervisor_VDC'] = ""

                        st.success("Datos generales actualizados localmente.")
                        fragment_rerun()

    # Render lists and editors only if a session is selected
    if sel_sesion_id:
        st.divider()
        
        # 1. Guest & Attendance Table
        st.markdown("#### 👥 INVITADOS Y ASISTENCIA")
        df_inv = st.session_state.ice_invitados[st.session_state.ice_invitados['ID_Sesion'] == sel_sesion_id].copy()
        df_inv_editor = df_inv[['Nombre', 'Empresa', 'Cargo', 'Asistencia']].copy()
        df_inv_editor['Nombre'] = df_inv_editor['Nombre'].fillna("").astype(str)
        df_inv_editor['Empresa'] = df_inv_editor['Empresa'].fillna("").astype(str)
        df_inv_editor['Cargo'] = df_inv_editor['Cargo'].fillna("").astype(str)
        df_inv_editor['Asistencia'] = df_inv_editor['Asistencia'].fillna("Asistió").astype(str)
        
        base_inv_key = f"base_df_ice_inv_{sel_sesion_id}"
        if base_inv_key not in st.session_state:
            st.session_state[base_inv_key] = df_inv_editor.copy()
            
        updated_inv = st.data_editor(
            st.session_state[base_inv_key],
            use_container_width=True,
            hide_index=True,
            num_rows="dynamic",
            key=f"ed_ice_inv_{v_proy}_{sel_sesion_id}",
            column_config={
                "Nombre": st.column_config.TextColumn("Nombre", help="Nombre del invitado", width="medium"),
                "Empresa": st.column_config.TextColumn("Empresa", help="Empresa del invitado"),
                "Cargo": st.column_config.TextColumn("Cargo", help="Cargo o especialidad"),
                "Asistencia": st.column_config.SelectboxColumn(
                    "Asistencia",
                    help="Estado de asistencia",
                    options=["Asistió", "Faltó", "Tarde", "Licencia"],
                    required=True
                )
            },
            on_change=on_table_edit,
            args=(
                base_inv_key,
                f"ed_ice_inv_{v_proy}_{sel_sesion_id}",
                'ice_invitados',
                ['ID_Sesion'],
                [sel_sesion_id],
                None,  # date_cols
                ['Nombre', 'Empresa', 'Cargo', 'Asistencia'],  # text_cols
                None,  # bool_cols
                None,  # num_cols
                None   # id_col
            )
        )
        
        st.divider()
        
        # 2. Activities & Agreements Table
        st.markdown("#### 📋 ACTIVIDADES Y ACUERDOS")
        df_acu = st.session_state.ice_acuerdos[st.session_state.ice_acuerdos['ID_Sesion'] == sel_sesion_id].copy()
        df_acu_editor = df_acu[['Descripcion', 'Responsable', 'Acuerdo']].copy()
        df_acu_editor['Descripcion'] = df_acu_editor['Descripcion'].fillna("").astype(str)
        df_acu_editor['Responsable'] = df_acu_editor['Responsable'].fillna("").astype(str)
        df_acu_editor['Acuerdo'] = df_acu_editor['Acuerdo'].fillna("").astype(str)
        
        base_acu_key = f"base_df_ice_acu_{sel_sesion_id}"
        if base_acu_key not in st.session_state:
            st.session_state[base_acu_key] = df_acu_editor.copy()
            
        updated_acu = st.data_editor(
            st.session_state[base_acu_key],
            use_container_width=True,
            hide_index=True,
            num_rows="dynamic",
            key=f"ed_ice_acu_{v_proy}_{sel_sesion_id}",
            column_config={
                "Descripcion": st.column_config.TextColumn("Descripción / Actividad", help="Descripción de la actividad o tema tratado", width="large"),
                "Responsable": st.column_config.TextColumn("Responsable", help="Persona o especialidad responsable"),
                "Acuerdo": st.column_config.TextColumn("Acuerdo", help="Acuerdo alcanzado", width="large")
            },
            on_change=on_table_edit,
            args=(
                base_acu_key,
                f"ed_ice_acu_{v_proy}_{sel_sesion_id}",
                'ice_acuerdos',
                ['ID_Sesion'],
                [sel_sesion_id],
                None,  # date_cols
                ['Descripcion', 'Responsable', 'Acuerdo'],  # text_cols
                None,  # bool_cols
                None,  # num_cols
                None   # id_col
            )
        )
        
        # Action buttons
        st.divider()
        c_btns1, c_btns2, c_btns3, c_btns4 = st.columns([2.5, 2.5, 2.5, 2.5])
        
        with c_btns1:
            st.info("💡 Cambios en memoria local. Sincroniza al terminar.")
            
        with c_btns2:
            with st.popover("✨ IMPORTAR NOTAS MEET", use_container_width=True):
                st.markdown("<p style='font-size: 12px; color: #cbd5e1;'>Incorpora asistentes y acuerdos a esta sesión desde un resumen de Gemini en Meet:</p>", unsafe_allow_html=True)
                f_append = st.file_uploader("Archivo (.txt, .docx, .pdf, .md)", type=['txt', 'docx', 'pdf', 'md'], key=f"f_app_{sel_sesion_id}")
                t_append = st.text_area("O texto del resumen:", height=100, key=f"t_app_{sel_sesion_id}")
                if st.button("PROCESAR E INCORPORAR", key=f"btn_app_{sel_sesion_id}", type="primary", use_container_width=True):
                    txt_proc = extraer_texto_documento(f_append) if f_append else (t_append or "")
                    if txt_proc.strip():
                        with st.spinner("Analizando con Gemini..."):
                            d_app = procesar_resumen_meet_gemini(txt_proc, v_proy, len(proj_sessions))
                        if d_app:
                            # Append guests
                            new_g = []
                            for g in d_app.get('invitados', []):
                                if g.get('nombre', '').strip():
                                    new_g.append({'ID_Sesion': sel_sesion_id, 'Nombre': g.get('nombre', ''), 'Empresa': g.get('empresa', ''), 'Cargo': g.get('cargo', ''), 'Asistencia': g.get('asistencia', 'Asistió')})
                            if new_g:
                                st.session_state.ice_invitados = pd.concat([st.session_state.ice_invitados, pd.DataFrame(new_g)], ignore_index=True)
                            # Append agreements
                            new_a = []
                            for a in d_app.get('acuerdos', []):
                                if a.get('descripcion', '').strip() or a.get('acuerdo', '').strip():
                                    new_a.append({'ID_Sesion': sel_sesion_id, 'Descripcion': a.get('descripcion', ''), 'Responsable': a.get('responsable', ''), 'Acuerdo': a.get('acuerdo', '')})
                            if new_a:
                                st.session_state.ice_acuerdos = pd.concat([st.session_state.ice_acuerdos, pd.DataFrame(new_a)], ignore_index=True)
                            st.success("Datos incorporados a la sesión actual.")
                            fragment_rerun()
            
        with c_btns3:
            try:
                pdf_bytes = generar_pdf_sesion(r_ses, updated_inv, updated_acu)
                st.download_button(
                    label="📥 EXPORTAR A PDF",
                    data=pdf_bytes,
                    file_name=f"{r_ses['Nombre_Sesion'].replace(' ', '_')}.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                    key=f"btn_pdf_ice_{v_proy}_{sel_sesion_id}"
                )
            except Exception as e_pdf:
                st.error(f"Error al generar PDF: {str(e_pdf)}")
                
        with c_btns4:
            if st.button("🗑️ ELIMINAR SESIÓN", type="secondary", use_container_width=True, key=f"btn_del_ice_{v_proy}_{sel_sesion_id}"):
                # Remove session and its children
                st.session_state.ice_sesiones = st.session_state.ice_sesiones[st.session_state.ice_sesiones['ID_Sesion'] != sel_sesion_id]
                st.session_state.ice_invitados = st.session_state.ice_invitados[st.session_state.ice_invitados['ID_Sesion'] != sel_sesion_id]
                st.session_state.ice_acuerdos = st.session_state.ice_acuerdos[st.session_state.ice_acuerdos['ID_Sesion'] != sel_sesion_id]
                
                st.warning("Sesión eliminada de la memoria local. Recuerda sincronizar para guardar en Google Sheets.")
                fragment_rerun()



@fragment_decorator
def render_seccion_ejecucion_obra(proy_name, proy_data):
    df_gg_proy = st.session_state.gg_log[st.session_state.gg_log['Proyecto'] == proy_name].copy()
    gg_gastado_total = pd.to_numeric(df_gg_proy['Monto'], errors='coerce').sum() if not df_gg_proy.empty else 0.0
    gg_restante = safe_float(proy_data.get('GG_Presupuesto', 0)) - gg_gastado_total
    
    df_pagos_proy = st.session_state.pagos[st.session_state.pagos['Proyecto'] == proy_name].copy()
    total_pagado_maestros = pd.to_numeric(df_pagos_proy['Monto'], errors='coerce').sum() if not df_pagos_proy.empty else 0.0
    presupuesto_restante = safe_float(proy_data.get('Presupuesto', 0)) - total_pagado_maestros

    util_meta = safe_float(proy_data.get('Utilidades', 0))
    df_ing_proy = st.session_state.ingresos_socios[(st.session_state.ingresos_socios['Proyecto'] == proy_name) & (st.session_state.ingresos_socios['Categoria'] == 'Utilidades')]
    utilidad_real = pd.to_numeric(df_ing_proy['Monto'], errors='coerce').sum() if not df_ing_proy.empty else 0.0

    tareas_proy = st.session_state.tasks[st.session_state.tasks['Proyecto'] == proy_name]
    t_totales = len(tareas_proy)
    t_a_tiempo = len(tareas_proy[tareas_proy['Estado'].astype(str).str.contains('Completado', case=False, na=False)])
    ppc = (t_a_tiempo / t_totales * 100) if t_totales > 0 else 0.0
    
    c1, c2, c3, c4, c5, c6 = st.columns(6)
    c1.markdown(f"<div class='metric-card'><h4>Presupuesto Total</h4><h2>S/ {safe_float(proy_data.get('Presupuesto', 0)):,.2f}</h2><small style='color:#3b82f6;'>Pagado: S/ {total_pagado_maestros:,.2f}</small></div>", unsafe_allow_html=True)
    c2.markdown(f"<div class='metric-card'><h4>Ppto. Restante</h4><h2 style='color:#00E5FF;'>S/ {presupuesto_restante:,.2f}</h2><small>Saldo a favor</small></div>", unsafe_allow_html=True)
    
    col_ut = "#10b981" if utilidad_real >= util_meta else "#00E5FF" 
    pct_util = (utilidad_real / util_meta * 100) if util_meta > 0 else 0.0
    c3.markdown(f"<div class='metric-card'><h4>Utilidades</h4><h2 style='color: {col_ut};'>S/ {utilidad_real:,.2f}</h2><small style='color:#94a3b8;'>Meta: S/ {util_meta:,.2f} ({pct_util:.1f}%)</small></div>", unsafe_allow_html=True)
    
    c4.markdown(f"<div class='metric-card'><h4>Presupuesto Caja Chica</h4><h2>S/ {safe_float(proy_data.get('GG_Presupuesto', 0)):,.2f}</h2></div>", unsafe_allow_html=True)
    c5.markdown(f"<div class='metric-card'><h4>Caja Chica Restante</h4><h2 style='color: {'#ef4444' if gg_restante < 0 else '#10b981'};'>S/ {gg_restante:,.2f}</h2><small>Gastado: S/ {gg_gastado_total:,.2f}</small></div>", unsafe_allow_html=True)
    c6.markdown(f"<div class='metric-card'><h4>PPC (Plan Cumplido)</h4><h2 style='color: #00E5FF;'>{ppc:.1f}%</h2><small>Tareas completadas</small></div>", unsafe_allow_html=True)
    
    st.divider()
    st.markdown("#### REGISTRO DE CAJA CHICA (SALIDAS)")
    with st.form(f"form_nuevo_gg_{proy_name}", clear_on_submit=True):
        col_g1, col_g2 = st.columns([6, 2])
        n_conc = col_g1.text_input("Concepto de Gasto")
        n_monto = col_g2.number_input("Monto (S/)", value=0.0, format="%.2f")
        if st.form_submit_button("REGISTRAR SALIDA CAJA CHICA") and n_conc and n_monto != 0.0:
            nueva_fila = pd.DataFrame([{'Proyecto': proy_name, 'Fecha': fecha_hoy.strftime("%d/%m/%Y"), 'Responsable': '', 'Concepto': n_conc, 'Monto': n_monto}])
            st.session_state.gg_log = pd.concat([st.session_state.gg_log, nueva_fila], ignore_index=True)
            if f"base_df_gg_{proy_name}" in st.session_state:
                del st.session_state[f"base_df_gg_{proy_name}"]
            fragment_rerun()
            
    if not df_gg_proy.empty:
        df_gg_proy_display = df_gg_proy.drop(columns=['Proyecto', 'Responsable']) if 'Responsable' in df_gg_proy.columns else df_gg_proy.drop(columns=['Proyecto'])
        for col in ['Fecha', 'Concepto', 'Monto']:
            if col not in df_gg_proy_display.columns:
                df_gg_proy_display[col] = ""
        df_gg_proy_display = df_gg_proy_display[['Fecha', 'Concepto', 'Monto']]
        
        base_key = f"base_df_gg_{proy_name}"
        if base_key not in st.session_state:
            st.session_state[base_key] = df_gg_proy_display.copy()
            
        updated_gg = st.data_editor(
            st.session_state[base_key],
            use_container_width=True,
            hide_index=True,
            num_rows="dynamic",
            key=f"ed_gg_{proy_name}",
            on_change=on_table_edit,
            args=(
                base_key,
                f"ed_gg_{proy_name}",
                'gg_log',
                ['Proyecto'],
                [proy_name],
                None,  # date_cols
                ['Fecha', 'Concepto', 'Responsable'],  # text_cols
                None,  # bool_cols
                ['Monto'],  # num_cols
                None   # id_col
            )
        )

    # LOGÍSTICA DE MATERIALES EN DASHBOARD
    st.divider()
    st.markdown("#### LOGÍSTICA DE MATERIALES")
    
    # Limpiar y alinear la base en memoria con el nuevo esquema
    if 'materiales' in st.session_state and not st.session_state.materiales.empty:
        # Migrar valores de columnas viejas si aún quedan en la sesión activa
        if 'Ppto_Estimado' in st.session_state.materiales.columns:
            if 'Costo_Estimado' not in st.session_state.materiales.columns:
                st.session_state.materiales['Costo_Estimado'] = st.session_state.materiales['Ppto_Estimado']
            else:
                st.session_state.materiales['Costo_Estimado'] = st.session_state.materiales['Costo_Estimado'].combine_first(st.session_state.materiales['Ppto_Estimado'])
        if 'Ppto_Real' in st.session_state.materiales.columns:
            if 'Costo_Real' not in st.session_state.materiales.columns:
                st.session_state.materiales['Costo_Real'] = st.session_state.materiales['Ppto_Real']
            else:
                st.session_state.materiales['Costo_Real'] = st.session_state.materiales['Costo_Real'].combine_first(st.session_state.materiales['Ppto_Real'])
        if 'Estado' in st.session_state.materiales.columns:
            st.session_state.materiales['Estado'] = st.session_state.materiales['Estado'].replace({
                'No Comprado': '🔴 No Comprado',
                'Comprado': '🟢 Comprado'
            }).fillna('🔴 No Comprado')
            
        # Asegurar que existan todas las columnas
        for col in cols_mat:
            if col not in st.session_state.materiales.columns:
                if col == 'Estado':
                    st.session_state.materiales[col] = '🔴 No Comprado'
                else:
                    st.session_state.materiales[col] = ''
        # Mantener solo columnas correctas
        st.session_state.materiales = st.session_state.materiales[cols_mat]

    df_mat = st.session_state.materiales[st.session_state.materiales['Proyecto'] == proy_name].copy()
    df_mat_edit = df_mat.drop(columns=['Proyecto', 'ID']) if not df_mat.empty else pd.DataFrame(columns=['Material', 'Cantidad', 'Unidad', 'Costo_Estimado', 'Costo_Real', 'Proveedor', 'Estado'])
    
    # Asegurar que todas las columnas necesarias existan con sus valores por defecto
    for c in ['Material', 'Cantidad', 'Unidad', 'Costo_Estimado', 'Costo_Real', 'Proveedor', 'Estado']:
        if c not in df_mat_edit.columns: 
            if c == 'Estado':
                df_mat_edit[c] = "🔴 No Comprado"
            else:
                df_mat_edit[c] = ""
            
    base_key = f"base_df_mat_{proy_name}"
    if base_key not in st.session_state:
        st.session_state[base_key] = df_mat_edit.copy()

    upd_mat = st.data_editor(
        st.session_state[base_key], use_container_width=True, hide_index=True, num_rows="dynamic", key=f"m_dash_{proy_name}",
        column_config={
            "Cantidad": st.column_config.NumberColumn("Cant."), 
            "Unidad": st.column_config.SelectboxColumn("Und", options=['und', 'kg', 'm2', 'm3', 'glb', 'm', 'bolsa']),
            "Costo_Estimado": st.column_config.NumberColumn("Costo Estimado", format="S/ %.2f"),
            "Costo_Real": st.column_config.NumberColumn("Costo Real", format="S/ %.2f"),
            "Estado": st.column_config.SelectboxColumn("Estado", options=['🔴 No Comprado', '🟢 Comprado'], default='🔴 No Comprado')
        },
        on_change=on_table_edit,
        args=(
            base_key,
            f"m_dash_{proy_name}",
            'materiales',
            ['Proyecto'],
            [proy_name],
            None,  # date_cols
            ['Material', 'Unidad', 'Proveedor', 'Estado'],  # text_cols
            None,  # bool_cols
            ['Cantidad', 'Costo_Estimado', 'Costo_Real'],  # num_cols
            'ID'   # id_col
        )
    )

    # REGISTRO DE INGRESOS (UTILIDADES DEL PROYECTO)
    st.divider()
    st.markdown("#### REGISTRO DE INGRESOS (UTILIDADES DEL PROYECTO)")
    with st.form(f"form_nuevo_ing_{proy_name}", clear_on_submit=True):
        col_i1, col_i2 = st.columns([6, 2])
        n_notas = col_i1.text_input("Nota / Detalle de Utilidad")
        n_monto_ing = col_i2.number_input("Monto de Utilidad (S/)", value=0.0, format="%.2f")
        if st.form_submit_button("REGISTRAR UTILIDAD") and n_monto_ing != 0.0:
            nueva_fila_ing = pd.DataFrame([{
                'ID': datetime.now().timestamp(),
                'Fecha': fecha_hoy.strftime("%d/%m/%Y"),
                'Proyecto': proy_name,
                'Socio': 'Venti',
                'Categoria': 'Utilidades',
                'Monto': n_monto_ing,
                'Notas': n_notas
            }])
            st.session_state.ingresos_socios = pd.concat([st.session_state.ingresos_socios, nueva_fila_ing], ignore_index=True)
            if f"base_df_ing_proy_{proy_name}" in st.session_state:
                del st.session_state[f"base_df_ing_proy_{proy_name}"]
            fragment_rerun()
            
    if not df_ing_proy.empty:
        df_ing_proy_display = df_ing_proy[['Fecha', 'Monto', 'Notas']].copy()
        df_ing_proy_display['Monto'] = pd.to_numeric(df_ing_proy_display['Monto'], errors='coerce').fillna(0.0).astype(float)
        df_ing_proy_display['Fecha'] = df_ing_proy_display['Fecha'].fillna("").astype(str)
        df_ing_proy_display['Notas'] = df_ing_proy_display['Notas'].fillna("").astype(str)
        
        base_key = f"base_df_ing_proy_{proy_name}"
        if base_key not in st.session_state:
            st.session_state[base_key] = df_ing_proy_display.copy()
            
        updated_ing = st.data_editor(
            st.session_state[base_key],
            use_container_width=True,
            hide_index=True,
            num_rows="dynamic",
            key=f"ed_ing_proy_{proy_name}",
            column_config={
                "Fecha": st.column_config.TextColumn("Fecha"),
                "Monto": st.column_config.NumberColumn("Monto (S/)", format="%.2f"),
                "Notas": st.column_config.TextColumn("Nota / Detalle")
            },
            on_change=on_table_edit,
            args=(
                base_key,
                f"ed_ing_proy_{proy_name}",
                'ingresos_socios',
                ['Proyecto', 'Categoria', 'Socio'],
                [proy_name, 'Utilidades', 'Venti'],
                None,  # date_cols
                ['Fecha', 'Notas'],  # text_cols
                None,  # bool_cols
                ['Monto'],  # num_cols
                'ID'   # id_col
            )
        )

    st.divider()
    st.markdown("#### HISTORIAL DE PAGOS A CONTRATISTAS (POR PARTIDA)")
    partidas_proy = st.session_state.partidas[st.session_state.partidas['Proyecto'] == proy_name]
    if not partidas_proy.empty:
        partidas_list = partidas_proy['Nombre_Partida'].unique().tolist()
        partidas_options = ["Todas las partidas"] + partidas_list
        
        # Partida selector dropdown
        p_seleccionada = st.selectbox(
            "Seleccionar Partida",
            options=partidas_options,
            key=f"sb_partida_pagos_{proy_name}"
        )
        
        if p_seleccionada:
            if p_seleccionada == "Todas las partidas":
                df_pact_proy = st.session_state.p_info[st.session_state.p_info['Proyecto'] == proy_name].dropna(subset=['Nombre_Partida'])
                df_pact_proy = df_pact_proy.drop_duplicates(subset=['Nombre_Partida'], keep='first')
                monto_pactado_dash = pd.to_numeric(df_pact_proy['Monto_Pactado'], errors='coerce').sum()
                
                df_p = df_pagos_proy if not df_pagos_proy.empty else pd.DataFrame(columns=cols_pagos)
                monto_pagado_dash = pd.to_numeric(df_p['Monto'], errors='coerce').sum()
                saldo_dash = monto_pactado_dash - monto_pagado_dash
                color_saldo_dash = "#10b981" if saldo_dash >= 0 else "#ef4444"
            else:
                df_pact_dash = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy_name) & (st.session_state.p_info['Nombre_Partida'] == p_seleccionada)]
                monto_pactado_dash = float(df_pact_dash['Monto_Pactado'].iloc[0]) if not df_pact_dash.empty else 0.0
                
                df_p = df_pagos_proy[df_pagos_proy['Partida'] == p_seleccionada] if not df_pagos_proy.empty else pd.DataFrame(columns=cols_pagos)
                monto_pagado_dash = pd.to_numeric(df_p['Monto'], errors='coerce').sum()
                saldo_dash = monto_pactado_dash - monto_pagado_dash
                color_saldo_dash = "#10b981" if saldo_dash >= 0 else "#ef4444"
            
            # Layout with two columns: Left for cumulative curve, Right for circular budget indicator
            col_chart, col_circle = st.columns([5, 5])
            
            with col_chart:
                df_p_chart = df_p.copy()
                df_p_chart['Fecha_DT'] = pd.to_datetime(df_p_chart['Fecha'], format='%d/%m/%Y', errors='coerce')
                df_p_chart = df_p_chart.dropna(subset=['Fecha_DT', 'Monto'])
                df_p_chart['Monto'] = pd.to_numeric(df_p_chart['Monto'], errors='coerce').fillna(0.0)
                
                if not df_p_chart.empty:
                    df_p_chart = df_p_chart.sort_values('Fecha_DT')
                    df_grouped = df_p_chart.groupby('Fecha_DT').agg({'Monto': 'sum', 'Fecha': 'first'}).reset_index()
                    df_grouped = df_grouped.sort_values('Fecha_DT')
                    df_grouped['Acumulado'] = df_grouped['Monto'].cumsum()
                    
                    # Prepend a S/ 0 point to make the curve start nicely
                    if p_seleccionada == "Todas las partidas":
                        df_pact_proy = st.session_state.p_info[st.session_state.p_info['Proyecto'] == proy_name].dropna(subset=['Nombre_Partida'])
                        df_pact_proy = df_pact_proy.drop_duplicates(subset=['Nombre_Partida'], keep='first')
                        df_pact_proy['Inicio_DT'] = pd.to_datetime(df_pact_proy['Inicio'], format='%d/%m/%Y', errors='coerce')
                        partida_start_dt = df_pact_proy['Inicio_DT'].min()
                        partida_start = partida_start_dt.strftime('%d/%m/%Y') if pd.notna(partida_start_dt) else None
                    else:
                        partida_start = df_pact_dash['Inicio'].iloc[0] if not df_pact_dash.empty else None
                        
                    if partida_start:
                        partida_start_dt = pd.to_datetime(partida_start, format='%d/%m/%Y', errors='coerce')
                        if pd.notna(partida_start_dt) and partida_start_dt < df_grouped['Fecha_DT'].min():
                            start_row = pd.DataFrame([{'Fecha_DT': partida_start_dt, 'Fecha': partida_start, 'Monto': 0.0, 'Acumulado': 0.0}])
                            df_grouped = pd.concat([start_row, df_grouped], ignore_index=True)
                        else:
                            start_date = df_grouped['Fecha_DT'].min() - pd.Timedelta(days=1)
                            start_row = pd.DataFrame([{'Fecha_DT': start_date, 'Fecha': start_date.strftime('%d/%m/%Y'), 'Monto': 0.0, 'Acumulado': 0.0}])
                            df_grouped = pd.concat([start_row, df_grouped], ignore_index=True)
                    else:
                        start_date = df_grouped['Fecha_DT'].min() - pd.Timedelta(days=1)
                        start_row = pd.DataFrame([{'Fecha_DT': start_date, 'Fecha': start_date.strftime('%d/%m/%Y'), 'Monto': 0.0, 'Acumulado': 0.0}])
                        df_grouped = pd.concat([start_row, df_grouped], ignore_index=True)
                    
                    # Convert to datetime and sort
                    df_grouped['Fecha_DT'] = pd.to_datetime(df_grouped['Fecha_DT'])
                    df_grouped = df_grouped.sort_values('Fecha_DT')
                    
                    # Use Streamlit's native line chart with DatetimeIndex
                    chart_data = df_grouped.set_index('Fecha_DT')[['Acumulado']]
                    st.line_chart(chart_data, color="#00E5FF", height=210)
                else:
                    st.info("No hay pagos registrados para graficar en esta partida." if p_seleccionada != "Todas las partidas" else "No hay pagos registrados para graficar en el proyecto.")
                    
            with col_circle:
                pct = (monto_pagado_dash / monto_pactado_dash) if monto_pactado_dash > 0 else 0.0
                pct_capped = min(pct, 1.0)
                pct_display = pct * 100
                
                # SVG Circular Progress Ring (Enlarged Row layout, height matched to chart)
                svg_donut_html = f"""
                <div style="
                    background-color: var(--card-bg);
                    padding: 15px;
                    border-radius: 6px;
                    border: 1px solid var(--border-color);
                    box-shadow: var(--card-shadow);
                    display: flex;
                    flex-direction: row;
                    align-items: center;
                    justify-content: space-between;
                    height: 210px;
                    box-sizing: border-box;
                ">
                  <!-- Left Side: Large Circular indicator -->
                  <div style="position: relative; width: 180px; height: 180px; display: flex; align-items: center; justify-content: center; flex-shrink: 0;">
                    <svg width="180" height="180" viewBox="0 0 180 180">
                      <!-- Background Circle (Grey) -->
                      <circle cx="90" cy="90" r="74" fill="transparent" stroke="rgba(255, 255, 255, 0.1)" stroke-width="12" />
                      <!-- Foreground Circle (Cyan) -->
                      <circle cx="90" cy="90" r="74" fill="transparent" stroke="#00E5FF" stroke-width="12"
                              stroke-dasharray="464.96"
                              stroke-dashoffset="{464.96 * (1 - pct_capped)}"
                              stroke-linecap="round"
                              transform="rotate(-90 90 90)"
                              style="transition: stroke-dashoffset 0.8s ease-in-out; filter: drop-shadow(0 0 5px rgba(0, 229, 255, 0.45));" />
                    </svg>
                    <div style="position: absolute; display: flex; flex-direction: column; align-items: center; justify-content: center;">
                      <span style="font-size: 26px; font-weight: bold; color: var(--text-color); font-family: 'Inter', sans-serif;">{pct_display:.1f}%</span>
                      <span style="font-size: 10px; color: #00E5FF; text-transform: uppercase; font-weight: 500; letter-spacing: 0.5px; margin-top: 2px;">Pagado</span>
                    </div>
                  </div>
                  
                  <!-- Right Side: Info breakdown -->
                  <div style="
                      display: flex;
                      flex-direction: column;
                      justify-content: center;
                      flex-grow: 1;
                      margin-left: 20px;
                      font-family: monospace;
                      font-size: 13px;
                      color: var(--text-color);
                  ">
                    <div style="font-size: 11px; font-weight: 600; color: var(--text-muted); margin-bottom: 15px; text-transform: uppercase; letter-spacing: 0.5px; font-family: sans-serif; text-align: left;">Presupuesto vs Pago</div>
                    <div style="display: flex; justify-content: space-between; margin-bottom: 4px;">
                      <span style="color: var(--text-muted);">Pactado:</span>
                      <span style="color: var(--text-color); font-weight: bold;">S/{monto_pactado_dash:,.2f}</span>
                    </div>
                    <div style="display: flex; justify-content: space-between; margin-bottom: 4px;">
                      <span style="color: var(--text-muted);">Pagado:</span>
                      <span style="color: #00E5FF; font-weight: bold;">S/{monto_pagado_dash:,.2f}</span>
                    </div>
                    <div style="height: 1px; background-color: var(--border-color); margin: 6px 0;"></div>
                    <div style="display: flex; justify-content: space-between;">
                      <span style="color: var(--text-muted);">Saldo:</span>
                      <span style="color: {color_saldo_dash}; font-weight: bold;">S/{saldo_dash:,.2f}</span>
                    </div>
                  </div>
                </div>
                """
                st.markdown(svg_donut_html, unsafe_allow_html=True)
            
            st.write("") # small spacing
            
            # Payment editor for the selected partida
            if p_seleccionada == "Todas las partidas":
                df_p_clean = df_p.drop(columns=['Proyecto', 'ID'], errors='ignore') if not df_p.empty else pd.DataFrame(columns=['Partida', 'Fecha', 'Contratista', 'Monto', 'Tipo', 'Notas'])
                for col in ['Partida', 'Fecha', 'Contratista', 'Monto', 'Tipo', 'Notas']:
                    if col not in df_p_clean.columns: df_p_clean[col] = ""
                cols_order = ['Partida', 'Fecha', 'Contratista', 'Monto', 'Tipo', 'Notas']
                df_p_display = df_p_clean[[c for c in cols_order if c in df_p_clean.columns]].reset_index(drop=True)
            else:
                df_p_clean = df_p.drop(columns=['Proyecto', 'ID', 'Partida'], errors='ignore') if not df_p.empty else pd.DataFrame(columns=['Fecha', 'Contratista', 'Monto', 'Tipo', 'Notas'])
                for col in ['Fecha', 'Contratista', 'Monto', 'Tipo', 'Notas']:
                    if col not in df_p_clean.columns: df_p_clean[col] = ""
                cols_order = ['Fecha', 'Contratista', 'Monto', 'Tipo', 'Notas']
                df_p_display = df_p_clean[[c for c in cols_order if c in df_p_clean.columns]].reset_index(drop=True)
                
            df_p_display['Notas'] = df_p_display['Notas'].fillna("").astype(str).replace('nan', '').replace('None', '')
            df_p_display['Contratista'] = df_p_display['Contratista'].fillna("").astype(str)
            df_p_display['Tipo'] = df_p_display['Tipo'].fillna("").astype(str)
            df_p_display['Fecha'] = df_p_display['Fecha'].fillna("").astype(str)
            df_p_display['Monto'] = pd.to_numeric(df_p_display['Monto'], errors='coerce').fillna(0.0)
            if 'Partida' in df_p_display.columns:
                df_p_display['Partida'] = df_p_display['Partida'].fillna("").astype(str)

            key_editor = f"ed_pagos_dash_{proy_name}_{p_seleccionada}"

            if p_seleccionada == "Todas las partidas":
                upd_p = st.data_editor(
                    df_p_display,
                    use_container_width=True,
                    hide_index=True,
                    num_rows="dynamic",
                    key=key_editor,
                    column_config={
                        "Partida": st.column_config.SelectboxColumn("Partida", options=partidas_list),
                        "Contratista": st.column_config.SelectboxColumn("Contratista", options=st.session_state.maestro['ID_Contratista'].tolist()),
                        "Tipo": st.column_config.SelectboxColumn("Tipo", options=["Adelanto", "Valorización", "Liquidación", "Adicional", "Otros"]),
                        "Monto": st.column_config.NumberColumn("Monto (S/)", format="%.2f"),
                        "Notas": st.column_config.TextColumn("Notas")
                    }
                )
                if df_p_display.astype(str).to_dict() != upd_p.astype(str).to_dict():
                    st.session_state.pagos = st.session_state.pagos[st.session_state.pagos['Proyecto'] != proy_name]
                    if not upd_p.empty:
                        n_p = upd_p.copy()
                        n_p['Proyecto'] = proy_name
                        n_p['ID'] = [int(datetime.now().timestamp() * 1000) + i for i in range(len(n_p))]
                        for c in cols_pagos:
                            if c not in n_p.columns: n_p[c] = ""
                        st.session_state.pagos = pd.concat([st.session_state.pagos, n_p[cols_pagos]], ignore_index=True)
                    clear_base_dfs()
                    st.rerun()
            else:
                upd_p = st.data_editor(
                    df_p_display,
                    use_container_width=True,
                    hide_index=True,
                    num_rows="dynamic",
                    key=key_editor,
                    column_config={
                        "Contratista": st.column_config.SelectboxColumn("Contratista", options=st.session_state.maestro['ID_Contratista'].tolist()),
                        "Tipo": st.column_config.SelectboxColumn("Tipo", options=["Adelanto", "Valorización", "Liquidación", "Adicional", "Otros"]),
                        "Monto": st.column_config.NumberColumn("Monto (S/)", format="%.2f"),
                        "Notas": st.column_config.TextColumn("Notas")
                    }
                )
                if df_p_display.astype(str).to_dict() != upd_p.astype(str).to_dict():
                    mask_del = (st.session_state.pagos['Proyecto'] == proy_name) & (st.session_state.pagos['Partida'] == p_seleccionada)
                    st.session_state.pagos = st.session_state.pagos[~mask_del]
                    if not upd_p.empty:
                        n_p = upd_p.copy()
                        n_p['Proyecto'] = proy_name
                        n_p['Partida'] = p_seleccionada
                        n_p['ID'] = [int(datetime.now().timestamp() * 1000) + i for i in range(len(n_p))]
                        for c in cols_pagos:
                            if c not in n_p.columns: n_p[c] = ""
                        st.session_state.pagos = pd.concat([st.session_state.pagos, n_p[cols_pagos]], ignore_index=True)
                    clear_base_dfs()
                    st.rerun()
    else: st.info("No hay partidas (secciones) creadas en este proyecto para registrar pagos. Agrégalas en la Matriz de Operaciones.")


    is_multi_proy = proy_data.get('Tipo_Proyecto') == 'Multifamiliar'
    
    if is_multi_proy:
        st.divider()
        st.markdown("#### TREN DE TRABAJO (AVANCE POR ALTURA)")
        
        num_pisos = int(proy_data.get('Num_Pisos', 5))
        num_sotanos = int(proy_data.get('Num_Sotanos', 1))
        
        df_p_info_proy = st.session_state.p_info[st.session_state.p_info['Proyecto'] == proy_name].copy()
        if not df_p_info_proy.empty:
            df_p_info_proy = df_p_info_proy.dropna(subset=['Nombre_Partida'])
            df_p_info_proy['has_data'] = df_p_info_proy['Nivel_Inicio'].notna() & (df_p_info_proy['Nivel_Inicio'] != '')
            df_p_info_proy = df_p_info_proy.sort_values(by=['has_data', 'Estado'], ascending=[False, False])
            df_p_info_proy = df_p_info_proy.drop_duplicates(subset=['Nombre_Partida'], keep='first').drop(columns=['has_data'])
        
        if not df_p_info_proy.empty:
            # 1. Build sorted floors list
            floors = [f"Sótano {i}" for i in range(num_sotanos, 0, -1)] + [f"Piso {i}" for i in range(1, num_pisos + 1)]
            name_to_numeric = {name: idx for idx, name in enumerate(floors)}
            
            # Map tick values to the center of each floor block
            numeric_to_name = {idx + 0.5: name for idx, name in enumerate(floors)}
            
            df_chart = df_p_info_proy.copy()
            y_starts = []
            y_ends = []
            y_vals = []
            base_y_vals = []
            
            for _, row in df_chart.iterrows():
                nivel = row.get('Nivel_Actual', 'No iniciado')
                etapa = row.get('Etapa', 'Subestructura')
                n_ini = row.get('Nivel_Inicio', '')
                n_fin = row.get('Nivel_Fin', '')
                estado = row.get('Estado', 'No iniciado')
                
                if etapa == 'Subestructura':
                    default_ini = "Sótano 1" if num_sotanos > 0 else "Piso 1"
                    default_fin = f"Sótano {num_sotanos}" if num_sotanos > 0 else "Piso 1"
                else:
                    default_ini = "Piso 1"
                    default_fin = f"Piso {num_pisos}"
                
                if not n_ini or n_ini not in name_to_numeric:
                    n_ini = default_ini
                if not n_fin or n_fin not in name_to_numeric:
                    n_fin = default_fin
                
                idx_ini = name_to_numeric[n_ini]
                idx_fin = name_to_numeric[n_fin]
                
                # Determine background bounds
                if idx_ini <= idx_fin:
                    y_start = idx_ini
                    y_end = idx_fin + 1
                else:
                    y_start = idx_fin
                    y_end = idx_ini + 1
                
                # Determine progress bounds (y_val and Base_Y)
                if estado == 'No iniciado' or nivel == 'No iniciado' or pd.isna(nivel) or not nivel:
                    base_y = y_start
                    y_val = y_start
                else:
                    if idx_ini <= idx_fin: # Upwards
                        base_y = idx_ini
                        if estado == 'Completado':
                            y_val = idx_fin + 1
                        else:
                            idx_act = name_to_numeric.get(nivel, idx_ini)
                            y_val = idx_act + 1
                    else: # Downwards
                        base_y = idx_ini + 1
                        if estado == 'Completado':
                            y_val = idx_fin
                        else:
                            idx_act = name_to_numeric.get(nivel, idx_ini)
                            y_val = idx_act
                
                y_starts.append(y_start)
                y_ends.append(y_end)
                y_vals.append(y_val)
                base_y_vals.append(base_y)
                
            df_chart['y_start'] = y_starts
            df_chart['y_end'] = y_ends
            df_chart['y_val'] = y_vals
            df_chart['Base_Y'] = base_y_vals
            
            # Sort order for nominal X axis by start date (Inicio)
            df_chart['Inicio_DT'] = pd.to_datetime(df_chart['Inicio'], errors='coerce')
            sorted_partidas = df_chart.sort_values('Inicio_DT')['Nombre_Partida'].drop_duplicates().tolist()
            
            js_dict = ", ".join([f"'{val}': '{name}'" for val, name in numeric_to_name.items()])
            js_map = "{" + js_dict + "}"
            label_expr = f"({js_map}[datum.value] || '')"
            
            color_scale = alt.Scale(
                domain=['No iniciado', 'En Proceso', 'Completado'],
                range=['#475569', '#00E5FF', '#10b981']
            )
            
            chart_bg = alt.Chart(df_chart).mark_bar(
                color='#1e293b',
                opacity=0.3,
                cornerRadius=4,
                stroke='#334155',
                strokeWidth=1
            ).encode(
                y=alt.Y('Nombre_Partida:N', title='Partida', sort=sorted_partidas, axis=alt.Axis(labelColor='#cbd5e1')),
                x=alt.X('y_end:Q', title='Nivel de Altura', scale=alt.Scale(domain=[0, len(floors)], nice=False), axis=alt.Axis(
                    values=list(numeric_to_name.keys()),
                    labelExpr=label_expr,
                    labelOverlap=False,
                    gridColor='#1e293b'
                )),
                x2='y_start:Q',
                tooltip=['Nombre_Partida', 'Etapa', 'Nivel_Inicio', 'Nivel_Fin', 'Nivel_Actual', 'Contratista']
            )
            
            chart_fg = alt.Chart(df_chart).mark_bar(cornerRadius=4).encode(
                y=alt.Y('Nombre_Partida:N', title='Partida', sort=sorted_partidas),
                x=alt.X('y_val:Q'),
                x2='Base_Y:Q',
                color=alt.Color('Estado:N', scale=color_scale, legend=alt.Legend(title="Estado", labelColor='#cbd5e1', titleColor='#cbd5e1')),
                tooltip=['Nombre_Partida', 'Etapa', 'Nivel_Inicio', 'Nivel_Fin', 'Nivel_Actual', 'Contratista', 'Inicio', 'Fin', 'Estado']
            )
            
            ground_level_idx = next((idx for idx, name in enumerate(floors) if name.startswith("Piso")), 0)
            df_cero = pd.DataFrame({'x': [ground_level_idx]})
            linea_cero = alt.Chart(df_cero).mark_rule(color='#ef4444', strokeWidth=1.5, strokeDash=[4, 4]).encode(x='x:Q')
            
            chart_altura_final = alt.layer(chart_bg, chart_fg, linea_cero).properties(height=400)
            
            st.altair_chart(chart_altura_final, use_container_width=True)
            
            st.divider()
            st.markdown("#### CRONOGRAMA GENERAL DE PARTIDAS (GANTT)")
            df_gantt_part = df_p_info_proy.copy().dropna(subset=['Inicio', 'Fin'])
            df_gantt_part = df_gantt_part[(df_gantt_part['Inicio'] != '') & (df_gantt_part['Fin'] != '')]
            
            if not df_gantt_part.empty:
                df_gantt_part['Inicio_DT'] = pd.to_datetime(df_gantt_part['Inicio'], errors='coerce')
                df_gantt_part['Fin_Visual'] = pd.to_datetime(df_gantt_part['Fin'], errors='coerce')
                df_gantt_part = df_gantt_part.dropna(subset=['Inicio_DT', 'Fin_Visual'])
                mask_1day = df_gantt_part['Inicio_DT'].dt.date == df_gantt_part['Fin_Visual'].dt.date
                df_gantt_part.loc[mask_1day, 'Fin_Visual'] = df_gantt_part['Fin_Visual'] + timedelta(days=1)
                
                chart_gantt_part = alt.Chart(df_gantt_part).mark_bar(cornerRadius=2).encode(
                    x=alt.X('Inicio_DT:T', title='Timeline', axis=alt.Axis(gridColor='#1e293b', domainColor='#1e293b', format='%d %b')),
                    x2='Fin_Visual:T',
                    y=alt.Y('Nombre_Partida:N', sort='x', title='', axis=alt.Axis(labelColor='#cbd5e1')),
                    color=alt.Color('Nombre_Partida:N', scale=alt.Scale(scheme='dark2'), legend=None),
                    tooltip=['Nombre_Partida', 'Contratista', 'Inicio', 'Fin', 'Estado']
                )
                linea_hoy_part = alt.Chart(df_hoy).mark_rule(color='#00E5FF', strokeWidth=2).encode(x='Hoy:T')
                gantt_final_part = alt.layer(chart_gantt_part, linea_hoy_part).properties(height=250).configure_view(strokeOpacity=0).configure_axis(grid=True).interactive(bind_y=False)
                st.altair_chart(gantt_final_part, use_container_width=True)
            else:
                st.info("Sin fechas definidas para las partidas. Asigne fechas en la Matriz de Operaciones para ver el cronograma.")
        else:
            st.info("No hay partidas configuradas en este proyecto para graficar.")
            
    else:
        st.divider()
        st.markdown("#### DIAGRAMA DE SECUENCIA (GANTT)")
        df_gantt_p = st.session_state.tasks[st.session_state.tasks['Proyecto'] == proy_name].copy()
        df_gantt_p['Inicio'] = pd.to_datetime(df_gantt_p['Inicio'], errors='coerce')
        df_gantt_p['Fin_Visual'] = pd.to_datetime(df_gantt_p['Fin'], errors='coerce')
        df_gantt_p = df_gantt_p.dropna(subset=['Inicio', 'Fin_Visual'])
        if not df_gantt_p.empty:
            mask_1day = df_gantt_p['Inicio'].dt.date == df_gantt_p['Fin_Visual'].dt.date
            df_gantt_p.loc[mask_1day, 'Fin_Visual'] = df_gantt_p['Fin_Visual'] + timedelta(days=1)
            
            chart_p = alt.Chart(df_gantt_p).mark_bar(cornerRadius=2).encode(
                x=alt.X('Inicio:T', title='Timeline', axis=alt.Axis(gridColor='#1e293b', domainColor='#1e293b', format='%d %b')), 
                x2='Fin_Visual:T', 
                y=alt.Y('Tarea:N', sort='x', title='', axis=alt.Axis(labelColor='#cbd5e1')),
                color=alt.Color('Partida:N', scale=alt.Scale(scheme='dark2')), 
                tooltip=['Partida', 'Tarea', 'Contratista', 'Inicio', 'Fin']
            )
            linea_hoy = alt.Chart(df_hoy).mark_rule(color='#00E5FF', strokeWidth=2).encode(x='Hoy:T')
            gantt_final = alt.layer(chart_p, linea_hoy).properties(height=350).configure_view(strokeOpacity=0).configure_axis(grid=True).interactive(bind_y=False)
            st.altair_chart(gantt_final, use_container_width=True)


df_tareas_activas = st.session_state.tasks.copy()
df_tareas_activas['Inicio_DT'] = pd.to_datetime(df_tareas_activas['Inicio'], errors='coerce')
df_tareas_activas['Fin_DT'] = pd.to_datetime(df_tareas_activas['Fin'], errors='coerce')
df_tareas_activas = df_tareas_activas.dropna(subset=['Inicio_DT', 'Fin_DT'])
df_tareas_activas['Inicio_Date'] = df_tareas_activas['Inicio_DT'].dt.date
df_tareas_activas['Fin_Date'] = df_tareas_activas['Fin_DT'].dt.date
estado_dinamico = {}
carga_dinamica = {}

if not df_tareas_activas.empty:
    for _, m_row in st.session_state.maestro.iterrows():
        nombre_m = m_row['ID_Contratista']
        t_m = df_tareas_activas[df_tareas_activas['Contratista'] == nombre_m]
        en_obra_hoy = t_m[(t_m['Inicio_DT'] <= fecha_hoy) & (t_m['Fin_DT'] >= fecha_hoy)]
        estado_dinamico[nombre_m] = "En Ejecución" if not en_obra_hoy.empty else "Disponible"
        
        t_7d = t_m[(t_m['Inicio_DT'] <= fecha_hoy + timedelta(days=7)) & (t_m['Fin_DT'] >= fecha_hoy)]
        n_frentes = t_7d['Proyecto'].nunique()
        if n_frentes == 0: carga = "Libre"
        elif n_frentes == 1: carga = "Carga Óptima"
        else: carga = f"Saturado ({n_frentes} Frentes)"
        carga_dinamica[nombre_m] = carga
else:
    for _, m_row in st.session_state.maestro.iterrows():
        estado_dinamico[m_row['ID_Contratista']] = "Disponible"
        carga_dinamica[m_row['ID_Contratista']] = "Libre"

# ==========================================
# 4. SIDEBAR & HEADER
# ==========================================
with st.sidebar:
    st.markdown(f"<h2 style='color:#00E5FF; font-family:monospace; text-align:center;'>{st.secrets.get('APP_TITLE', 'COTA APP')} <span style='font-size: 12px; color: #FF9800; background-color: rgba(255, 152, 0, 0.15); border: 1px solid #FF9800; border-radius: 4px; padding: 1px 6px; vertical-align: middle; font-family: monospace;'>BETA</span></h2>", unsafe_allow_html=True)
    socios_input = st.text_input("Directores / Socios (separados por coma)", value="Socio 1, Socio 2")
    lista_socios = [s.strip() for s in socios_input.split(",") if s.strip()]
    if not lista_socios:
        lista_socios = ["Socio 1", "Socio 2"]
    SOCIOS = lista_socios
    st.divider()
    if st.session_state.get('hubo_error_conexion', False):
        st.warning("⚠️ **Conexión Bloqueada**")
        st.info("Se detectaron datos vacíos o un fallo de conexión. Para evitar la pérdida de información en Google Sheets, la sincronización está deshabilitada temporalmente.")
        if st.button("🔄 REINTENTAR CONEXIÓN", type="primary", use_container_width=True):
            with st.spinner("Reintentando conexión..."):
                forzar_importacion()
                st.rerun()
    else:
        if st.button("SINCRONIZAR DATOS", type="primary", use_container_width=True):
            with st.spinner("Actualizando servidores..."):
                g1 = guardar_tabla(st.session_state.proyectos, 'proyectos')
                g2 = guardar_tabla(st.session_state.tasks, 'tasks')
                g3 = guardar_tabla(st.session_state.maestro, 'maestro')
                g_hist = guardar_tabla(st.session_state.df_historial, 'historial')
                g4 = guardar_tabla(st.session_state.gg_log, 'gg_log')
                g5 = guardar_tabla(st.session_state.todo_list, 'logistica')
                g6 = guardar_tabla(st.session_state.pagos, 'pagos_maestros')
                g7 = guardar_tabla(st.session_state.p_info, 'partidas_info')
                g8 = guardar_tabla(st.session_state.ingresos_socios, 'ingresos_socios')
                g9 = guardar_tabla(st.session_state.materiales, 'materiales')
                g13 = guardar_tabla(st.session_state.ice_sesiones, 'ice_sesiones')
                g14 = guardar_tabla(st.session_state.ice_invitados, 'ice_invitados')
                g15 = guardar_tabla(st.session_state.ice_acuerdos, 'ice_acuerdos')
                g16 = guardar_tabla(st.session_state.entregables_planos, 'entregables_planos')
                g17 = guardar_tabla(st.session_state.observaciones_municipales, 'observaciones_municipales')
                g18 = guardar_tabla(st.session_state.entrega_dossier, 'entrega_dossier')
                g19 = guardar_tabla(st.session_state.entrega_unidades, 'entrega_unidades')
                
                # Limpiar claves de dataframes base estables en memoria
                for k in list(st.session_state.keys()):
                    if k.startswith("base_df_"):
                        del st.session_state[k]
                
                # Robust status validation: Critical vs. Secondary
                critical_ok = all([g1, g2, g3])
                secondary_failed = []
                for name, status in [('historial', g_hist), ('gg_log', g4), ('logistica', g5), ('pagos_maestros', g6), 
                                     ('partidas_info', g7), ('ingresos_socios', g8), ('materiales', g9), 
                                     ('ice_sesiones', g13), ('ice_invitados', g14), ('ice_acuerdos', g15),
                                     ('entregables_planos', g16), ('observaciones_municipales', g17),
                                     ('entrega_dossier', g18), ('entrega_unidades', g19)]:
                    if not status:
                        secondary_failed.append(name)

                
                if critical_ok:
                    if secondary_failed:
                        st.warning(f"⚠️ Datos principales sincronizados. Falló guardado secundario en: {', '.join(secondary_failed)}")
                    else:
                        st.success("Sincronización exitosa.")
                else:
                    st.error("❌ Fallo en sincronización de datos críticos (proyectos, tareas o contratistas).")
app_title = st.secrets.get("APP_TITLE", "COTA APP")
logo_src = get_image_base64(st.secrets.get("LOGO_IMAGE", "Logo2.png"))
logo_html = f'<img src="{logo_src}" class="header-logo">' if logo_src else ""
st.markdown(f'''
    <div class="header-container">
        {logo_html}
        <div class="header-divider"></div>
        <div class="header-title-box">
            <h1 class="header-title">{app_title} <span style='font-size: 13px; color: #FF9800; background-color: rgba(255, 152, 0, 0.15); border: 1px solid #FF9800; border-radius: 6px; padding: 2px 8px; margin-left: 10px; vertical-align: middle; font-family: monospace; font-weight: bold;'>BETA</span></h1>
            <div class="header-subtitle">PROJECT & DESIGN MANAGEMENT</div>
        </div>
    </div>''', unsafe_allow_html=True)


def render_proyecto_operaciones_tab(proy_name, proy_data):
    """Renders the operations tab content for a single project in unified view."""
    t_p = st.session_state.tasks[st.session_state.tasks['Proyecto'] == proy_name]
    completadas = t_p['Estado'].astype(str).str.contains('Completado', case=False, na=False)
    ava = f"{int((completadas.sum() / len(t_p)) * 100)}%" if len(t_p) > 0 else "0%"
    
    st.markdown(f"<h4 style='color:#00E5FF; font-family:monospace; margin-bottom:15px;'>DETALLE OPERATIVO // AVANCE GENERAL: {ava}</h4>", unsafe_allow_html=True)
    
    # Materiales del proyecto
    st.markdown("<h5 style='color:#00E5FF; font-family:monospace; margin-top:15px;'>LOGÍSTICA DE MATERIALES</h5>", unsafe_allow_html=True)
    df_mat_pv = st.session_state.materiales[st.session_state.materiales['Proyecto'] == proy_name].copy()
    base_mat_key = f"base_df_mat_{proy_name}"
    widget_mat_key = f"ed_mat_{proy_name}"
    if base_mat_key not in st.session_state:
        df_m_clean = df_mat_pv.drop(columns=['Proyecto'], errors='ignore').reset_index(drop=True)
        for c in ['Material', 'Unidad', 'Proveedor', 'Estado']:
            if c in df_m_clean.columns: df_m_clean[c] = df_m_clean[c].fillna("").astype(str)
            elif c == 'Estado': df_m_clean[c] = "🔴 No Comprado"
        st.session_state[base_mat_key] = df_m_clean
    
    st.data_editor(
        st.session_state[base_mat_key],
        use_container_width=True, hide_index=True, num_rows="dynamic", key=widget_mat_key,
        on_change=on_table_edit,
        args=(base_mat_key, widget_mat_key, 'materiales', ['Proyecto'], [proy_name], ['Fecha'], ['Material', 'Unidad', 'Proveedor', 'Estado'], None, ['Cantidad', 'Ppto_Estimado', 'Ppto_Real'], 'ID'),
        column_config={
            "ID": None,
            "Cantidad": st.column_config.NumberColumn("Cant."),
            "Unidad": st.column_config.SelectboxColumn("Und", options=['und', 'kg', 'm2', 'm3', 'glb', 'm', 'bolsa']),
            "Estado": st.column_config.SelectboxColumn("Estado", options=['🔴 No Comprado', '🟢 Comprado'], default='🔴 No Comprado')
        }
    )

    st.divider()
    # Partidas del proyecto
    with st.popover("➕ AGREGAR PARTIDA"):
        with st.form(f"pv_new_part_{proy_name}"):
            n_pa = st.text_input("Código/Nombre Partida")
            if st.form_submit_button("CREAR") and n_pa:
                st.session_state.partidas = pd.concat([st.session_state.partidas, pd.DataFrame([{'Proyecto': proy_name, 'Nombre_Partida': n_pa}])], ignore_index=True)
                st.session_state.p_info = pd.concat([st.session_state.p_info, pd.DataFrame([{'Proyecto': proy_name, 'Nombre_Partida': n_pa, 'Monto_Pactado': 0.0}])], ignore_index=True)
                st.rerun()

    parts_proy = st.session_state.partidas[st.session_state.partidas['Proyecto'] == proy_name]
    if parts_proy.empty:
        st.info("Sin partidas registradas en este proyecto. Crea una partida arriba para comenzar.")
    else:
        for idx_part, part_r in parts_proy.iterrows():
            p_nom = part_r['Nombre_Partida']
            st.markdown(f"<div style='font-family:monospace; font-size:14px; margin-top:20px; margin-bottom:5px; border-bottom:1px solid #1e293b; padding-bottom:5px;'><b style='color:#00E5FF;'>SECCIÓN: {p_nom.upper()}</b></div>", unsafe_allow_html=True)

            df_pactado = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy_name) & (st.session_state.p_info['Nombre_Partida'] == p_nom)]
            monto_pactado = float(df_pactado['Monto_Pactado'].iloc[0]) if not df_pactado.empty else 0.0
            df_pagos_part = st.session_state.pagos[(st.session_state.pagos['Proyecto'] == proy_name) & (st.session_state.pagos['Partida'] == p_nom)]
            monto_pagado = pd.to_numeric(df_pagos_part['Monto'], errors='coerce').sum() if not df_pagos_part.empty else 0.0
            saldo = monto_pactado - monto_pagado
            col_s = '#10b981' if saldo >= 0 else '#ef4444'

            cf1, cf2 = st.columns([2, 8])
            with cf1.popover("AJUSTAR PRESUPUESTO"):
                with st.form(f"pact_{proy_name}_{p_nom}_{idx_part}"):
                    n_pac = st.number_input("Pactado (S/)", value=monto_pactado)
                    if st.form_submit_button("FIJAR"):
                        if df_pactado.empty:
                            st.session_state.p_info = pd.concat([st.session_state.p_info, pd.DataFrame([{'Proyecto': proy_name, 'Nombre_Partida': p_nom, 'Monto_Pactado': n_pac}])], ignore_index=True)
                        else:
                            st.session_state.p_info.at[df_pactado.index[0], 'Monto_Pactado'] = n_pac
                        st.rerun()
            cf2.markdown(f"**Pactado:** S/{monto_pactado:,.2f} | **Pagado:** S/{monto_pagado:,.2f} | **Saldo:** <span style='color:{col_s};'>S/{saldo:,.2f}</span>", unsafe_allow_html=True)

            mask_p = (st.session_state.tasks['Proyecto'] == proy_name) & (st.session_state.tasks['Partida'] == p_nom)
            df_v = st.session_state.tasks[mask_p].drop(columns=['Proyecto', 'Partida'], errors='ignore').reset_index(drop=True)
            for c in ['Tarea', 'Especialidad', 'Contratista', 'Estado', 'Causa_Cumplimiento', 'Dependencia', 'Notas']:
                if c in df_v.columns: df_v[c] = df_v[c].fillna("").astype(str).replace('nan', '')
            for c in ['Culpa_Contratista', 'Permitir_Cruce']:
                if c in df_v.columns: df_v[c] = df_v[c].apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y']).astype(bool)
            for c in ['Inicio', 'Fin']:
                if c in df_v.columns: df_v[c] = pd.to_datetime(df_v[c], errors='coerce').dt.date
            if 'Orden' in df_v.columns: df_v['Orden'] = pd.to_numeric(df_v['Orden'], errors='coerce').fillna(0).astype(int)
            df_v = df_v.sort_values(by='Inicio', na_position='last').reset_index(drop=True)

            base_tasks_key = f"base_df_tasks_{proy_name}_{p_nom}_{idx_part}"
            widget_tasks_key = f"t_{proy_name}_{p_nom}_{idx_part}"
            if base_tasks_key not in st.session_state:
                st.session_state[base_tasks_key] = df_v.copy()

            st.data_editor(
                st.session_state[base_tasks_key], use_container_width=True, hide_index=True, num_rows="dynamic", key=widget_tasks_key,
                on_change=on_table_edit,
                args=(base_tasks_key, widget_tasks_key, 'tasks', ['Proyecto', 'Partida'], [proy_name, p_nom], ['Inicio', 'Fin'], ['Tarea', 'Especialidad', 'Contratista', 'Estado', 'Causa_Cumplimiento', 'Dependencia', 'Notas'], ['Culpa_Contratista', 'Permitir_Cruce'], ['Orden'], None),
                column_config={
                    "Especialidad": st.column_config.SelectboxColumn(options=extraer_esp()),
                    "Contratista": st.column_config.SelectboxColumn(options=st.session_state.maestro['ID_Contratista'].tolist()),
                    "Estado": st.column_config.SelectboxColumn(options=status_tasks),
                    "Inicio": st.column_config.DateColumn("Inicio", format="DD/MM/YYYY"),
                    "Fin": st.column_config.DateColumn("Fin", format="DD/MM/YYYY")
                }
            )

            if st.button("ASISTENTE AI", key=f"ai_{proy_name}_{p_nom}_{idx_part}"):
                for _, row_ai in st.session_state[base_tasks_key].iterrows():
                    if pd.notna(row_ai.get('Especialidad')) and row_ai.get('Especialidad') != "Ninguna":
                        sug = sugerir_contratista(row_ai['Especialidad'], proy_data.get('Ubicación', 'Lima'), row_ai.get('Inicio'), row_ai.get('Fin'))
                        if sug: st.success(f"{row_ai.get('Tarea')}: {sug}")


# ==========================================
# 5. VISTA DE CONTROL FINANCIERO FULLSCREEN
# ==========================================
if st.session_state.vista_actual in ['proyecto', 'dashboard'] and st.session_state.proy_activo:
    proy_name = st.session_state.proy_activo
    proy_data = st.session_state.proyectos[st.session_state.proyectos['Nombre'] == proy_name].iloc[0]
    
    if st.button("⬅ VOLVER A MATRIZ PRINCIPAL"): st.session_state.vista_actual = 'main'; st.rerun()

    lbl_m_pv = ""
    st.markdown(f"<h3 style='color: {proy_data.get('Color', '#00E5FF')}; margin-top: 15px; font-family: monospace;'>PROYECTO: {str(proy_name).upper()}{lbl_m_pv}</h3>", unsafe_allow_html=True)
    st.divider()
    
    tab_pv_ops, tab_pv_dash, tab_pv_ice = st.tabs(["[ OPERACIONES ]", "[ DASHBOARD FINANCIERO ]", "[ SESIONES ICE ]"])
    
    with tab_pv_ops:
        render_proyecto_operaciones_tab(proy_name, proy_data)
        
    with tab_pv_dash:
        render_seccion_ejecucion_obra(proy_name, proy_data)
        
    with tab_pv_ice:
        render_gestor_ice_sesiones(proy_name)

# ==========================================
# 7. MATRIZ GENERAL Y FINANZAS
# ==========================================
elif st.session_state.vista_actual == 'main':
    # Listas Dinámicas
    exist_divs = sorted(st.session_state.proyectos['Division'].dropna().unique().tolist())
    if "General" not in exist_divs: exist_divs.insert(0, "General")
    
    exist_ubis = sorted(st.session_state.proyectos['Ubicación'].dropna().unique().tolist())
    if not exist_ubis: exist_ubis = ["Lima", "Mala"]

    t_mat, t_con, t_fin, t_ope, t_cal, t_arc = st.tabs(["[ OPERACIONES ]", "[ BASE CONTRATISTAS ]", "[ FINANZAS ]", "[ CONTROL OPERATIVO ]", "[ CALENDARIO ]", "[ ARCHIVO ]"])

    with t_mat:
        with st.popover("NUEVO PROYECTO"):
            with st.form("f_new_p", clear_on_submit=True):
                pn = st.text_input("Nombre de Proyecto")
                sel_div = st.selectbox("Carpeta", exist_divs + ["+ Crear Nueva Carpeta"])
                p_div = st.text_input("Nombre de Nueva Carpeta") if sel_div == "+ Crear Nueva Carpeta" else sel_div
                sel_ubi = st.selectbox("Ubicación", exist_ubis + ["+ Agregar Ubicación"])
                p_ubi = st.text_input("Nueva Ubicación") if sel_ubi == "+ Agregar Ubicación" else sel_ubi
                pe = st.selectbox("Líder Asignado", lista_socios)
                ps = st.selectbox("Estado Actual", status_proyectos)
                
                p_tipo = st.selectbox("Tipo de Proyecto", ["Unifamiliar", "Multifamiliar", "Comercial (Próximamente)"])
                p_pisos = 0
                p_sotanos = 0
                if p_tipo == "Multifamiliar":
                    col_p1, col_p2 = st.columns(2)
                    p_pisos = col_p1.number_input("Número de Pisos", min_value=1, value=5, step=1)
                    p_sotanos = col_p2.number_input("Número de Sótanos", min_value=0, value=1, step=1)
                elif p_tipo == "Comercial (Próximamente)":
                    st.warning("⚠️ Los proyectos comerciales no están disponibles por ahora.")
                
                p_cliente = st.text_input("Cliente")
                p_color = st.color_picker("Color Identificador (Calendario)", value="#00E5FF")
                
                submit_disabled = p_tipo == "Comercial (Próximamente)"
                if st.form_submit_button("CONFIRMAR CREACIÓN", disabled=submit_disabled):
                    n_p = pd.DataFrame([{
                        'Nombre': pn, 
                        'Encargado': pe, 
                        'Ubicación': p_ubi, 
                        'Estado': ps, 
                        'Inicio': fecha_hoy.strftime('%Y-%m-%d'), 
                        'Fin': fecha_hoy.strftime('%Y-%m-%d'), 
                        'Color': p_color, 
                        'Presupuesto': 0, 
                        'Utilidades': 0, 
                        'GG_Presupuesto': 0, 
                        'Archivado': False, 
                        'Metodologia': 'Tradicional', 
                        'Division': p_div,
                        'Tipo_Proyecto': p_tipo,
                        'Num_Pisos': int(p_pisos),
                        'Num_Sotanos': int(p_sotanos),
                        'Cliente': p_cliente
                    }])
                    st.session_state.proyectos = pd.concat([st.session_state.proyectos, n_p], ignore_index=True); st.rerun()

        st.divider()
        search_query = st.text_input("🔍 Buscar proyecto", placeholder="Buscar por nombre o cliente...")
        
        p_vis = st.session_state.proyectos[st.session_state.proyectos['Archivado'] != True].copy()
        
        if 'Cliente' not in p_vis.columns:
            p_vis['Cliente'] = ""
            
        if search_query:
            sq = search_query.lower()
            p_vis = p_vis[
                p_vis['Nombre'].astype(str).str.lower().str.contains(sq, na=False, regex=False) |
                p_vis['Cliente'].astype(str).str.lower().str.contains(sq, na=False, regex=False)
            ]
        
        if 'Division' not in p_vis.columns: p_vis['Division'] = "General"
        p_vis['Division'] = p_vis['Division'].fillna("General").replace("", "General")
        
        divisiones = sorted(p_vis['Division'].unique().tolist())
        if not divisiones: divisiones = ["General"]
        
        tabs_div = st.tabs([f"[ CARPETA: {d.upper()} ]" for d in divisiones])
        
        for t_div, d_name in zip(tabs_div, divisiones):
            with t_div:
                p_div = p_vis[p_vis['Division'] == d_name]
                
                for est in status_proyectos:
                    est_clean = est.split(" ", 1)[-1] if " " in est else est
                    p_est = p_div[p_div['Estado'].astype(str).str.contains(est_clean, case=False, na=False)]
                    
                    show_state = False
                    if not p_est.empty:
                        show_state = st.toggle(f"{est.upper()} ({len(p_est)} proyectos)", value=False, key=f"tgl_{d_name}_{est}")
                    if show_state:
                        color_status = "#cbd5e1"
                        if "Planificación" in est: color_status = "#eab308"
                        elif "Proceso" in est: color_status = "#00E5FF"
                        elif "Completado" in est: color_status = "#10b981"
                        
                        st.markdown(f"""
                            <h4 class='xian-status-header' style='color: {color_status};'>
                                <span class='xian-status-header-bar' style='background-color: {color_status};'></span>
                                ESTADO: {est.upper()}
                            </h4>
                        """, unsafe_allow_html=True)
                        
                        for idx_p, proy in p_est.iterrows():
                            t_p = st.session_state.tasks[st.session_state.tasks['Proyecto'] == proy['Nombre']]
                            completadas = t_p['Estado'].astype(str).str.contains('Completado', case=False, na=False)
                            ava = f"{int((completadas.sum() / len(t_p)) * 100)}%" if len(t_p) > 0 else "0%"
                            
                            lbl_m = ""
                            if "Planificación" in str(proy['Estado']):
                                st_col = "#eab308"
                            elif "Completado" in str(proy['Estado']):
                                st_col = "#10b981"
                            else:
                                st_col = "#00E5FF"
                            tit = f"> PROYECTO: {str(proy['Nombre']).upper()} {lbl_m} | LÍDER: {proy.get('Encargado', 'N/A')} | AVANCE: {ava}"
                            
                            with st.expander(tit, expanded=False, key=f"exp_proy_{proy['Nombre']}"):
                                st.markdown(f"<div style='border-left: 3px solid {st_col}; background: {st_col}06; border-top: 1px solid var(--border-color); border-right: 1px solid var(--border-color); border-bottom: 1px solid var(--border-color); border-radius: 0 4px 4px 0; padding: 12px 15px; margin-bottom: 20px;'>", unsafe_allow_html=True)
                                cb1, cb2, cb3, cb4, cb5, _ = st.columns([2, 1.5, 1.5, 1.5, 2, 1.5])
                                
                                if cb1.button("ABRIR PROYECTO", key=f"d_{idx_p}"): st.session_state.proy_activo = proy['Nombre']; st.session_state.vista_actual = 'proyecto'; st.rerun()
                                
                                with cb2.popover("INFO"):
                                    with st.form(f"e_{idx_p}"):
                                        n_n = st.text_input("Identificador", value=proy['Nombre'])
                                        
                                        val_cliente = proy.get('Cliente', '')
                                        if pd.isna(val_cliente): val_cliente = ""
                                        n_cliente = st.text_input("Cliente", value=str(val_cliente))
                                        
                                        val_color = proy.get('Color', '#00E5FF')
                                        if pd.isna(val_color) or not str(val_color).startswith('#'): val_color = '#00E5FF'
                                        n_color = st.color_picker("Color Identificador (Calendario)", value=str(val_color))
                                        
                                        sel_d = st.selectbox("Carpeta", exist_divs + ["+ Crear Nueva Carpeta"], index=exist_divs.index(proy.get('Division', 'General')) if proy.get('Division', 'General') in exist_divs else 0)
                                        n_div = st.text_input("Nombre de Nueva Carpeta") if sel_d == "+ Crear Nueva Carpeta" else sel_d
                                        
                                        sel_u = st.selectbox("Ubicación", exist_ubis + ["+ Agregar Ubicación"], index=exist_ubis.index(proy.get('Ubicación', 'Lima')) if proy.get('Ubicación', 'Lima') in exist_ubis else 0)
                                        n_ubi = st.text_input("Nueva Ubicación") if sel_u == "+ Agregar Ubicación" else sel_u

                                        curr_est = proy['Estado'] if proy['Estado'] in status_proyectos else [s for s in status_proyectos if est_clean in s][0]
                                        n_s = st.selectbox("Estatus", status_proyectos, index=status_proyectos.index(curr_est))
                                        
                                        tipo_opts = ["Unifamiliar", "Multifamiliar", "Comercial (Próximamente)"]
                                        curr_tipo = proy.get('Tipo_Proyecto', 'Unifamiliar')
                                        if "Comercial" in str(curr_tipo):
                                            curr_tipo = "Comercial (Próximamente)"
                                        tipo_idx = tipo_opts.index(curr_tipo) if curr_tipo in tipo_opts else 0
                                        n_tipo = st.selectbox("Tipo de Proyecto", tipo_opts, index=tipo_idx)
                                        
                                        n_pisos = 0
                                        n_sotanos = 0
                                        if n_tipo == "Multifamiliar":
                                            col_e1, col_e2 = st.columns(2)
                                            try:
                                                curr_pisos = int(proy.get('Num_Pisos', 5))
                                                if pd.isna(curr_pisos) or curr_pisos < 1: curr_pisos = 5
                                            except:
                                                curr_pisos = 5
                                            try:
                                                curr_sotanos = int(proy.get('Num_Sotanos', 1))
                                                if pd.isna(curr_sotanos) or curr_sotanos < 0: curr_sotanos = 1
                                            except:
                                                curr_sotanos = 1
                                            
                                            n_pisos = col_e1.number_input("Número de Pisos", min_value=1, value=curr_pisos, step=1, key=f"edit_p_pisos_{idx_p}")
                                            n_sotanos = col_e2.number_input("Número de Sótanos", min_value=0, value=curr_sotanos, step=1, key=f"edit_p_sotanos_{idx_p}")
                                        elif n_tipo == "Comercial (Próximamente)":
                                            st.warning("⚠️ Los proyectos comerciales no están disponibles por ahora.")
                                        
                                        f_i, f_f = st.columns(2)
                                        d_i = f_i.date_input("Inicio", value=pd.to_datetime(proy.get('Inicio', fecha_hoy)).date())
                                        d_f = f_f.date_input("Fin", value=pd.to_datetime(proy.get('Fin', fecha_hoy)).date())
                                        v_p = st.number_input("Ppto Venta", value=safe_float(proy.get('Presupuesto', 0)))
                                        v_u = st.number_input("Utilidad Meta", value=safe_float(proy.get('Utilidades', 0)))
                                        v_g = st.number_input("Ppto Caja Chica", value=safe_float(proy.get('GG_Presupuesto', 0)))
                                        
                                        submit_edit_disabled = n_tipo == "Comercial (Próximamente)"
                                        if st.form_submit_button("ACTUALIZAR PARAMETROS", disabled=submit_edit_disabled):
                                            st.session_state.proyectos.at[idx_p, 'Nombre'] = n_n
                                            st.session_state.proyectos.at[idx_p, 'Division'], st.session_state.proyectos.at[idx_p, 'Ubicación'] = n_div, n_ubi
                                            st.session_state.proyectos.at[idx_p, 'Estado'], st.session_state.proyectos.at[idx_p, 'Inicio'], st.session_state.proyectos.at[idx_p, 'Fin'] = n_s, d_i.strftime('%Y-%m-%d'), d_f.strftime('%Y-%m-%d')
                                            st.session_state.proyectos.at[idx_p, 'Presupuesto'], st.session_state.proyectos.at[idx_p, 'Utilidades'], st.session_state.proyectos.at[idx_p, 'GG_Presupuesto'] = v_p, v_u, v_g
                                            st.session_state.proyectos.at[idx_p, 'Tipo_Proyecto'] = n_tipo
                                            st.session_state.proyectos.at[idx_p, 'Num_Pisos'] = int(n_pisos) if n_tipo == 'Multifamiliar' else 0
                                            st.session_state.proyectos.at[idx_p, 'Num_Sotanos'] = int(n_sotanos) if n_tipo == 'Multifamiliar' else 0
                                            st.session_state.proyectos.at[idx_p, 'Cliente'] = n_cliente
                                            st.session_state.proyectos.at[idx_p, 'Color'] = n_color
                                            st.rerun()
                                            
                                if cb3.button("ARCHIVAR", key=f"a_{idx_p}", type="primary"): 
                                    st.session_state.proyectos.at[idx_p, 'Archivado'] = True
                                    st.rerun()
                                
                                with cb4.popover("AGREGAR PARTIDA"):
                                    with st.form(f"ap_{idx_p}"):
                                        n_pa = st.text_input("Código/Nombre Partida")
                                        if st.form_submit_button("CREAR") and n_pa:
                                            st.session_state.partidas = pd.concat([st.session_state.partidas, pd.DataFrame([{'Proyecto': proy['Nombre'], 'Nombre_Partida': n_pa}])], ignore_index=True)
                                            st.session_state.p_info = pd.concat([st.session_state.p_info, pd.DataFrame([{'Proyecto': proy['Nombre'], 'Nombre_Partida': n_pa, 'Monto_Pactado': 0.0}])], ignore_index=True); st.rerun()

                                with cb5.popover("AGREGAR NOTAS"):
                                    with st.form(f"f_notes_{idx_p}"):
                                        curr_notas = proy.get('Notas', '')
                                        if pd.isna(curr_notas): curr_notas = ""
                                        n_notas = st.text_area("Notas del Proyecto (Spicy Gala AI)", value=str(curr_notas), height=150, key=f"notes_ta_{idx_p}")
                                        if st.form_submit_button("GUARDAR NOTAS"):
                                            st.session_state.proyectos.at[idx_p, 'Notas'] = n_notas
                                            st.rerun()

                                st.markdown("</div>", unsafe_allow_html=True)

                                # TAREAS Y PARTIDAS
                                parts = st.session_state.partidas[st.session_state.partidas['Proyecto'] == proy['Nombre']]
                                is_multi = proy.get('Tipo_Proyecto') == 'Multifamiliar'

                                if is_multi:
                                    proy_name = proy['Nombre']
                                    proy_data = proy
                                    
                                    t_exp_diseno, t_exp_ejecucion, t_exp_entrega = st.tabs([
                                        "📐 ETAPA 1: DISEÑO",
                                        "🏗️ ETAPA 2: CONSTRUCCIÓN",
                                        "🔑 ETAPA 3: ENTREGA"
                                    ])
                                    
                                    with t_exp_diseno:
                                        render_gestor_ice_sesiones(proy_name)
                                        
                                        st.divider()
                                        st.markdown("### 📋 CONTROL DE ENTREGABLES (PLANOS Y DOCUMENTOS)")
                                        with st.popover("➕ REGISTRAR NUEVO ENTREGABLE", use_container_width=True):
                                            with st.form(f"form_nuevo_entregable_{proy_name}", clear_on_submit=True):
                                                col_e1, col_e2 = st.columns(2)
                                                e_esp = col_e1.selectbox("Especialidad", ["Arquitectura", "Estructuras", "Inst. Sanitarias", "Inst. Eléctricas", "Inst. Mecánicas", "Inst. Especiales", "Otros"], key=f"ent_esp_add_{proy_name}")
                                                e_nom = col_e2.text_input("Nombre del Entregable", key=f"ent_nom_add_{proy_name}")
                                                
                                                col_e3, col_e4 = st.columns(2)
                                                e_rev = col_e3.text_input("Revisión", value="0", key=f"ent_rev_add_{proy_name}")
                                                e_est = col_e4.selectbox("Estado Inicial", ["En desarrollo", "En revisión", "Aprobado", "Observado"], key=f"ent_est_add_{proy_name}")
                                                
                                                col_e5, col_e6 = st.columns(2)
                                                e_f_lim = col_e5.date_input("Fecha Límite", value=datetime.today(), key=f"ent_flim_add_{proy_name}")
                                                e_resp = col_e6.text_input("Responsable", key=f"ent_resp_add_{proy_name}")
                                                
                                                if st.form_submit_button("REGISTRAR ENTREGABLE") and e_nom:
                                                    new_id = f"ENT_{int(datetime.now().timestamp())}_{uuid.uuid4().hex[:6]}"
                                                    new_row = pd.DataFrame([{
                                                        'ID': new_id,
                                                        'Proyecto': proy_name,
                                                        'Especialidad': e_esp,
                                                        'Nombre_Entregable': e_nom,
                                                        'Revision': e_rev,
                                                        'Estado': e_est,
                                                        'Fecha_Limite': e_f_lim.strftime("%d/%m/%Y"),
                                                        'Responsable': e_resp
                                                    }])
                                                    st.session_state.entregables_planos = pd.concat([st.session_state.entregables_planos, new_row], ignore_index=True)
                                                    st.success("Entregable registrado localmente. Recuerde sincronizar.")
                                                    st.rerun()
                                                    
                                        df_ent = st.session_state.entregables_planos[st.session_state.entregables_planos['Proyecto'] == proy_name].copy()
                                        df_ent_edit = df_ent.drop(columns=['Proyecto']) if not df_ent.empty else pd.DataFrame(columns=['ID', 'Especialidad', 'Nombre_Entregable', 'Revision', 'Estado', 'Fecha_Limite', 'Responsable'])
                                        for c in ['Especialidad', 'Nombre_Entregable', 'Revision', 'Estado', 'Fecha_Limite', 'Responsable']:
                                            if c not in df_ent_edit.columns: df_ent_edit[c] = ""
                                            
                                        df_ent_edit['Fecha_Limite'] = pd.to_datetime(df_ent_edit['Fecha_Limite'], format="%d/%m/%Y", errors='coerce').dt.date
                                        
                                        base_key = f"base_df_ent_{proy_name}"
                                        if base_key not in st.session_state:
                                            st.session_state[base_key] = df_ent_edit.copy()
                                            
                                        upd_ent = st.data_editor(
                                            st.session_state[base_key], use_container_width=True, hide_index=True, num_rows="dynamic", key=f"ent_editor_{proy_name}",
                                            column_config={
                                                "ID": None,
                                                "Estado": st.column_config.SelectboxColumn("Estado", options=["En desarrollo", "En revisión", "Aprobado", "Observado"]),
                                                "Especialidad": st.column_config.SelectboxColumn("Especialidad", options=["Arquitectura", "Estructuras", "Inst. Sanitarias", "Inst. Eléctricas", "Inst. Mecánicas", "Inst. Especiales", "Otros"]),
                                                "Fecha_Limite": st.column_config.DateColumn("Fecha Límite", format="DD/MM/YYYY")
                                            },
                                            on_change=on_table_edit,
                                            args=(
                                                base_key,
                                                f"ent_editor_{proy_name}",
                                                'entregables_planos',
                                                ['Proyecto'],
                                                [proy_name],
                                                ['Fecha_Limite'],
                                                ['Especialidad', 'Nombre_Entregable', 'Revision', 'Estado', 'Responsable'],
                                                None,
                                                None,
                                                'ID'
                                            )
                                        )

                                        st.divider()
                                        st.markdown("### ⚠️ LEVANTAMIENTO DE OBSERVACIONES (MUNICIPALIDAD / OTROS)")
                                        with st.popover("➕ REGISTRAR NUEVA OBSERVACIÓN", use_container_width=True):
                                            with st.form(f"form_nueva_observacion_{proy_name}", clear_on_submit=True):
                                                col_o1, col_o2 = st.columns(2)
                                                o_nro = col_o1.text_input("Nro. Observación (Ej: OBS-01)", key=f"obs_nro_add_{proy_name}")
                                                o_esp = col_o2.selectbox("Especialidad", ["Arquitectura", "Estructuras", "Inst. Sanitarias", "Inst. Eléctricas", "Inst. Mecánicas", "Inst. Especiales", "Otros"], key=f"obs_esp_add_{proy_name}")
                                                
                                                o_desc = st.text_area("Descripción de la Observación", key=f"obs_desc_add_{proy_name}")
                                                
                                                col_o3, col_o4 = st.columns(2)
                                                o_est = col_o3.selectbox("Estado Inicial", ["Pendiente", "Enviado", "Subsanado", "Aprobado"], key=f"obs_est_add_{proy_name}")
                                                o_f_ing = col_o4.date_input("Fecha de Ingreso", value=datetime.today(), key=f"obs_fing_add_{proy_name}")
                                                
                                                col_o5, col_o6 = st.columns(2)
                                                o_resp = col_o5.text_input("Responsable", key=f"obs_resp_add_{proy_name}")
                                                o_notas = col_o6.text_input("Notas Adicionales", key=f"obs_notas_add_{proy_name}")
                                                
                                                if st.form_submit_button("REGISTRAR OBSERVACIÓN") and o_desc:
                                                    new_id = f"OBS_{int(datetime.now().timestamp())}_{uuid.uuid4().hex[:6]}"
                                                    new_row = pd.DataFrame([{
                                                        'ID': new_id,
                                                        'Proyecto': proy_name,
                                                        'Nro_Observacion': o_nro,
                                                        'Especialidad': o_esp,
                                                        'Descripcion': o_desc,
                                                        'Estado': o_est,
                                                        'Fecha_Ingreso': o_f_ing.strftime("%d/%m/%Y"),
                                                        'Fecha_Subsanacion': "",
                                                        'Responsable': o_resp,
                                                        'Notas': o_notas
                                                    }])
                                                    st.session_state.observaciones_municipales = pd.concat([st.session_state.observaciones_municipales, new_row], ignore_index=True)
                                                    st.success("Observación registrada localmente. Recuerde sincronizar.")
                                                    st.rerun()
                                                    
                                        df_obs = st.session_state.observaciones_municipales[st.session_state.observaciones_municipales['Proyecto'] == proy_name].copy()
                                        df_obs_edit = df_obs.drop(columns=['Proyecto']) if not df_obs.empty else pd.DataFrame(columns=['ID', 'Nro_Observacion', 'Especialidad', 'Descripcion', 'Estado', 'Fecha_Ingreso', 'Fecha_Subsanacion', 'Responsable', 'Notas'])
                                        for c in ['Nro_Observacion', 'Especialidad', 'Descripcion', 'Estado', 'Fecha_Ingreso', 'Fecha_Subsanacion', 'Responsable', 'Notas']:
                                            if c not in df_obs_edit.columns: df_obs_edit[c] = ""
                                            
                                        df_obs_edit['Fecha_Ingreso'] = pd.to_datetime(df_obs_edit['Fecha_Ingreso'], format="%d/%m/%Y", errors='coerce').dt.date
                                        df_obs_edit['Fecha_Subsanacion'] = pd.to_datetime(df_obs_edit['Fecha_Subsanacion'], format="%d/%m/%Y", errors='coerce').dt.date
                                            
                                        base_key = f"base_df_obs_{proy_name}"
                                        if base_key not in st.session_state:
                                            st.session_state[base_key] = df_obs_edit.copy()
                                            
                                        upd_obs = st.data_editor(
                                            st.session_state[base_key], use_container_width=True, hide_index=True, num_rows="dynamic", key=f"obs_editor_{proy_name}",
                                            column_config={
                                                "ID": None,
                                                "Estado": st.column_config.SelectboxColumn("Estado", options=["Pendiente", "Enviado", "Subsanado", "Aprobado"]),
                                                "Especialidad": st.column_config.SelectboxColumn("Especialidad", options=["Arquitectura", "Estructuras", "Inst. Sanitarias", "Inst. Eléctricas", "Inst. Mecánicas", "Inst. Especiales", "Otros"]),
                                                "Fecha_Ingreso": st.column_config.DateColumn("Fecha de Ingreso", format="DD/MM/YYYY"),
                                                "Fecha_Subsanacion": st.column_config.DateColumn("Fecha de Subsanación", format="DD/MM/YYYY")
                                            },
                                            on_change=on_table_edit,
                                            args=(
                                                base_key,
                                                f"obs_editor_{proy_name}",
                                                'observaciones_municipales',
                                                ['Proyecto'],
                                                [proy_name],
                                                ['Fecha_Ingreso', 'Fecha_Subsanacion'],  # date_cols
                                                ['Nro_Observacion', 'Especialidad', 'Descripcion', 'Estado', 'Responsable', 'Notas'],  # text_cols
                                                None,  # bool_cols
                                                None,  # num_cols
                                                'ID'   # id_col
                                            )
                                        )
                                            
                                    with t_exp_ejecucion:
                                        # 1. Finanzas por Partida
                                        for idx_part, part in parts.iterrows():
                                            p_nom = part['Nombre_Partida']
                                            st.markdown(f"""
                                                <div class='xian-section-header'>
                                                    <span class='xian-section-header-icon'>■</span> PARTIDA // {p_nom.upper()}
                                                </div>
                                            """, unsafe_allow_html=True)
                                            
                                            df_pactado = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy['Nombre']) & (st.session_state.p_info['Nombre_Partida'] == p_nom)]
                                            monto_pactado = float(df_pactado['Monto_Pactado'].iloc[0]) if not df_pactado.empty else 0.0
                                            df_pagos_p = st.session_state.pagos[(st.session_state.pagos['Proyecto'] == proy['Nombre']) & (st.session_state.pagos['Partida'] == p_nom)]
                                            monto_pagado = pd.to_numeric(df_pagos_p['Monto'], errors='coerce').sum() if not df_pagos_p.empty else 0.0
                                            saldo = monto_pactado - monto_pagado
                                            col_s = '#10b981' if saldo >= 0 else '#ef4444'
                                            
                                            cf1, cf2, cf3, cf4, cf5 = st.columns([1.5, 1.8, 1.8, 1.5, 5.4])
                                            with cf1.popover("✏️ RENOMBRAR"):
                                                with st.form(f"ren_{idx_p}_{p_nom}_{idx_part}"):
                                                    nuevo_nombre = st.text_input("Nuevo Nombre de Partida", value=p_nom)
                                                    if st.form_submit_button("GUARDAR") and nuevo_nombre and nuevo_nombre != p_nom:
                                                        st.session_state.partidas.at[idx_part, 'Nombre_Partida'] = nuevo_nombre
                                                        p_info_matches = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy['Nombre']) & (st.session_state.p_info['Nombre_Partida'] == p_nom)]
                                                        for idx in p_info_matches.index:
                                                            st.session_state.p_info.at[idx, 'Nombre_Partida'] = nuevo_nombre
                                                        tasks_matches = st.session_state.tasks[(st.session_state.tasks['Proyecto'] == proy['Nombre']) & (st.session_state.tasks['Partida'] == p_nom)]
                                                        for idx in tasks_matches.index:
                                                            st.session_state.tasks.at[idx, 'Partida'] = nuevo_nombre
                                                        pagos_matches = st.session_state.pagos[(st.session_state.pagos['Proyecto'] == proy['Nombre']) & (st.session_state.pagos['Partida'] == p_nom)]
                                                        for idx in pagos_matches.index:
                                                            st.session_state.pagos.at[idx, 'Partida'] = nuevo_nombre
                                                        clear_base_dfs()
                                                        st.rerun()
                                            
                                            with cf2.popover("AJUSTAR PRESUPUESTO"):
                                                with st.form(f"pact_{idx_p}_{p_nom}_{idx_part}"):
                                                    n_pac = st.number_input("Pactado (S/)", value=monto_pactado)
                                                    if st.form_submit_button("FIJAR"):
                                                        if df_pactado.empty: st.session_state.p_info = pd.concat([st.session_state.p_info, pd.DataFrame([{'Proyecto': proy['Nombre'], 'Nombre_Partida': p_nom, 'Monto_Pactado': n_pac}])], ignore_index=True)
                                                        else: st.session_state.p_info.at[df_pactado.index[0], 'Monto_Pactado'] = n_pac
                                                        clear_base_dfs()
                                                        st.rerun()
                                            
                                            with cf3.popover("INGRESAR PAGO"):
                                                with st.form(f"pago_{idx_p}_{p_nom}_{idx_part}"):
                                                    p_contra = st.selectbox("Contratista", st.session_state.maestro['ID_Contratista'].tolist(), key=f"sel_p_contra_{idx_p}_{p_nom}_{idx_part}")
                                                    p_fecha = st.date_input("Fecha", value=pd.to_datetime(fecha_hoy).date(), key=f"sel_p_fecha_{idx_p}_{p_nom}_{idx_part}")
                                                    p_monto = st.number_input("Monto (S/)", value=0.0, format="%.2f", key=f"sel_p_monto_{idx_p}_{p_nom}_{idx_part}")
                                                    p_tipo = st.selectbox("Tipo de Pago", ["Adelanto", "Valorización", "Liquidación", "Adicional", "Otros"], key=f"sel_p_tipo_{idx_p}_{p_nom}_{idx_part}")
                                                    p_notas = st.text_input("Notas", value="", key=f"sel_p_notas_{idx_p}_{p_nom}_{idx_part}")
                                                    if st.form_submit_button("INGRESAR PAGO") and p_monto != 0.0:
                                                        n_pago_row = pd.DataFrame([{
                                                            'ID': datetime.now().timestamp(),
                                                            'Proyecto': proy['Nombre'],
                                                            'Partida': p_nom,
                                                            'Contratista': p_contra,
                                                            'Fecha': p_fecha.strftime("%d/%m/%Y"),
                                                            'Monto': p_monto,
                                                            'Tipo': p_tipo,
                                                            'Notas': p_notas
                                                        }])
                                                        st.session_state.pagos = pd.concat([st.session_state.pagos, n_pago_row[cols_pagos]], ignore_index=True)
                                                        clear_base_dfs()
                                                        st.rerun()
                                            
                                            with cf4.popover("🗑️ ELIMINAR"):
                                                st.warning(f"¿Seguro que desea eliminar la partida '{p_nom}'? Se eliminarán también sus tareas y presupuestos.")
                                                if st.button("CONFIRMAR ELIMINACIÓN", key=f"del_part_{idx_p}_{p_nom}_{idx_part}", type="primary", use_container_width=True):
                                                    st.session_state.partidas = st.session_state.partidas.drop(idx_part)
                                                    rem_count = len(st.session_state.partidas[(st.session_state.partidas['Proyecto'] == proy['Nombre']) & (st.session_state.partidas['Nombre_Partida'] == p_nom)])
                                                    p_info_matches = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy['Nombre']) & (st.session_state.p_info['Nombre_Partida'] == p_nom)]
                                                    if len(p_info_matches) > rem_count:
                                                        drop_indices = p_info_matches.index[rem_count:]
                                                        st.session_state.p_info = st.session_state.p_info.drop(drop_indices)
                                                    if rem_count == 0:
                                                        st.session_state.tasks = st.session_state.tasks[~((st.session_state.tasks['Proyecto'] == proy['Nombre']) & (st.session_state.tasks['Partida'] == p_nom))]
                                                    st.rerun()
                                            
                                            cls_saldo = 'xian-badge-saldo-ok' if saldo >= 0 else 'xian-badge-saldo-error'
                                            badge_html = f"""
                                            <div class='xian-badge-container'>
                                                <span class='xian-badge xian-badge-pactado'>PACTADO: S/ {monto_pactado:,.2f}</span>
                                                <span class='xian-badge xian-badge-pagado'>PAGADO: S/ {monto_pagado:,.2f}</span>
                                                <span class='xian-badge {cls_saldo}'>SALDO: S/ {saldo:,.2f}</span>
                                            </div>
                                            """
                                            cf5.markdown(badge_html, unsafe_allow_html=True)
                                            
                                        # 2. Formulario Tren de Trabajo
                                        if not parts.empty:
                                            st.markdown("<hr><h4 style='color:#00E5FF; font-family:monospace; margin-top: 25px;'>🎛️ CONTROL DE PROGRESO (TREN DE TRABAJO)</h4>", unsafe_allow_html=True)
                                            num_pisos = int(proy.get('Num_Pisos', 5))
                                            num_sotanos = int(proy.get('Num_Sotanos', 1))
                                            niveles_opts = ["No iniciado"] + [f"Sótano {i}" for i in range(num_sotanos, 0, -1)] + [f"Piso {i}" for i in range(1, num_pisos + 1)] + ["Terminado"]
                                            niveles_rango_opts = [f"Sótano {i}" for i in range(num_sotanos, 0, -1)] + [f"Piso {i}" for i in range(1, num_pisos + 1)]
                                            
                                            contractors_list = ["Ninguno"] + st.session_state.maestro['ID_Contratista'].tolist()
                                            estado_opts = ["No iniciado", "En Proceso", "Completado"]
                                            etapas_opts = ["Subestructura", "Superestructura", "Acabados", "Instalaciones", "General"]
                                            
                                            with st.form(key=f"form_tren_{proy['Nombre']}"):
                                                col_h1, col_h1b, col_h1c, col_h1d, col_h2, col_h3, col_h4, col_h5, col_h6 = st.columns([1.5, 1.5, 1.5, 1.5, 1.5, 1.5, 1.2, 1.2, 1.6])
                                                col_h1.markdown("**Partida**")
                                                col_h1b.markdown("**Etapa**")
                                                col_h1c.markdown("**Nivel Inicio**")
                                                col_h1d.markdown("**Nivel Fin**")
                                                col_h2.markdown("**Nivel Actual**")
                                                col_h3.markdown("**Contratista**")
                                                col_h4.markdown("**Inicio**")
                                                col_h5.markdown("**Fin**")
                                                col_h6.markdown("**Estado**")
                                                
                                                updated_data = []
                                                
                                                for idx_part, part in parts.iterrows():
                                                    p_nom = part['Nombre_Partida']
                                                    df_p_info = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy['Nombre']) & (st.session_state.p_info['Nombre_Partida'] == p_nom)]
                                                    
                                                    if not df_p_info.empty:
                                                        curr_nivel = df_p_info['Nivel_Actual'].iloc[0]
                                                        curr_contra = df_p_info['Contratista'].iloc[0]
                                                        curr_ini = df_p_info['Inicio'].iloc[0]
                                                        curr_fin = df_p_info['Fin'].iloc[0]
                                                        curr_est = df_p_info['Estado'].iloc[0]
                                                        curr_etapa = df_p_info['Etapa'].iloc[0] if 'Etapa' in df_p_info.columns else "Subestructura"
                                                        curr_n_ini = df_p_info['Nivel_Inicio'].iloc[0] if 'Nivel_Inicio' in df_p_info.columns else ""
                                                        curr_n_fin = df_p_info['Nivel_Fin'].iloc[0] if 'Nivel_Fin' in df_p_info.columns else ""
                                                    else:
                                                        curr_nivel = "No iniciado"
                                                        curr_contra = "Ninguno"
                                                        curr_ini = ""
                                                        curr_fin = ""
                                                        curr_est = "No iniciado"
                                                        curr_etapa = "Subestructura"
                                                        curr_n_ini = ""
                                                        curr_n_fin = ""
                                                        
                                                    if not curr_nivel or curr_nivel not in niveles_opts: curr_nivel = "No iniciado"
                                                    if not curr_contra or curr_contra not in contractors_list: curr_contra = "Ninguno"
                                                    if not curr_est or curr_est not in estado_opts: curr_est = "No iniciado"
                                                    if not curr_etapa or curr_etapa not in etapas_opts: curr_etapa = "Subestructura"
                                                    
                                                    if not curr_n_ini or curr_n_ini not in niveles_rango_opts:
                                                        curr_n_ini = "Sótano 1" if curr_etapa == "Subestructura" and num_sotanos > 0 else "Piso 1"
                                                    if not curr_n_fin or curr_n_fin not in niveles_rango_opts:
                                                        curr_n_fin = f"Sótano {num_sotanos}" if curr_etapa == "Subestructura" and num_sotanos > 0 else f"Piso {num_pisos}"
                                                        
                                                    if curr_n_ini not in niveles_rango_opts: curr_n_ini = niveles_rango_opts[0]
                                                    if curr_n_fin not in niveles_rango_opts: curr_n_fin = niveles_rango_opts[-1]
                                                    
                                                    try:
                                                        ini_val = pd.to_datetime(curr_ini).date() if curr_ini and pd.notna(curr_ini) else pd.to_datetime(fecha_hoy).date()
                                                    except:
                                                        ini_val = pd.to_datetime(fecha_hoy).date()
                                                    try:
                                                        fin_val = pd.to_datetime(curr_fin).date() if curr_fin and pd.notna(curr_fin) else pd.to_datetime(fecha_hoy).date()
                                                    except:
                                                        fin_val = pd.to_datetime(fecha_hoy).date()
                                                        
                                                    c1, c1b, c1c, c1d, c2, c3, c4, c5, c6 = st.columns([1.5, 1.5, 1.5, 1.5, 1.5, 1.5, 1.2, 1.2, 1.6])
                                                    c1.markdown(f"<div style='padding-top:8px; font-weight:bold;'>{p_nom}</div>", unsafe_allow_html=True)
                                                    
                                                    sel_etapa = c1b.selectbox("Etapa", etapas_opts, index=etapas_opts.index(curr_etapa), key=f"sel_eta_{proy['Nombre']}_{p_nom}_{idx_part}", label_visibility="collapsed")
                                                    sel_n_ini = c1c.selectbox("Nivel Inicio", niveles_rango_opts, index=niveles_rango_opts.index(curr_n_ini), key=f"sel_n_ini_{proy['Nombre']}_{p_nom}_{idx_part}", label_visibility="collapsed")
                                                    sel_n_fin = c1d.selectbox("Nivel Fin", niveles_rango_opts, index=niveles_rango_opts.index(curr_n_fin), key=f"sel_n_fin_{proy['Nombre']}_{p_nom}_{idx_part}", label_visibility="collapsed")
                                                    
                                                    sel_nivel = c2.selectbox("Nivel Actual", niveles_opts, index=niveles_opts.index(curr_nivel), key=f"sel_niv_{proy['Nombre']}_{p_nom}_{idx_part}", label_visibility="collapsed")
                                                    sel_contra = c3.selectbox("Contratista", contractors_list, index=contractors_list.index(curr_contra), key=f"sel_con_{proy['Nombre']}_{p_nom}_{idx_part}", label_visibility="collapsed")
                                                    sel_ini = c4.date_input("Inicio", value=ini_val, key=f"sel_ini_{proy['Nombre']}_{p_nom}_{idx_part}", label_visibility="collapsed")
                                                    sel_fin = c5.date_input("Fin", value=fin_val, key=f"sel_fin_{proy['Nombre']}_{p_nom}_{idx_part}", label_visibility="collapsed")
                                                    sel_est = c6.selectbox("Estado", estado_opts, index=estado_opts.index(curr_est), key=f"sel_est_{proy['Nombre']}_{p_nom}_{idx_part}", label_visibility="collapsed")
                                                    
                                                    updated_data.append({
                                                        'Proyecto': proy['Nombre'],
                                                        'Nombre_Partida': p_nom,
                                                        'Nivel_Actual': sel_nivel,
                                                        'Contratista': sel_contra if sel_contra != "Ninguno" else "",
                                                        'Inicio': sel_ini.strftime('%Y-%m-%d'),
                                                        'Fin': sel_fin.strftime('%Y-%m-%d'),
                                                        'Estado': sel_est,
                                                        'Etapa': sel_etapa,
                                                        'Nivel_Inicio': sel_n_ini,
                                                        'Nivel_Fin': sel_n_fin
                                                    })
                                                    
                                                if st.form_submit_button("💾 GUARDAR TREN DE TRABAJO", use_container_width=True):
                                                    for item in updated_data:
                                                        proj = item['Proyecto']
                                                        part_name = item['Nombre_Partida']
                                                        idx_list = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proj) & (st.session_state.p_info['Nombre_Partida'] == part_name)].index
                                                        if not idx_list.empty:
                                                            idx = idx_list[0]
                                                            st.session_state.p_info.at[idx, 'Nivel_Actual'] = item['Nivel_Actual']
                                                            st.session_state.p_info.at[idx, 'Contratista'] = item['Contratista']
                                                            st.session_state.p_info.at[idx, 'Inicio'] = item['Inicio']
                                                            st.session_state.p_info.at[idx, 'Fin'] = item['Fin']
                                                            st.session_state.p_info.at[idx, 'Estado'] = item['Estado']
                                                            st.session_state.p_info.at[idx, 'Etapa'] = item['Etapa']
                                                            st.session_state.p_info.at[idx, 'Nivel_Inicio'] = item['Nivel_Inicio']
                                                            st.session_state.p_info.at[idx, 'Nivel_Fin'] = item['Nivel_Fin']
                                                        else:
                                                            new_row = pd.DataFrame([{
                                                                'Proyecto': proj,
                                                                'Nombre_Partida': part_name,
                                                                'Monto_Pactado': 0.0,
                                                                'Nivel_Actual': item['Nivel_Actual'],
                                                                'Contratista': item['Contratista'],
                                                                'Inicio': item['Inicio'],
                                                                'Fin': item['Fin'],
                                                                'Estado': item['Estado'],
                                                                'Etapa': item['Etapa'],
                                                                'Nivel_Inicio': item['Nivel_Inicio'],
                                                                'Nivel_Fin': item['Nivel_Fin']
                                                            }])
                                                            st.session_state.p_info = pd.concat([st.session_state.p_info, new_row[cols_p_info]], ignore_index=True)
                                                    st.rerun()
                                    
                                    with t_exp_entrega:
                                        st.info("Disponible Pronto")
                                else:
                                    for idx_part, part in parts.iterrows():
                                        p_nom = part['Nombre_Partida']
                                        st.markdown(f"""
                                            <div class='xian-section-header'>
                                                <span class='xian-section-header-icon'>■</span> PARTIDA // {p_nom.upper()}
                                            </div>
                                        """, unsafe_allow_html=True)
                                        
                                        df_pactado = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy['Nombre']) & (st.session_state.p_info['Nombre_Partida'] == p_nom)]
                                        monto_pactado = float(df_pactado['Monto_Pactado'].iloc[0]) if not df_pactado.empty else 0.0
                                        df_pagos_p = st.session_state.pagos[(st.session_state.pagos['Proyecto'] == proy['Nombre']) & (st.session_state.pagos['Partida'] == p_nom)]
                                        monto_pagado = pd.to_numeric(df_pagos_p['Monto'], errors='coerce').sum() if not df_pagos_p.empty else 0.0
                                        saldo = monto_pactado - monto_pagado
                                        col_s = '#10b981' if saldo >= 0 else '#ef4444'
                                        
                                        cf1, cf2, cf3, cf4, cf5 = st.columns([1.5, 1.8, 1.8, 1.5, 5.4])
                                        with cf1.popover("✏️ RENOMBRAR"):
                                            with st.form(f"ren_{idx_p}_{p_nom}_{idx_part}"):
                                                nuevo_nombre = st.text_input("Nuevo Nombre de Partida", value=p_nom)
                                                if st.form_submit_button("GUARDAR") and nuevo_nombre and nuevo_nombre != p_nom:
                                                    st.session_state.partidas.at[idx_part, 'Nombre_Partida'] = nuevo_nombre
                                                    p_info_matches = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy['Nombre']) & (st.session_state.p_info['Nombre_Partida'] == p_nom)]
                                                    for idx in p_info_matches.index:
                                                        st.session_state.p_info.at[idx, 'Nombre_Partida'] = nuevo_nombre
                                                    tasks_matches = st.session_state.tasks[(st.session_state.tasks['Proyecto'] == proy['Nombre']) & (st.session_state.tasks['Partida'] == p_nom)]
                                                    for idx in tasks_matches.index:
                                                        st.session_state.tasks.at[idx, 'Partida'] = nuevo_nombre
                                                    pagos_matches = st.session_state.pagos[(st.session_state.pagos['Proyecto'] == proy['Nombre']) & (st.session_state.pagos['Partida'] == p_nom)]
                                                    for idx in pagos_matches.index:
                                                        st.session_state.pagos.at[idx, 'Partida'] = nuevo_nombre
                                                    clear_base_dfs()
                                                    st.rerun()
                                        with cf2.popover("AJUSTAR PRESUPUESTO"):
                                            with st.form(f"pact_{idx_p}_{p_nom}_{idx_part}"):
                                                n_pac = st.number_input("Pactado (S/)", value=monto_pactado)
                                                if st.form_submit_button("FIJAR"):
                                                    if df_pactado.empty: st.session_state.p_info = pd.concat([st.session_state.p_info, pd.DataFrame([{'Proyecto': proy['Nombre'], 'Nombre_Partida': p_nom, 'Monto_Pactado': n_pac}])], ignore_index=True)
                                                    else: st.session_state.p_info.at[df_pactado.index[0], 'Monto_Pactado'] = n_pac
                                                    clear_base_dfs()
                                                    st.rerun()
            
                                        with cf3.popover("INGRESAR PAGO"):
                                            with st.form(f"pago_{idx_p}_{p_nom}_{idx_part}"):
                                                p_contra = st.selectbox("Contratista", st.session_state.maestro['ID_Contratista'].tolist(), key=f"sel_p_contra_{idx_p}_{p_nom}_{idx_part}")
                                                p_fecha = st.date_input("Fecha", value=pd.to_datetime(fecha_hoy).date(), key=f"sel_p_fecha_{idx_p}_{p_nom}_{idx_part}")
                                                p_monto = st.number_input("Monto (S/)", value=0.0, format="%.2f", key=f"sel_p_monto_{idx_p}_{p_nom}_{idx_part}")
                                                p_tipo = st.selectbox("Tipo de Pago", ["Adelanto", "Valorización", "Liquidación", "Adicional", "Otros"], key=f"sel_p_tipo_{idx_p}_{p_nom}_{idx_part}")
                                                p_notas = st.text_input("Notas", value="", key=f"sel_p_notas_{idx_p}_{p_nom}_{idx_part}")
                                                if st.form_submit_button("INGRESAR PAGO") and p_monto != 0.0:
                                                    n_pago_row = pd.DataFrame([{
                                                        'ID': datetime.now().timestamp(),
                                                        'Proyecto': proy['Nombre'],
                                                        'Partida': p_nom,
                                                        'Contratista': p_contra,
                                                        'Fecha': p_fecha.strftime("%d/%m/%Y"),
                                                        'Monto': p_monto,
                                                        'Tipo': p_tipo,
                                                        'Notas': p_notas
                                                    }])
                                                    st.session_state.pagos = pd.concat([st.session_state.pagos, n_pago_row[cols_pagos]], ignore_index=True)
                                                    clear_base_dfs()
                                                    st.rerun()
            
                                        with cf4.popover("🗑️ ELIMINAR"):
                                            st.warning(f"¿Seguro que desea eliminar la partida '{p_nom}'? Se eliminarán también sus tareas y presupuestos.")
                                            if st.button("CONFIRMAR ELIMINACIÓN", key=f"del_part_{idx_p}_{p_nom}_{idx_part}", type="primary", use_container_width=True):
                                                st.session_state.partidas = st.session_state.partidas.drop(idx_part)
                                                rem_count = len(st.session_state.partidas[(st.session_state.partidas['Proyecto'] == proy['Nombre']) & (st.session_state.partidas['Nombre_Partida'] == p_nom)])
                                                p_info_matches = st.session_state.p_info[(st.session_state.p_info['Proyecto'] == proy['Nombre']) & (st.session_state.p_info['Nombre_Partida'] == p_nom)]
                                                if len(p_info_matches) > rem_count:
                                                    drop_indices = p_info_matches.index[rem_count:]
                                                    st.session_state.p_info = st.session_state.p_info.drop(drop_indices)
                                                if rem_count == 0:
                                                    st.session_state.tasks = st.session_state.tasks[~((st.session_state.tasks['Proyecto'] == proy['Nombre']) & (st.session_state.tasks['Partida'] == p_nom))]
                                                clear_base_dfs()
                                                st.rerun()
        
                                        cls_saldo = 'xian-badge-saldo-ok' if saldo >= 0 else 'xian-badge-saldo-error'
                                        badge_html = f"""
                                        <div class='xian-badge-container'>
                                            <span class='xian-badge xian-badge-pactado'>PACTADO: S/ {monto_pactado:,.2f}</span>
                                            <span class='xian-badge xian-badge-pagado'>PAGADO: S/ {monto_pagado:,.2f}</span>
                                            <span class='xian-badge {cls_saldo}'>SALDO: S/ {saldo:,.2f}</span>
                                        </div>
                                        """
                                        cf5.markdown(badge_html, unsafe_allow_html=True)

                                        mask = (st.session_state.tasks['Proyecto'] == proy['Nombre']) & (st.session_state.tasks['Partida'] == p_nom)
                                        df_v = st.session_state.tasks[mask].drop(columns=['Proyecto', 'Partida']).reset_index(drop=True)
                                        for c in ['Tarea', 'Especialidad', 'Contratista', 'Estado', 'Causa_Cumplimiento', 'Dependencia', 'Notas']:
                                            if c in df_v.columns: df_v[c] = df_v[c].fillna("").astype(str).replace('nan', '')
                                        for c in ['Culpa_Contratista', 'Permitir_Cruce']:
                                            if c in df_v.columns: df_v[c] = df_v[c].apply(lambda x: str(x).strip().lower() in ['true', '1', 'verdadero', 't', 'y']).astype(bool)
                                        for c in ['Inicio', 'Fin']:
                                            if c in df_v.columns: df_v[c] = pd.to_datetime(df_v[c], errors='coerce').dt.date
                                        if 'Orden' in df_v.columns: df_v['Orden'] = pd.to_numeric(df_v['Orden'], errors='coerce').fillna(0).astype(int)
                                        df_v = df_v.sort_values(by='Inicio', na_position='last').reset_index(drop=True)

                                        base_key = f"base_df_tasks_{proy['Nombre']}_{p_nom}_{idx_part}"
                                        if base_key not in st.session_state:
                                            st.session_state[base_key] = df_v.copy()
                                        counter_key = f"c_{base_key}"
                                        w_key = f"t_{idx_p}_{p_nom}_{idx_part}_{st.session_state.get(counter_key, 0)}"

                                        upd_t = st.data_editor(
                                            st.session_state[base_key], use_container_width=True, hide_index=True, num_rows="dynamic", key=w_key,
                                            column_config={
                                                "Especialidad": st.column_config.SelectboxColumn(options=extraer_esp()),
                                                "Contratista": st.column_config.SelectboxColumn(options=st.session_state.maestro['ID_Contratista'].tolist()),
                                                "Estado": st.column_config.SelectboxColumn(options=status_tasks),
                                                "Inicio": st.column_config.DateColumn("Inicio", format="DD/MM/YYYY"),
                                                "Fin": st.column_config.DateColumn("Fin", format="DD/MM/YYYY")
                                            },
                                            on_change=on_table_edit,
                                            args=(
                                                base_key,
                                                w_key,
                                                'tasks',
                                                ['Proyecto', 'Partida'],
                                                [proy['Nombre'], p_nom],
                                                ['Inicio', 'Fin'],
                                                ['Tarea', 'Especialidad', 'Contratista', 'Estado', 'Causa_Cumplimiento', 'Dependencia', 'Notas'],
                                                ['Culpa_Contratista', 'Permitir_Cruce'],
                                                ['Orden'],
                                                None
                                            )
                                        )
                                        
                                        if st.button("ASISTENTE AI", key=f"ai_{proy['Nombre']}_{p_nom}_{idx_part}"):
                                            for _, row in upd_t.iterrows():
                                                if pd.notna(row['Especialidad']) and row['Especialidad'] != "Ninguna":
                                                    sug = sugerir_contratista(row['Especialidad'], proy.get('Ubicación', 'Lima'), row.get('Inicio'), row.get('Fin'))
                                                    if sug: st.success(f"{row.get('Tarea')}: {sug}")

    with t_con:
        df_r = st.session_state.maestro.copy()
        
        df_r['Estatus'] = df_r['ID_Contratista'].map(estado_dinamico).fillna("Disponible")
        df_r['Capacidad (7D)'] = df_r['ID_Contratista'].map(carga_dinamica).fillna("Libre")
        df_r['Rating_Interno'] = [calcular_rating(n) for n in df_r['ID_Contratista']]
        
        df_display = df_r[['ID_Contratista', 'Especialidad', 'Ubicación_Base', 'Estatus', 'Capacidad (7D)', 'Rating_Interno', 'Notas']]
        base_key = "base_df_maestro"
        if base_key not in st.session_state:
            st.session_state[base_key] = df_display.copy()
            
        editado = st.data_editor(
            st.session_state[base_key], 
            use_container_width=True, hide_index=True,
            column_config={"Estatus": st.column_config.Column(disabled=True), "Capacidad (7D)": st.column_config.Column(disabled=True), "Rating_Interno": st.column_config.Column(disabled=True)},
            key="ed_maestro"
        )
        
        if st.session_state[base_key].astype(str).to_dict() != editado.astype(str).to_dict():
            st.session_state.maestro = editado[['ID_Contratista', 'Especialidad', 'Ubicación_Base', 'Notas']].copy()
            st.session_state[base_key] = editado.copy()


        st.divider()
        st.markdown("#### BITÁCORA DE MAESTROS")
        if not df_r.empty:
            c_nota_sel = st.selectbox("Seleccionar Maestro:", df_r['ID_Contratista'].tolist())
            idx_c = st.session_state.maestro[st.session_state.maestro['ID_Contratista'] == c_nota_sel].index[0]
            notas_c = st.session_state.maestro.at[idx_c, 'Notas']
            
            with st.container(border=True):
                if pd.notna(notas_c) and str(notas_c).strip() != "":
                    html_notas = str(notas_c).replace("[+]", "<span style='color:#10b981;'>[+]</span>").replace("[-]", "<span style='color:#ef4444;'>[-]</span>")
                    st.markdown(f"<div style='font-family:monospace;'>{html_notas}</div>", unsafe_allow_html=True)
                else: st.info("Sin historial de reportes.")
                
            n_nota_c = st.text_input(f"Añadir reporte sobre {c_nota_sel}:")
            cb1, cb2, _ = st.columns([2, 2, 6])
            
            if cb1.button("REPORTE POSITIVO (+)") and n_nota_c:
                n_lin = f"**[{fecha_hoy.strftime('%d/%m %H:%M')}]** [+] {n_nota_c}"
                st.session_state.maestro.at[idx_c, 'Notas'] = f"{notas_c}\n<br>{n_lin}" if pd.notna(notas_c) and str(notas_c).strip() != "" else n_lin
                st.rerun()
            if cb2.button("REPORTE NEGATIVO (-)") and n_nota_c:
                n_lin = f"**[{fecha_hoy.strftime('%d/%m %H:%M')}]** [-] {n_nota_c}"
                st.session_state.maestro.at[idx_c, 'Notas'] = f"{notas_c}\n<br>{n_lin}" if pd.notna(notas_c) and str(notas_c).strip() != "" else n_lin
                st.rerun()

        st.divider()
        st.markdown("#### CRONOGRAMA DE ASIGNACIONES")
        df_tasks_all = st.session_state.tasks.copy().dropna(subset=['Contratista'])
        df_tasks_all = df_tasks_all[df_tasks_all['Contratista'].astype(str).str.strip() != '']
        df_tasks_all['Inicio'] = pd.to_datetime(df_tasks_all['Inicio'], errors='coerce')
        df_tasks_all['Fin_Visual'] = pd.to_datetime(df_tasks_all['Fin'], errors='coerce')
        df_tasks_all = df_tasks_all.dropna(subset=['Inicio', 'Fin_Visual'])
        if not df_tasks_all.empty:
            mask_1day_c = df_tasks_all['Inicio'].dt.date == df_tasks_all['Fin_Visual'].dt.date
            df_tasks_all.loc[mask_1day_c, 'Fin_Visual'] = df_tasks_all['Fin_Visual'] + timedelta(days=1)
            
            chart_c = alt.Chart(df_tasks_all).mark_bar(cornerRadius=2).encode(
                x=alt.X('Inicio:T', title='Fechas', axis=alt.Axis(gridColor='#1e293b', domainColor='#1e293b', format='%d %b')), 
                x2='Fin_Visual:T', 
                y=alt.Y('Contratista:N', sort='x', title='', axis=alt.Axis(labelColor='#cbd5e1')),
                color=alt.Color('Proyecto:N', scale=alt.Scale(scheme='set2')), 
                tooltip=['Contratista', 'Proyecto', 'Partida', 'Tarea', 'Inicio', 'Fin']
            )
            linea_hoy_c = alt.Chart(df_hoy).mark_rule(color='#ffffff', strokeWidth=2).encode(x='Hoy:T')
            gantt_final_c = alt.layer(chart_c, linea_hoy_c).properties(height=max(300, len(df_tasks_all['Contratista'].unique()) * 40)).configure_view(strokeOpacity=0).configure_axis(grid=True).interactive(bind_y=False)
            st.altair_chart(gantt_final_c, use_container_width=True)

    with t_fin:
        st.markdown("<h3 style='color:#00E5FF; font-family:monospace; margin-bottom:15px;'>CONSOLIDADO FINANCIERO Y LIQUIDACIÓN</h3>", unsafe_allow_html=True)
        
        # ── EXTRACCIÓN Y FILTRO POR AÑO ──
        df_p_all = st.session_state.proyectos[st.session_state.proyectos['Archivado'] != True].copy()
        df_p_all['_Inicio_dt'] = pd.to_datetime(df_p_all['Inicio'], errors='coerce')
        df_p_all['_Año'] = df_p_all['_Inicio_dt'].dt.year.fillna(0).astype(int)
        
        # Lista de años ordenados descendentemente (ej. 2026, 2025, 2024...)
        años_disp = sorted([int(a) for a in df_p_all['_Año'].unique() if a > 2000], reverse=True)
        
        col_filtro_año, col_info_año = st.columns([1.5, 4.5])
        with col_filtro_año:
            opciones_año = ["Todos"] + [str(a) for a in años_disp]
            año_sel = st.selectbox("📅 Filtrar por Año", opciones_año, key="fin_filtro_año")
            
        if año_sel != "Todos":
            df_proy_filtrado = df_p_all[df_p_all['_Año'] == int(año_sel)]
            with col_info_año:
                st.markdown(f"<div style='margin-top:28px; color:#94a3b8; font-size:13px;'>Mostrando datos financieros del año <b style='color:#00E5FF;'>{año_sel}</b> ({len(df_proy_filtrado)} proyectos)</div>", unsafe_allow_html=True)
        else:
            df_proy_filtrado = df_p_all
            with col_info_año:
                st.markdown(f"<div style='margin-top:28px; color:#94a3b8; font-size:13px;'>Mostrando datos de <b style='color:#00E5FF;'>Todos los años</b> ({len(df_proy_filtrado)} proyectos)</div>", unsafe_allow_html=True)

        nombres_proy_filtrados = df_proy_filtrado['Nombre'].tolist()
        
        # KPIs filtrados por año
        t_presupuesto = sum(safe_float(v) for v in df_proy_filtrado['Presupuesto'])
        t_gg_presup = sum(safe_float(v) for v in df_proy_filtrado['GG_Presupuesto'])
        
        # Gastos y Pagos correspondientes a los proyectos filtrados
        if not st.session_state.gg_log.empty:
            df_gg_filtrado = st.session_state.gg_log[st.session_state.gg_log['Proyecto'].isin(nombres_proy_filtrados)]
            t_gg_real = pd.to_numeric(df_gg_filtrado['Monto'], errors='coerce').sum()
        else:
            t_gg_real = 0.0
            
        if not st.session_state.pagos.empty:
            df_pagos_filtrado = st.session_state.pagos[st.session_state.pagos['Proyecto'].isin(nombres_proy_filtrados)]
            t_pagado_total = pd.to_numeric(df_pagos_filtrado['Monto'], errors='coerce').sum()
        else:
            t_pagado_total = 0.0
            
        t_saldo_ppto = t_presupuesto - t_pagado_total
        
        t_utilidad_proyectada = 0.0
        t_utilidad_recibida = 0.0
        datos_tabla = []
        df_ing = st.session_state.ingresos_socios.copy()
        
        for _, p in df_proy_filtrado.iterrows():
            if not st.session_state.gg_log.empty:
                gg_p_real = pd.to_numeric(st.session_state.gg_log[st.session_state.gg_log['Proyecto'] == p['Nombre']]['Monto'], errors='coerce').sum()
            else:
                gg_p_real = 0.0
            
            util_meta = safe_float(p.get('Utilidades', 0))
            if not df_ing.empty:
                s_ing_p = df_ing[(df_ing['Proyecto'] == p['Nombre']) & (df_ing['Categoria'] == 'Utilidades')]
                util_real_p = pd.to_numeric(s_ing_p['Monto'], errors='coerce').sum()
            else:
                util_real_p = 0.0
            
            t_utilidad_proyectada += util_meta
            t_utilidad_recibida += util_real_p
            
            if not st.session_state.pagos.empty:
                pagos_proy = st.session_state.pagos[st.session_state.pagos['Proyecto'] == p['Nombre']]
                total_pagado_proy = pd.to_numeric(pagos_proy['Monto'], errors='coerce').sum()
            else:
                total_pagado_proy = 0.0
                
            saldo_proy = safe_float(p.get('Presupuesto', 0)) - total_pagado_proy
            
            datos_tabla.append({
                "Proyecto": p['Nombre'],
                "Ppto Venta": f"S/ {safe_float(p.get('Presupuesto', 0)):,.2f}",
                "Ppto Pagado": f"S/ {total_pagado_proy:,.2f}",
                "Saldo Restante": f"S/ {saldo_proy:,.2f}",
                "Caja Chica Límite": f"S/ {safe_float(p.get('GG_Presupuesto', 0)):,.2f}",
                "Caja Chica Ejecutado": f"S/ {gg_p_real:,.2f}",
                "Utilidad Esperada": f"S/ {util_meta:,.2f}",
                "Utilidades": f"S/ {util_real_p:,.2f}"
            })
            
        mf1, mf2, mf3, mf4, mf5 = st.columns(5)
        mf1.markdown(f"<div class='metric-card'><h4>Total Ventas</h4><h2>S/ {t_presupuesto:,.2f}</h2></div>", unsafe_allow_html=True)
        mf2.markdown(f"<div class='metric-card'><h4>Pagos Totales (Salidas)</h4><h2 style='color:#00E5FF;'>S/ {t_pagado_total:,.2f}</h2></div>", unsafe_allow_html=True)
        mf3.markdown(f"<div class='metric-card'><h4>Caja Chica Ejecutada</h4><h2>S/ {t_gg_real:,.2f} <small style='font-size:12px; color:#64748b;'>/ Límite S/ {t_gg_presup:,.2f}</small></h2></div>", unsafe_allow_html=True)
        
        mf4.markdown(f"<div class='metric-card'><h4>Utilidades Proyectadas</h4><h2 style='color:#10b981;'>S/ {t_utilidad_proyectada:,.2f}</h2></div>", unsafe_allow_html=True)
        
        color_u_total = "#10b981" if t_utilidad_recibida >= 0 else "#ef4444"
        mf5.markdown(f"<div class='metric-card'><h4>Utilidades Recibidas</h4><h2 style='color: {color_u_total};'>S/ {t_utilidad_recibida:,.2f}</h2></div>", unsafe_allow_html=True)
        
        st.divider()
        if datos_tabla: 
            df_export = pd.DataFrame(datos_tabla)
            st.dataframe(df_export, use_container_width=True, hide_index=True)
            
            # Exportador de consolidado financiero
            csv_data = df_export.to_csv(index=False).encode('utf-8')
            st.download_button(
                label="📥 EXPORTAR CONSOLIDADO A CSV",
                data=csv_data,
                file_name=f"consolidado_financiero_{datetime.now().strftime('%Y%m%d')}.csv",
                mime="text/csv",
                use_container_width=True
            )

        # ── GRÁFICA COMPARATIVA POR AÑO ──
        st.divider()
        st.markdown("<h4 style='color:#00E5FF; font-family:monospace; margin-bottom:10px;'>📊 COMPARATIVA FINANCIERA POR AÑO</h4>", unsafe_allow_html=True)
        
        if len(años_disp) > 0:
            # Construir dataset comparativo cronológico
            datos_comparativa = []
            for año in sorted(años_disp):
                proys_año = df_p_all[df_p_all['_Año'] == año]
                nombres_año = proys_año['Nombre'].tolist()
                
                ventas_año = sum(safe_float(v) for v in proys_año['Presupuesto'])
                
                if not st.session_state.gg_log.empty:
                    gg_año = pd.to_numeric(st.session_state.gg_log[st.session_state.gg_log['Proyecto'].isin(nombres_año)]['Monto'], errors='coerce').sum()
                else:
                    gg_año = 0.0
                    
                util_proy_año = sum(safe_float(u) for u in proys_año['Utilidades'])
                
                if not df_ing.empty:
                    util_recibida_año = pd.to_numeric(df_ing[(df_ing['Proyecto'].isin(nombres_año)) & (df_ing['Categoria'] == 'Utilidades')]['Monto'], errors='coerce').sum()
                else:
                    util_recibida_año = 0.0
                    
                datos_comparativa.append({
                    "Año": str(año),
                    "Total Ventas": ventas_año,
                    "Gastos Generales": gg_año,
                    "Utilidades Proyectadas": util_proy_año,
                    "Utilidades Recibidas": util_recibida_año
                })
                
            df_comp = pd.DataFrame(datos_comparativa)
            
            c_filtro_m1, c_filtro_m2 = st.columns([3, 1])
            with c_filtro_m1:
                metricas_disponibles = ["Total Ventas", "Gastos Generales", "Utilidades Proyectadas", "Utilidades Recibidas"]
                metricas_sel = st.multiselect(
                    "Filtrar métricas para la comparativa:",
                    metricas_disponibles,
                    default=["Total Ventas", "Gastos Generales", "Utilidades Recibidas"],
                    key="fin_chart_metricas_sel"
                )
            
            if metricas_sel:
                df_melted = df_comp.melt(id_vars=['Año'], value_vars=metricas_sel, var_name='Métrica', value_name='Monto')
                
                color_map = alt.Scale(
                    domain=["Total Ventas", "Gastos Generales", "Utilidades Proyectadas", "Utilidades Recibidas"],
                    range=["#00E5FF", "#FFB300", "#34d399", "#10b981"]
                )
                
                chart_comp = alt.Chart(df_melted).mark_bar(cornerRadius=4, size=24).encode(
                    x=alt.X('Año:N', title='Año', axis=alt.Axis(labelColor='#cbd5e1', labelAngle=0, titleColor='#94a3b8')),
                    y=alt.Y('Monto:Q', title='Monto (S/)', axis=alt.Axis(labelColor='#94a3b8', titleColor='#94a3b8', format='~s')),
                    color=alt.Color('Métrica:N', scale=color_map, legend=alt.Legend(orient='top', title=None, labelColor='#cbd5e1')),
                    xOffset='Métrica:N',
                    tooltip=[
                        alt.Tooltip('Año:N', title='Año'),
                        alt.Tooltip('Métrica:N', title='Métrica'),
                        alt.Tooltip('Monto:Q', title='Monto (S/)', format=',.2f')
                    ]
                ).properties(height=360).configure_view(strokeOpacity=0).configure_axis(grid=True, gridColor='#1e293b')
                
                st.altair_chart(chart_comp, use_container_width=True)
            else:
                st.info("💡 Selecciona al menos una métrica para visualizar la gráfica comparativa.")
        
        st.divider()
        st.markdown("#### BALANCE DE INGRESOS")
        with st.form("f_ing", clear_on_submit=True):
            i1, i3, i4 = st.columns([4, 3, 3])
            i_proy = i1.selectbox("Obra", ["General"] + p_vis['Nombre'].tolist())
            i_cat = i3.selectbox("Tipo de Ingreso", ["Caja Chica", "Utilidades"])
            i_m = i4.number_input("Monto (S/)", value=0.0)
            if st.form_submit_button("INGRESAR CAPITAL A CAJA") and i_m != 0.0:
                n_ing = pd.DataFrame([{'ID': datetime.now().timestamp(), 'Fecha': fecha_hoy.strftime("%d/%m"), 'Proyecto': i_proy, 'Socio': 'Venti', 'Categoria': i_cat, 'Monto': i_m, 'Notas': ''}])
                st.session_state.ingresos_socios = pd.concat([st.session_state.ingresos_socios, n_ing], ignore_index=True); st.rerun()

        df_gg = st.session_state.gg_log.copy()
        caja_recibido = pd.to_numeric(df_ing[df_ing['Categoria'].isin(['Gastos Generales', 'Caja Chica'])]['Monto'], errors='coerce').sum() if not df_ing.empty else 0.0
        caja_gastado = pd.to_numeric(df_gg['Monto'], errors='coerce').sum() if not df_gg.empty else 0.0
        caja_balance = caja_recibido - caja_gastado
        utilidades_cobradas = pd.to_numeric(df_ing[df_ing['Categoria'] == 'Utilidades']['Monto'], errors='coerce').sum() if not df_ing.empty else 0.0

        bal_data = [{
            "Fondo Caja Chica Recibido": f"S/ {caja_recibido:,.2f}",
            "Caja Chica Gastado (Salidas)": f"S/ {caja_gastado:,.2f}",
            "Balance Caja Chica": f"S/ {caja_balance:,.2f}",
            "Utilidades Globales Cobradas": f"S/ {utilidades_cobradas:,.2f}"
        }]
            
        df_bal = pd.DataFrame(bal_data)
        st.dataframe(df_bal, use_container_width=True, hide_index=True)
        
        # Exportador de balance global
        csv_bal = df_bal.to_csv(index=False).encode('utf-8')
        st.download_button(
            label="📥 EXPORTAR BALANCE GLOBAL A CSV",
            data=csv_bal,
            file_name=f"balance_global_{datetime.now().strftime('%Y%m%d')}.csv",
            mime="text/csv",
            use_container_width=True
        )
        
        st.divider()
        st.markdown("#### AUDITORÍA: ENTRADAS VS SALIDAS GLOBALES")
        c_aud1, c_aud2 = st.columns(2)
        with c_aud1:
            st.markdown("##### Entradas (Ingresos Registrados)")
            if not df_ing.empty:
                df_ing_display = df_ing.drop(columns=['ID', 'Socio']) if 'Socio' in df_ing.columns else df_ing.drop(columns=['ID'])
                for col in ['Fecha', 'Proyecto', 'Categoria', 'Monto', 'Notas']:
                    if col not in df_ing_display.columns:
                        df_ing_display[col] = ""
                df_ing_display = df_ing_display[['Fecha', 'Proyecto', 'Categoria', 'Monto', 'Notas']]
                
                base_key = "base_df_ingresos_globales"
                if base_key not in st.session_state:
                    st.session_state[base_key] = df_ing_display.copy()
                    
                ed_ing = st.data_editor(
                    st.session_state[base_key],
                    use_container_width=True,
                    hide_index=True,
                    num_rows="dynamic",
                    key="ed_ingresos",
                    on_change=on_table_edit,
                    args=(
                        base_key,
                        "ed_ingresos",
                        'ingresos_socios',
                        ['Socio'],
                        ['Venti'],
                        None,  # date_cols
                        ['Fecha', 'Proyecto', 'Categoria', 'Notas'],  # text_cols
                        None,  # bool_cols
                        ['Monto'],  # num_cols
                        'ID'   # id_col
                    )
                )
            else: st.info("Sin registros.")
        with c_aud2:
            st.markdown("##### Salidas (Caja Chica Ejecutada)")
            if not df_gg.empty:
                df_gg_display = df_gg.drop(columns=['Responsable']) if 'Responsable' in df_gg.columns else df_gg
                st.dataframe(df_gg_display, use_container_width=True, hide_index=True)
            else: st.info("Sin registros.")

    with t_ope:
        co1, co2, co3 = st.columns([1.2, 1, 1])

        with co1:
            st.markdown("#### LOGÍSTICA DIARIA")
            with st.popover("NUEVO REQUERIMIENTO", use_container_width=True):
                with st.form("n_td", clear_on_submit=True):
                    t_p = st.selectbox("Asignar Obra", ["General"] + p_vis['Nombre'].tolist())
                    t_r = st.selectbox("Delegar a", lista_socios)
                    t_cat = st.selectbox("Clasificación", ["Compras", "Pagos", "Proveedores", "Coordinación"])
                    t_txt = st.text_input("Detalle de Actividad")
                    if st.form_submit_button("REGISTRAR ACTIVIDAD") and t_txt:
                        n_i = pd.DataFrame([{'ID': datetime.now().timestamp(), 'Proyecto': t_p, 'Categoria': t_cat, 'Tarea': t_txt, 'Responsable': t_r, 'Check': False}])
                        st.session_state.todo_list = pd.concat([st.session_state.todo_list, n_i], ignore_index=True); st.rerun()

            for i, row in st.session_state.todo_list.iterrows():
                if not row.get('Check', False):
                    ci, cb = st.columns([8.5, 1.5])
                    try:
                        idx_s = [s.lower() for s in lista_socios].index(str(row['Responsable']).lower())
                    except:
                        idx_s = 0
                    colors_tag = [
                        "background-color: #1e3a8a; color: #bfdbfe; padding: 2px 8px; border-radius: 2px; font-size: 10px; font-weight: bold; font-family: monospace;",
                        "background-color: #14532d; color: #bbf7d0; padding: 2px 8px; border-radius: 2px; font-size: 10px; font-weight: bold; font-family: monospace;",
                        "background-color: #581c87; color: #f3e8ff; padding: 2px 8px; border-radius: 2px; font-size: 10px; font-weight: bold; font-family: monospace;",
                        "background-color: #7c2d12; color: #ffedd5; padding: 2px 8px; border-radius: 2px; font-size: 10px; font-weight: bold; font-family: monospace;"
                    ]
                    style_tag = colors_tag[idx_s % len(colors_tag)]
                    ci.markdown(f"<div style='background:var(--card-bg); padding:10px; border:1px solid var(--border-color); border-radius:2px; margin-bottom:5px;'><b style='color:var(--text-muted);'>[{row['Categoria']}]</b> {row['Tarea']}<br><span style='{style_tag}'>USR: {str(row['Responsable']).upper()}</span> <span style='color:#00E5FF; font-family:monospace; margin-left:10px;'>PRY: {row['Proyecto']}</span></div>", unsafe_allow_html=True)
                    if cb.button("LISTO", key=f"c_{row['ID']}", use_container_width=True):
                        st.session_state.todo_list = st.session_state.todo_list.drop(i).reset_index(drop=True)
                        guardar_tabla(st.session_state.todo_list, 'logistica')
                        st.cache_data.clear()
                        st.rerun()

        with co2:
            st.markdown("#### EN EJECUCIÓN (HOY)")
            if not df_tareas_activas.empty:
                t_hoy = df_tareas_activas[(df_tareas_activas['Inicio_DT'] <= fecha_hoy) & (df_tareas_activas['Fin_DT'] >= fecha_hoy)]
                if not t_hoy.empty:
                    for p_n in t_hoy['Proyecto'].unique():
                        st.markdown(f"<div class='proj-label'>{p_n}</div>", unsafe_allow_html=True)
                        for _, t in t_hoy[t_hoy['Proyecto'] == p_n].iterrows():
                            c_n = str(t.get('Contratista', 'N/A')).strip()
                            if c_n.lower() in ['nan', 'none', '']: c_n = 'SIN ASIGNAR'
                            st.markdown(f"<div style='border-left: 2px solid #00E5FF; padding-left: 10px; margin-bottom:8px; background:var(--card-bg); padding:5px;'><b style='color:var(--text-muted);'>{t['Tarea']}</b><br><small style='color:var(--text-secondary); font-family:monospace;'>CONTRATISTA: {c_n}</small></div>", unsafe_allow_html=True)
                else: st.info("Sin registros de ejecución activa para hoy.")
            else: st.info("Cronograma vacío.")
            
        with co3:
            st.markdown("#### PRÓXIMOS ARRANQUES (7D)")
            if not df_tareas_activas.empty:
                lim_7d = fecha_hoy + timedelta(days=7)
                t_7 = df_tareas_activas[(df_tareas_activas['Inicio_DT'] > fecha_hoy) & (df_tareas_activas['Inicio_DT'] <= lim_7d)].sort_values('Inicio_DT')
                if not t_7.empty:
                    for p_n in t_7['Proyecto'].unique():
                        st.markdown(f"<div class='proj-label'>{p_n}</div>", unsafe_allow_html=True)
                        for _, t in t_7[t_7['Proyecto'] == p_n].iterrows():
                            c_n = str(t.get('Contratista', 'N/A')).strip()
                            f_s = t['Inicio_DT'].strftime('%d %b')
                            st.markdown(f"<div style='border-left: 2px solid #eab308; padding-left: 10px; margin-bottom:8px; background:var(--card-bg); padding:5px;'><b style='color:var(--text-muted);'>{t['Tarea']}</b><br><small style='color:var(--text-secondary); font-family:monospace;'>ARRANCA: {f_s} | CTR: {c_n}</small></div>", unsafe_allow_html=True)
                else: st.info("Sin arranques programados.")
            else: st.info("Cronograma vacío.")

    with t_cal:
        c_n1, c_n2, c_n3 = st.columns([1, 2, 1])
        if c_n1.button("MES PREVIO"): st.session_state.cal_month_offset -= 1; st.rerun()
        if c_n3.button("SIGUIENTE MES"): st.session_state.cal_month_offset += 1; st.rerun()
            
        t_d = fecha_hoy + pd.DateOffset(months=st.session_state.cal_month_offset)
        meses_es = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
        c_n2.markdown(f"<h3 style='text-align:center; color:#00E5FF; font-family:monospace;'>{meses_es[t_d.month].upper()} {t_d.year}</h3>", unsafe_allow_html=True)
        
        cal = calendar.monthcalendar(t_d.year, t_d.month)
        p_c = dict(zip(st.session_state.proyectos['Nombre'], st.session_state.proyectos['Color']))
        h_c = "<div class='calendar-container'>"
        for d in ["LUN", "MAR", "MIE", "JUE", "VIE", "SAB", "DOM"]: h_c += f"<div class='calendar-header-day'>{d}</div>"
            
        for w in cal:
            for d in w:
                if d == 0: h_c += "<div class='calendar-cell' style='background:transparent; border:none;'></div>"
                else:
                    c_dt = datetime(t_d.year, t_d.month, d).date()
                    cls = "today" if c_dt == fecha_hoy.date() else ""
                    h_c += f"<div class='calendar-cell {cls}'><div class='calendar-cell-date'>{d}</div>"
                    if not df_tareas_activas.empty:
                        t_t = df_tareas_activas[(df_tareas_activas['Inicio_Date'] <= c_dt) & (df_tareas_activas['Fin_Date'] >= c_dt)]
                        
                        # Group tasks by (Proyecto, Partida)
                        grupos_tareas = {}
                        for _, t in t_t.iterrows():
                            proj = str(t.get('Proyecto', 'N/A')).strip()
                            part = str(t.get('Partida', '')).strip()
                            if not part or part.lower() in ['nan', 'none']:
                                part = str(t.get('Tarea', 'General')).strip()
                            
                            key = (proj, part)
                            if key not in grupos_tareas:
                                grupos_tareas[key] = []
                            grupos_tareas[key].append(t)
                        
                        for (proj, part), list_t in grupos_tareas.items():
                            c = p_c.get(proj, '#3b82f6')
                            proj_esc = proj.replace('"', '&quot;').replace("'", "&#39;")
                            part_esc = part.replace('"', '&quot;').replace("'", "&#39;")
                            
                            # Construct tooltip content
                            tooltip_content = f"<b>PROYECTO:</b> {proj_esc}<br><b>PARTIDA:</b> {part_esc}"
                            
                            if len(list_t) == 1:
                                t = list_t[0]
                                t_esc = str(t.get('Tarea', 'N/A')).replace('"', '&quot;').replace("'", "&#39;")
                                c_n = str(t.get('Contratista', 'N/A')).strip()
                                if c_n.lower() in ['nan', 'none', '']: c_n = 'SIN ASIGNAR'
                                c_esc = c_n.replace('"', '&quot;').replace("'", "&#39;")
                                tooltip_content += f"<br><b>TAREA:</b> {t_esc}<br><b>CONTRATISTA:</b> {c_esc}"
                            else:
                                tooltip_content += f"<br><b>TAREAS ({len(list_t)}):</b>"
                                for item_t in list_t:
                                    t_esc = str(item_t.get('Tarea', 'N/A')).replace('"', '&quot;').replace("'", "&#39;")
                                    c_n = str(item_t.get('Contratista', 'N/A')).strip()
                                    if c_n.lower() in ['nan', 'none', '']: c_n = 'SIN ASIGNAR'
                                    c_esc = c_n.replace('"', '&quot;').replace("'", "&#39;")
                                    tooltip_content += f"<br>• {t_esc} ({c_esc})"
                                    
                            h_c += f"<div class='cal-task-item' style='background:{c}22; border-left: 2px solid {c}; color: {c};'><div style='overflow: hidden; text-overflow: ellipsis; white-space: nowrap;'>{part_esc[:15]}</div><span class='tooltip-text' style='border-left: 3px solid {c} !important;'>{tooltip_content}</span></div>"
                    h_c += "</div>"
        st.markdown(h_c + "</div>", unsafe_allow_html=True)
        
        # Leyenda de proyectos y colores
        p_activos = st.session_state.proyectos[st.session_state.proyectos['Archivado'] != True]
        if not p_activos.empty:
            legend_html = "<div class='calendar-legend'>"
            for _, p in p_activos.iterrows():
                p_name = p['Nombre']
                p_color = p.get('Color', '#3b82f6')
                legend_html += f"<div class='calendar-legend-item'><div class='calendar-legend-color' style='background-color:{p_color};'></div><span>{p_name}</span></div>"
            legend_html += "</div>"
            st.markdown(legend_html, unsafe_allow_html=True)

    with t_arc:
        df_a = st.session_state.proyectos[st.session_state.proyectos['Archivado'] == True]
        if not df_a.empty:
            for idx, r in df_a.iterrows():
                with st.container(border=True):
                    ca1, ca2, ca3 = st.columns([6, 2, 2])
                    ca1.markdown(f"**{r['Nombre']}**")
                    if ca2.button("RESTAURAR", key=f"rs_{idx}", use_container_width=True): st.session_state.proyectos.at[idx, 'Archivado'] = False; st.rerun()
                    if ca3.button("PURGAR", key=f"dl_{idx}", type="primary", use_container_width=True): st.session_state.proyectos = st.session_state.proyectos.drop(idx); st.rerun()
        else: st.info("Archivo histórico vacío.")
        

