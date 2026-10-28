"""
Módulo de Presupuesto y Cotización All-in-One de COTA APP.
Permite:
1. Construir y editar el presupuesto del proyecto 100% dentro de COTA (Pestaña 03. PRESUPUESTO):
   - Definir dinámicamente la cantidad de Partidas (1, 2, 3, 4...) y sus nombres.
   - Definir dinámicamente la cantidad de Tareas por cada Partida (1.01, 1.02, 1.03...) y editar
     constantemente su Descripción, Moneda, Unidad, Metrado y Valor Unitario.
   - Configurar Tipo de Cambio (TC), % de Gastos Generales, % de Gestión y Supervisión (Utilidades)
     y casilla de IGV (18% o 0%).
2. Sincronizar automáticamente todas las Partidas, Tareas y sus Presupuestos asignados hacia
   la pestaña 05. EJECUCIÓN & PARTIDAS y hacia la Ficha Financiera del Proyecto (Presupuesto Total,
   Utilidades y Caja Chica / Gastos Generales), preservando contratistas, estados y fechas de obra.
3. Exportar el presupuesto en cualquier momento a un archivo Excel (.xlsx) profesional con fórmulas
   automáticas para compartirlo con el cliente.
"""

from __future__ import annotations

import datetime
import io
import json
import math
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

from cota.db import COLS_P_INFO, COLS_PROY, COLS_TASKS, clear_base_dfs
from cota.utils import safe_float, safe_str, to_bool


# ==========================================
# 1. GESTIÓN EN VIVO DEL PRESUPUESTO DENTRO DE COTA (ALL-IN-ONE)
# ==========================================

ESTRUCTURA_COTIZACION_DEFAULT = [
    {"codigo": "1", "nombre": "Partida 1", "subitems": ["1.01", "1.02", "1.03"]},
    {"codigo": "2", "nombre": "Partida 2", "subitems": ["2.01", "2.02", "2.03"]},
    {"codigo": "3", "nombre": "Partida 3", "subitems": ["3.01", "3.02", "3.03"]},
]


def limpiar_nombre_partida(nombre_raw: str, idx_fallback: int = 1) -> str:
    """Extrae el nombre limpio de una partida sin el prefijo numérico '01. '."""
    s = safe_str(nombre_raw)
    clean = re.sub(r"^\d+[\.\-\s]+", "", s).strip()
    return clean if clean else f"Partida {idx_fallback}"


def formatear_nombre_partida(idx_1based: int, nombre_limpio: str) -> str:
    """Formatea el nombre canónico de la partida con prefijo de 2 dígitos: '01. Nombre'."""
    clean = limpiar_nombre_partida(nombre_limpio, idx_1based)
    return f"{idx_1based:02d}. {clean}"


def extraer_descripcion_tarea(tarea_raw: str, codigo_item: str = "") -> str:
    """Extrae la descripción limpia de una tarea quitando el prefijo '1.01 ' si lo tiene."""
    s = safe_str(tarea_raw)
    if not s:
        return ""
    if codigo_item and s.startswith(codigo_item):
        s = s[len(codigo_item):].lstrip(" -:")
    else:
        s = re.sub(r"^\d+\.\d+[\s\-:]*", "", s).strip()
    return s


def redondear_arriba_10(valor: float) -> float:
    """Redondea un monto hacia arriba al múltiplo de 10 más cercano (ej. 23.4 -> 30, 20.0 -> 20)."""
    v = safe_float(valor, 0.0)
    if v <= 0:
        return 0.0
    return float(math.ceil(round(v, 4) / 10.0) * 10)


def redondear_arriba_100(valor: float) -> float:
    """Redondea un monto hacia arriba al múltiplo de 100 más cercano (ej. 1435.5 -> 1500, 1400.0 -> 1400)."""
    v = safe_float(valor, 0.0)
    if v <= 0:
        return 0.0
    return float(math.ceil(round(v, 4) / 100.0) * 100)


def obtener_parametros_presupuesto_proyecto(proy_name: str) -> dict:
    """Obtiene TC, %/monto Gastos Generales, %/monto Gestión/Supervisión e IGV del proyecto."""
    params = {
        "tc": 3.42,
        "gg_pct": 2.0,
        "modo_gg": "pct",
        "monto_fijo_gg": 0.0,
        "utilidades_pct": 15.0,
        "modo_util": "pct",
        "monto_fijo_util": 0.0,
        "incluir_igv": False,
    }
    df_p = st.session_state.get("proyectos", pd.DataFrame())
    if df_p is None or df_p.empty or proy_name not in df_p["Nombre"].values:
        return params

    row = df_p[df_p["Nombre"] == proy_name].iloc[0]

    tc_val = safe_float(row.get("TC_Presupuesto"), 0.0)
    if tc_val > 0:
        params["tc"] = tc_val

    gg_raw = row.get("GG_Pct")
    if gg_raw is not None and str(gg_raw).strip() not in ["", "nan", "None"]:
        params["gg_pct"] = safe_float(gg_raw, 2.0)

    modo_gg_raw = safe_str(row.get("Modo_GG", "pct")).strip().lower()
    if modo_gg_raw in ["monto", "fijo", "exacto"]:
        params["modo_gg"] = "monto"
    else:
        params["modo_gg"] = "pct"
    params["monto_fijo_gg"] = safe_float(row.get("Monto_Fijo_GG", 0.0), 0.0)

    util_raw = row.get("Utilidades_Pct")
    if util_raw is not None and str(util_raw).strip() not in ["", "nan", "None"]:
        params["utilidades_pct"] = safe_float(util_raw, 15.0)

    modo_util_raw = safe_str(row.get("Modo_Utilidades", "pct")).strip().lower()
    if modo_util_raw in ["monto", "fijo", "exacto"]:
        params["modo_util"] = "monto"
    else:
        params["modo_util"] = "pct"
    params["monto_fijo_util"] = safe_float(row.get("Monto_Fijo_Utilidades", 0.0), 0.0)

    igv_raw = row.get("Incluir_IGV")
    if igv_raw is not None and str(igv_raw).strip() not in ["", "nan", "None"]:
        params["incluir_igv"] = to_bool(igv_raw)

    return params


def _sanitizar_dtypes_tasks_y_proyectos():
    """Garantiza que las columnas de texto sean object y las numéricas float en tasks, p_info y proyectos."""
    if "tasks" in st.session_state and isinstance(st.session_state.tasks, pd.DataFrame):
        for c_txt in ["Proyecto", "Partida", "Codigo_Item", "Tarea", "Especialidad", "Contratista", "Estado", "Supervisado", "Inicio", "Fin", "Causa_Cumplimiento", "Dependencia", "Notas", "Moneda", "Unidad"]:
            if c_txt in st.session_state.tasks.columns and st.session_state.tasks[c_txt].dtype != object:
                st.session_state.tasks[c_txt] = st.session_state.tasks[c_txt].fillna("").astype(str).replace({"nan": "", "None": ""}).astype(object)
        for c_num in ["Metrado", "Valor_Unitario", "Presupuesto"]:
            if c_num in st.session_state.tasks.columns:
                st.session_state.tasks[c_num] = pd.to_numeric(st.session_state.tasks[c_num], errors="coerce").fillna(0.0).astype(float)
        if "Orden" in st.session_state.tasks.columns:
            st.session_state.tasks["Orden"] = pd.to_numeric(st.session_state.tasks["Orden"], errors="coerce").fillna(0).astype(int)

    if "p_info" in st.session_state and isinstance(st.session_state.p_info, pd.DataFrame):
        if "Monto_Pactado" in st.session_state.p_info.columns:
            st.session_state.p_info["Monto_Pactado"] = pd.to_numeric(st.session_state.p_info["Monto_Pactado"], errors="coerce").fillna(0.0).astype(float)
        for c_pi in ["Proyecto", "Nombre_Partida", "Contratista", "Estado", "Etapa"]:
            if c_pi in st.session_state.p_info.columns and st.session_state.p_info[c_pi].dtype != object:
                st.session_state.p_info[c_pi] = st.session_state.p_info[c_pi].fillna("").astype(str).replace({"nan": "", "None": ""}).astype(object)

    if "proyectos" in st.session_state and isinstance(st.session_state.proyectos, pd.DataFrame):
        for c_num_p in ["Presupuesto", "Utilidades", "GG_Presupuesto", "Honorarios_Diseno", "Presupuesto_FFE", "TC_Presupuesto", "GG_Pct", "Utilidades_Pct", "Monto_Fijo_GG", "Monto_Fijo_Utilidades"]:
            if c_num_p in st.session_state.proyectos.columns:
                st.session_state.proyectos[c_num_p] = pd.to_numeric(st.session_state.proyectos[c_num_p], errors="coerce").fillna(0.0).astype(float)
        for c_txt_p in ["Link_Drive", "Pagos_Semanales_JSON", "Notas", "Modo_GG", "Modo_Utilidades"]:
            if c_txt_p in st.session_state.proyectos.columns and st.session_state.proyectos[c_txt_p].dtype != object:
                st.session_state.proyectos[c_txt_p] = st.session_state.proyectos[c_txt_p].fillna("").astype(str).replace({"nan": "", "None": ""}).astype(object)


# ==========================================
# GESTIÓN DE HISTORIAL / DESHACER (UNDO / REDO) DEL PRESUPUESTO
# ==========================================

def _crear_snapshot_presupuesto(proy_name: str) -> dict:
    """Crea una copia profunda del estado actual de tareas, partidas, p_info y parámetros del presupuesto."""
    _sanitizar_dtypes_tasks_y_proyectos()
    t_sub = st.session_state.tasks[st.session_state.tasks["Proyecto"] == proy_name].copy(deep=True) if "tasks" in st.session_state else pd.DataFrame()
    p_sub = st.session_state.partidas[st.session_state.partidas["Proyecto"] == proy_name].copy(deep=True) if "partidas" in st.session_state else pd.DataFrame()
    pi_sub = st.session_state.p_info[st.session_state.p_info["Proyecto"] == proy_name].copy(deep=True) if "p_info" in st.session_state else pd.DataFrame()

    proy_fields = {}
    if "proyectos" in st.session_state and not st.session_state.proyectos.empty:
        m = st.session_state.proyectos["Nombre"] == proy_name
        if m.any():
            r = st.session_state.proyectos[m].iloc[0]
            for f in ["Presupuesto", "Utilidades", "GG_Presupuesto", "TC_Presupuesto", "GG_Pct", "Utilidades_Pct", "Incluir_IGV", "Honorarios_Diseno"]:
                if f in r:
                    proy_fields[f] = r[f]

    return {
        "tasks": t_sub,
        "partidas": p_sub,
        "p_info": pi_sub,
        "proy_fields": proy_fields,
        "timestamp": datetime.datetime.now().isoformat(),
    }


def guardar_snapshot_presupuesto(proy_name: str):
    """Guarda el estado actual en la pila de deshacer (Undo history) antes de realizar una modificación."""
    try:
        snap = _crear_snapshot_presupuesto(proy_name)
        k_undo = f"ppto_undo_{proy_name}"
        k_redo = f"ppto_redo_{proy_name}"
        if k_undo not in st.session_state or not isinstance(st.session_state[k_undo], list):
            st.session_state[k_undo] = []
        # Evitar guardar snapshots idénticos consecutivos
        if st.session_state[k_undo]:
            ult = st.session_state[k_undo][-1]
            if ult["tasks"].equals(snap["tasks"]) and ult["partidas"].equals(snap["partidas"]) and ult["p_info"].equals(snap["p_info"]):
                return
        st.session_state[k_undo].append(snap)
        if len(st.session_state[k_undo]) > 25:
            st.session_state[k_undo].pop(0)
        # Cada nueva acción del usuario vacía la pila de rehacer (redo)
        st.session_state[k_redo] = []
    except Exception:
        pass


def puede_deshacer_presupuesto(proy_name: str) -> bool:
    """Indica si hay estados previos en el historial para deshacer."""
    k_undo = f"ppto_undo_{proy_name}"
    return bool(st.session_state.get(k_undo))


def puede_rehacer_presupuesto(proy_name: str) -> bool:
    """Indica si hay estados posteriores en la pila de rehacer."""
    k_redo = f"ppto_redo_{proy_name}"
    return bool(st.session_state.get(k_redo))


def _restaurar_snapshot_presupuesto(proy_name: str, snap: dict):
    """Restaura en session_state los datos guardados en un snapshot."""
    # 1. Restaurar tareas
    if "tasks" in st.session_state and not snap["tasks"].empty:
        st.session_state.tasks = st.session_state.tasks[st.session_state.tasks["Proyecto"] != proy_name].reset_index(drop=True)
        st.session_state.tasks = pd.concat([st.session_state.tasks, snap["tasks"]], ignore_index=True)
    elif "tasks" in st.session_state:
        st.session_state.tasks = st.session_state.tasks[st.session_state.tasks["Proyecto"] != proy_name].reset_index(drop=True)

    # 2. Restaurar partidas
    if "partidas" in st.session_state and not snap["partidas"].empty:
        st.session_state.partidas = st.session_state.partidas[st.session_state.partidas["Proyecto"] != proy_name].reset_index(drop=True)
        st.session_state.partidas = pd.concat([st.session_state.partidas, snap["partidas"]], ignore_index=True)

    # 3. Restaurar p_info
    if "p_info" in st.session_state and not snap["p_info"].empty:
        st.session_state.p_info = st.session_state.p_info[st.session_state.p_info["Proyecto"] != proy_name].reset_index(drop=True)
        st.session_state.p_info = pd.concat([st.session_state.p_info, snap["p_info"]], ignore_index=True)

    # 4. Restaurar campos del proyecto
    if "proyectos" in st.session_state and not st.session_state.proyectos.empty:
        m = st.session_state.proyectos["Nombre"] == proy_name
        if m.any():
            idx_p = st.session_state.proyectos[m].index[0]
            for f, val in snap.get("proy_fields", {}).items():
                if f in st.session_state.proyectos.columns:
                    st.session_state.proyectos.at[idx_p, f] = val

    # 5. Invalidar estados de data_editor para forzar recarga limpia en la UI
    clear_base_dfs()
    for k in list(st.session_state.keys()):
        if str(k).startswith(f"base_df_ppto_{proy_name}") or str(k).startswith(f"c_base_df_ppto_{proy_name}") or str(k).startswith(f"ed_ppto_{proy_name}"):
            del st.session_state[k]
        elif str(k).startswith(f"cfg_n_partidas_{proy_name}") or str(k).startswith(f"ppto_ntasks_{proy_name}") or str(k).startswith(f"ppto_pnom_{proy_name}"):
            del st.session_state[k]

    recalcular_totales_presupuesto_proyecto(proy_name)


def deshacer_presupuesto(proy_name: str) -> bool:
    """Deshace el último cambio en el presupuesto volviendo al estado inmediatamente anterior."""
    k_undo = f"ppto_undo_{proy_name}"
    k_redo = f"ppto_redo_{proy_name}"
    if not puede_deshacer_presupuesto(proy_name):
        return False
    try:
        # Guardar estado actual en la pila de rehacer
        actual = _crear_snapshot_presupuesto(proy_name)
        if k_redo not in st.session_state or not isinstance(st.session_state[k_redo], list):
            st.session_state[k_redo] = []
        st.session_state[k_redo].append(actual)

        # Recuperar estado anterior
        prev_snap = st.session_state[k_undo].pop()
        _restaurar_snapshot_presupuesto(proy_name, prev_snap)
        return True
    except Exception:
        return False


def rehacer_presupuesto(proy_name: str) -> bool:
    """Rehace el cambio que fue deshecho previamente."""
    k_undo = f"ppto_undo_{proy_name}"
    k_redo = f"ppto_redo_{proy_name}"
    if not puede_rehacer_presupuesto(proy_name):
        return False
    try:
        # Guardar estado actual en la pila de deshacer
        actual = _crear_snapshot_presupuesto(proy_name)
        st.session_state[k_undo].append(actual)

        # Recuperar estado posterior
        next_snap = st.session_state[k_redo].pop()
        _restaurar_snapshot_presupuesto(proy_name, next_snap)
        return True
    except Exception:
        return False


def guardar_parametros_presupuesto_proyecto(
    proy_name: str,
    tc: float = 3.42,
    gg_pct: float = 2.0,
    utilidades_pct: float = 15.0,
    incluir_igv: bool = False,
    modo_gg: str = "pct",
    monto_fijo_gg: float = 0.0,
    modo_util: str = "pct",
    monto_fijo_util: float = 0.0,
):
    """Guarda los parámetros globales del presupuesto en la fila del proyecto y recalcula montos."""
    if "proyectos" not in st.session_state or st.session_state.proyectos.empty:
        return
    guardar_snapshot_presupuesto(proy_name)
    for col in COLS_PROY:
        if col not in st.session_state.proyectos.columns:
            st.session_state.proyectos[col] = ""
    _sanitizar_dtypes_tasks_y_proyectos()

    mask = st.session_state.proyectos["Nombre"] == proy_name
    if mask.any():
        idx = st.session_state.proyectos[mask].index[0]
        st.session_state.proyectos.at[idx, "TC_Presupuesto"] = float(tc)
        st.session_state.proyectos.at[idx, "GG_Pct"] = float(gg_pct)
        st.session_state.proyectos.at[idx, "Utilidades_Pct"] = float(utilidades_pct)
        st.session_state.proyectos.at[idx, "Modo_GG"] = str(modo_gg)
        st.session_state.proyectos.at[idx, "Monto_Fijo_GG"] = float(monto_fijo_gg)
        st.session_state.proyectos.at[idx, "Modo_Utilidades"] = str(modo_util)
        st.session_state.proyectos.at[idx, "Monto_Fijo_Utilidades"] = float(monto_fijo_util)
        if st.session_state.proyectos["Incluir_IGV"].dtype != bool:
            st.session_state.proyectos["Incluir_IGV"] = st.session_state.proyectos["Incluir_IGV"].apply(to_bool).astype(bool)
        st.session_state.proyectos.at[idx, "Incluir_IGV"] = bool(incluir_igv)

    recalcular_totales_presupuesto_proyecto(proy_name)


def asegurar_estructura_presupuesto_proyecto(
    proy_name: str,
    n_partidas_default: int = 3,
    n_tareas_default: int = 3,
):
    """
    Garantiza que el proyecto tenga una estructura de Partidas y Tareas lista para editarse
    en '03. PRESUPUESTO' y sincronizada con '05. EJECUCIÓN & PARTIDAS'.
    """
    if "partidas" not in st.session_state:
        st.session_state.partidas = pd.DataFrame(columns=["Proyecto", "Nombre_Partida"])
    if "p_info" not in st.session_state:
        st.session_state.p_info = pd.DataFrame(columns=COLS_P_INFO)
    if "tasks" not in st.session_state:
        st.session_state.tasks = pd.DataFrame(columns=COLS_TASKS)

    for col in COLS_P_INFO:
        if col not in st.session_state.p_info.columns:
            st.session_state.p_info[col] = ""
    for col in COLS_TASKS:
        if col not in st.session_state.tasks.columns:
            st.session_state.tasks[col] = ""
    _sanitizar_dtypes_tasks_y_proyectos()

    parts_proy = st.session_state.partidas[st.session_state.partidas["Proyecto"] == proy_name]

    # Si el proyecto aún no tiene ninguna partida, inicializar con n_partidas_default (ej. Partida 1, 2, 3)
    if parts_proy.empty:
        ajustar_cantidad_partidas_proyecto(proy_name, n_partidas_default, n_tareas_default)
        return

    # Si ya tiene partidas, asegurar que cada partida tenga al menos 1 tarea en tasks
    nuevas_tasks = []
    for idx_p, (_, p_row) in enumerate(parts_proy.reset_index(drop=True).iterrows(), start=1):
        p_nom = safe_str(p_row["Nombre_Partida"])
        mask_t = (st.session_state.tasks["Proyecto"] == proy_name) & (st.session_state.tasks["Partida"] == p_nom)
        df_t_p = st.session_state.tasks[mask_t]

        if df_t_p.empty:
            # Verificar si tenía un Monto_Pactado previo en p_info
            mask_pi = (st.session_state.p_info["Proyecto"] == proy_name) & (st.session_state.p_info["Nombre_Partida"] == p_nom)
            monto_prev = safe_float(st.session_state.p_info.loc[mask_pi, "Monto_Pactado"].iloc[0], 0.0) if mask_pi.any() else 0.0

            for idx_t in range(1, n_tareas_default + 1):
                cod_t = f"{idx_p}.{idx_t:02d}"
                vu_init = monto_prev if (idx_t == 1 and monto_prev > 0) else 0.0
                desc_init = limpiar_nombre_partida(p_nom, idx_p) if (idx_t == 1 and monto_prev > 0) else ""
                tarea_txt = f"{cod_t} {desc_init}".strip() if desc_init else cod_t
                nuevas_tasks.append({
                    "Proyecto": proy_name,
                    "Partida": p_nom,
                    "Orden": idx_t,
                    "Codigo_Item": cod_t,
                    "Tarea": tarea_txt,
                    "Especialidad": "Ninguna",
                    "Contratista": "",
                    "Estado": "🟡 Planificado",
                    "Supervisado": "No",
                    "Inicio": datetime.date.today().strftime("%Y-%m-%d"),
                    "Fin": "",
                    "Causa_Cumplimiento": "",
                    "Culpa_Contratista": False,
                    "Permitir_Cruce": False,
                    "Dependencia": "",
                    "Notas": "",
                    "Moneda": "Soles",
                    "Unidad": "glb",
                    "Metrado": 1.0,
                    "Valor_Unitario": vu_init,
                    "Presupuesto": vu_init,
                })
        else:
            # Asegurar numeración Codigo_Item en tareas existentes
            for idx_t, r_idx in enumerate(df_t_p.index, start=1):
                cod_t = f"{idx_p}.{idx_t:02d}"
                st.session_state.tasks.at[r_idx, "Orden"] = idx_t
                st.session_state.tasks.at[r_idx, "Codigo_Item"] = cod_t
                if not safe_str(st.session_state.tasks.at[r_idx, "Moneda"]):
                    st.session_state.tasks.at[r_idx, "Moneda"] = "Soles"
                if not safe_str(st.session_state.tasks.at[r_idx, "Unidad"]):
                    st.session_state.tasks.at[r_idx, "Unidad"] = "glb"
                if safe_float(st.session_state.tasks.at[r_idx, "Metrado"], 0.0) <= 0:
                    st.session_state.tasks.at[r_idx, "Metrado"] = 1.0

    if nuevas_tasks:
        st.session_state.tasks = pd.concat(
            [st.session_state.tasks, pd.DataFrame(nuevas_tasks)[COLS_TASKS]],
            ignore_index=True,
        )
        clear_base_dfs()


def recalcular_totales_presupuesto_proyecto(proy_name: str, sincronizar_pactado: bool = True) -> dict:
    """
    Recalcula en cascada:
    1. El Valor Parcial (Presupuesto en S/) de cada tarea según su Metrado, Valor Unitario, Moneda y TC.
    2. El Sub Total (Monto_Pactado) de cada Partida en st.session_state.p_info.
    3. El Sub Total (Casco), Gastos Generales, Gestión/Supervisión (Utilidades), IGV y TOTAL del proyecto
       en st.session_state.proyectos.
    """
    _sanitizar_dtypes_tasks_y_proyectos()
    params = obtener_parametros_presupuesto_proyecto(proy_name)
    tc = params["tc"] if params["tc"] > 0 else 3.42
    gg_pct = params["gg_pct"]
    util_pct = params["utilidades_pct"]
    incluir_igv = params["incluir_igv"]

    parts_proy = st.session_state.partidas[st.session_state.partidas["Proyecto"] == proy_name] if "partidas" in st.session_state else pd.DataFrame()

    subtotal_casco = 0.0
    subtotales_por_partida = {}

    for _, p_row in parts_proy.iterrows():
        p_nom = safe_str(p_row["Nombre_Partida"])
        mask_t = (st.session_state.tasks["Proyecto"] == proy_name) & (st.session_state.tasks["Partida"] == p_nom)
        df_t_p = st.session_state.tasks[mask_t]

        sub_p = 0.0
        contratistas_p = []
        for r_idx, t_row in df_t_p.iterrows():
            mon = safe_str(t_row.get("Moneda", "Soles")) or "Soles"
            met = safe_float(t_row.get("Metrado", 1.0), 1.0)
            if met <= 0:
                met = 1.0
            vu = safe_float(t_row.get("Valor_Unitario", 0.0), 0.0)
            es_usd = any(d in mon.lower() for d in ["dolar", "dólar", "usd", "$"])
            factor_tc = tc if es_usd else 1.0
            parcial_soles = round(met * vu * factor_tc, 2)
            st.session_state.tasks.at[r_idx, "Presupuesto"] = float(parcial_soles)
            sub_p += parcial_soles
            c_t = safe_str(t_row.get("Contratista", "")).strip()
            if c_t and c_t not in contratistas_p:
                contratistas_p.append(c_t)

        sub_p = round(sub_p, 2)
        subtotales_por_partida[p_nom] = sub_p
        subtotal_casco += sub_p

        # Actualizar Monto_Pactado y Contratista en p_info
        mask_pi = (st.session_state.p_info["Proyecto"] == proy_name) & (st.session_state.p_info["Nombre_Partida"] == p_nom)
        contra_str = ", ".join(contratistas_p)
        if mask_pi.any():
            if sincronizar_pactado:
                st.session_state.p_info.loc[mask_pi, "Monto_Pactado"] = float(sub_p)
            if contra_str:
                st.session_state.p_info.loc[mask_pi, "Contratista"] = contra_str
        else:
            nueva_pi = pd.DataFrame([{
                "Proyecto": proy_name,
                "Nombre_Partida": p_nom,
                "Monto_Pactado": float(sub_p),
                "Nivel_Actual": "",
                "Contratista": contra_str,
                "Inicio": "",
                "Fin": "",
                "Estado": "🟡 Planificado",
                "Etapa": "Construcción",
                "Nivel_Inicio": "",
                "Nivel_Fin": "",
            }])
            st.session_state.p_info = pd.concat([st.session_state.p_info, nueva_pi[COLS_P_INFO]], ignore_index=True)

    subtotal_casco = round(subtotal_casco, 2)

    modo_gg = params.get("modo_gg", "pct")
    if modo_gg == "monto":
        gg_monto = redondear_arriba_10(params.get("monto_fijo_gg", 0.0))
        gg_pct = round((gg_monto / subtotal_casco * 100.0) if subtotal_casco > 0 else 0.0, 2)
    else:
        calc_gg = subtotal_casco * (gg_pct / 100.0)
        gg_monto = redondear_arriba_10(calc_gg)
        # Recalcular % efectivo tras redondeo hacia arriba al 10
        gg_pct = round((gg_monto / subtotal_casco * 100.0) if subtotal_casco > 0 else gg_pct, 2)

    modo_util = params.get("modo_util", "pct")
    if modo_util == "monto":
        util_monto = redondear_arriba_100(params.get("monto_fijo_util", 0.0))
        util_pct = round((util_monto / subtotal_casco * 100.0) if subtotal_casco > 0 else 0.0, 2)
    else:
        calc_util = subtotal_casco * (util_pct / 100.0)
        util_monto = redondear_arriba_100(calc_util)
        # Recalcular % efectivo tras redondeo hacia arriba al 100
        util_pct = round((util_monto / subtotal_casco * 100.0) if subtotal_casco > 0 else util_pct, 2)

    base_imp = subtotal_casco + gg_monto + util_monto
    igv_monto = round(base_imp * 0.18, 2) if incluir_igv else 0.0
    total_soles = round(base_imp + igv_monto, 2)
    total_usd = round(total_soles / tc, 2) if tc > 0 else 0.0

    # Actualizar fila del proyecto en st.session_state.proyectos
    if "proyectos" in st.session_state and not st.session_state.proyectos.empty:
        mask_p = st.session_state.proyectos["Nombre"] == proy_name
        if mask_p.any():
            idx_p = st.session_state.proyectos[mask_p].index[0]
            st.session_state.proyectos.at[idx_p, "Presupuesto"] = float(total_soles)
            st.session_state.proyectos.at[idx_p, "GG_Presupuesto"] = float(gg_monto)
            st.session_state.proyectos.at[idx_p, "Utilidades"] = float(util_monto)
            st.session_state.proyectos.at[idx_p, "GG_Pct"] = float(gg_pct)
            st.session_state.proyectos.at[idx_p, "Utilidades_Pct"] = float(util_pct)

    return {
        "tc": tc,
        "subtotal_casco": subtotal_casco,
        "gg_pct": gg_pct,
        "gg_monto": gg_monto,
        "modo_gg": modo_gg,
        "monto_fijo_gg": params.get("monto_fijo_gg", 0.0),
        "utilidades_pct": util_pct,
        "utilidades_monto": util_monto,
        "modo_util": modo_util,
        "monto_fijo_util": params.get("monto_fijo_util", 0.0),
        "incluir_igv": incluir_igv,
        "igv_pct": 18.0 if incluir_igv else 0.0,
        "igv_monto": igv_monto,
        "presupuesto_total": total_soles,
        "presupuesto_dolares": total_usd,
        "subtotales_por_partida": subtotales_por_partida,
    }


def ajustar_cantidad_partidas_proyecto(proy_name: str, nueva_cantidad: int, n_tareas_default: int = 3):
    """
    Ajusta dinámicamente el número de Partidas del proyecto (ej. de 3 a 4 partidas -> crea Partida 4
    con sus tareas iniciales; si disminuye, conserva las primeras `nueva_cantidad` partidas).
    """
    guardar_snapshot_presupuesto(proy_name)
    nueva_cantidad = max(1, int(nueva_cantidad))
    parts_proy = st.session_state.partidas[st.session_state.partidas["Proyecto"] == proy_name].reset_index(drop=True)
    actual_cant = len(parts_proy)

    if nueva_cantidad > actual_cant:
        nuevas_p = []
        nuevas_pi = []
        nuevas_t = []
        for idx_p in range(actual_cant + 1, nueva_cantidad + 1):
            p_nom = formatear_nombre_partida(idx_p, f"Partida {idx_p}")
            nuevas_p.append({"Proyecto": proy_name, "Nombre_Partida": p_nom})
            nuevas_pi.append({
                "Proyecto": proy_name,
                "Nombre_Partida": p_nom,
                "Monto_Pactado": 0.0,
                "Nivel_Actual": "",
                "Contratista": "",
                "Inicio": "",
                "Fin": "",
                "Estado": "🟡 Planificado",
                "Etapa": "Construcción",
                "Nivel_Inicio": "",
                "Nivel_Fin": "",
            })
            for idx_t in range(1, n_tareas_default + 1):
                cod_t = f"{idx_p}.{idx_t:02d}"
                nuevas_t.append({
                    "Proyecto": proy_name,
                    "Partida": p_nom,
                    "Orden": idx_t,
                    "Codigo_Item": cod_t,
                    "Tarea": cod_t,
                    "Especialidad": "Ninguna",
                    "Contratista": "",
                    "Estado": "🟡 Planificado",
                    "Supervisado": "No",
                    "Inicio": datetime.date.today().strftime("%Y-%m-%d"),
                    "Fin": "",
                    "Causa_Cumplimiento": "",
                    "Culpa_Contratista": False,
                    "Permitir_Cruce": False,
                    "Dependencia": "",
                    "Notas": "",
                    "Moneda": "Soles",
                    "Unidad": "glb",
                    "Metrado": 1.0,
                    "Valor_Unitario": 0.0,
                    "Presupuesto": 0.0,
                })
        st.session_state.partidas = pd.concat([st.session_state.partidas, pd.DataFrame(nuevas_p)], ignore_index=True)
        st.session_state.p_info = pd.concat([st.session_state.p_info, pd.DataFrame(nuevas_pi)[COLS_P_INFO]], ignore_index=True)
        st.session_state.tasks = pd.concat([st.session_state.tasks, pd.DataFrame(nuevas_t)[COLS_TASKS]], ignore_index=True)

    elif nueva_cantidad < actual_cant:
        partidas_a_eliminar = parts_proy.iloc[nueva_cantidad:]["Nombre_Partida"].tolist()
        st.session_state.partidas = st.session_state.partidas[
            ~((st.session_state.partidas["Proyecto"] == proy_name) & (st.session_state.partidas["Nombre_Partida"].isin(partidas_a_eliminar)))
        ].reset_index(drop=True)
        st.session_state.p_info = st.session_state.p_info[
            ~((st.session_state.p_info["Proyecto"] == proy_name) & (st.session_state.p_info["Nombre_Partida"].isin(partidas_a_eliminar)))
        ].reset_index(drop=True)
        st.session_state.tasks = st.session_state.tasks[
            ~((st.session_state.tasks["Proyecto"] == proy_name) & (st.session_state.tasks["Partida"].isin(partidas_a_eliminar)))
        ].reset_index(drop=True)

    clear_base_dfs()
    recalcular_totales_presupuesto_proyecto(proy_name)


def renombrar_partida_presupuesto(
    proy_name: str,
    idx_partida_1based,
    nombre_actual_completo: str,
    nuevo_nombre_limpio: str = None,
):
    """
    Renombra una Partida en el presupuesto y propaga el nuevo nombre automáticamente
    a partidas, p_info, tasks (05. Ejecución & Partidas) y pagos.
    Soporta llamada con 4 argumentos (proy_name, idx_p, nombre_actual, nuevo_limpio)
    o 3 argumentos (proy_name, nombre_actual, nuevo_nombre).
    """
    if nuevo_nombre_limpio is None:
        # Llamada con 3 argumentos: (proy_name, nombre_actual_completo, nuevo_nombre)
        nuevo_nombre_limpio = str(nombre_actual_completo)
        nombre_actual_completo = str(idx_partida_1based)
        m_idx = re.match(r"^\s*(\d+)", nombre_actual_completo)
        idx_partida_1based = int(m_idx.group(1)) if m_idx else 1

    nuevo_completo = formatear_nombre_partida(int(idx_partida_1based), nuevo_nombre_limpio)
    if nuevo_completo == nombre_actual_completo:
        return nuevo_completo

    guardar_snapshot_presupuesto(proy_name)

    st.session_state.partidas.loc[
        (st.session_state.partidas["Proyecto"] == proy_name) & (st.session_state.partidas["Nombre_Partida"] == nombre_actual_completo),
        "Nombre_Partida",
    ] = nuevo_completo

    st.session_state.p_info.loc[
        (st.session_state.p_info["Proyecto"] == proy_name) & (st.session_state.p_info["Nombre_Partida"] == nombre_actual_completo),
        "Nombre_Partida",
    ] = nuevo_completo

    st.session_state.tasks.loc[
        (st.session_state.tasks["Proyecto"] == proy_name) & (st.session_state.tasks["Partida"] == nombre_actual_completo),
        "Partida",
    ] = nuevo_completo

    if "pagos" in st.session_state and not st.session_state.pagos.empty:
        st.session_state.pagos.loc[
            (st.session_state.pagos["Proyecto"] == proy_name) & (st.session_state.pagos["Partida"] == nombre_actual_completo),
            "Partida",
        ] = nuevo_completo

    clear_base_dfs()
    return nuevo_completo


def _reemplazar_tareas_partida_in_place(proy_name: str, nombre_partida_completo: str, df_nuevas_partida: pd.DataFrame):
    """Reemplaza las tareas de una partida conservando la posición exacta de la partida dentro de st.session_state.tasks."""
    master_df = st.session_state.tasks.reset_index(drop=True)
    for c in master_df.columns:
        if c not in df_nuevas_partida.columns:
            df_nuevas_partida[c] = "" if master_df[c].dtype == object else 0
    df_nuevas_partida = df_nuevas_partida[master_df.columns].copy()

    mask_t = (master_df["Proyecto"] == proy_name) & (master_df["Partida"] == nombre_partida_completo)
    if mask_t.any():
        first_pos = int(master_df.index[mask_t][0])
        df_before = master_df.iloc[:first_pos][~mask_t.iloc[:first_pos]]
        df_after = master_df.iloc[first_pos:][~mask_t.iloc[first_pos:]]
        st.session_state.tasks = pd.concat([df_before, df_nuevas_partida, df_after], ignore_index=True)
    else:
        st.session_state.tasks = pd.concat([master_df, df_nuevas_partida], ignore_index=True)


def ajustar_cantidad_tareas_partida(
    proy_name: str,
    idx_partida_1based: int,
    nombre_partida_completo: str,
    nueva_cant_tareas: int,
):
    """
    Ajusta dinámicamente la cantidad de tareas (1.01, 1.02, 1.03...) dentro de una Partida,
    preservando las tareas existentes y su posición sin alterar el orden de las partidas.
    """
    guardar_snapshot_presupuesto(proy_name)
    _sanitizar_dtypes_tasks_y_proyectos()
    nueva_cant_tareas = max(1, int(nueva_cant_tareas))
    mask_t = (st.session_state.tasks["Proyecto"] == proy_name) & (st.session_state.tasks["Partida"] == nombre_partida_completo)
    df_t_p = st.session_state.tasks[mask_t].copy().reset_index(drop=True)
    actual_cant = len(df_t_p)

    if nueva_cant_tareas > actual_cant:
        nuevas_filas = []
        for idx_t in range(actual_cant + 1, nueva_cant_tareas + 1):
            cod_t = f"{idx_partida_1based}.{idx_t:02d}"
            nuevas_filas.append({
                "Proyecto": proy_name,
                "Partida": nombre_partida_completo,
                "Orden": idx_t,
                "Codigo_Item": cod_t,
                "Tarea": cod_t,
                "Especialidad": "Ninguna",
                "Contratista": "",
                "Estado": "🟡 Planificado",
                "Supervisado": "No",
                "Inicio": datetime.date.today().strftime("%Y-%m-%d"),
                "Fin": datetime.date.today().strftime("%Y-%m-%d"),
                "Causa_Cumplimiento": "",
                "Culpa_Contratista": False,
                "Permitir_Cruce": False,
                "Dependencia": "",
                "Notas": "",
                "Moneda": "Soles",
                "Unidad": "glb",
                "Metrado": 1.0,
                "Valor_Unitario": 0.0,
                "Presupuesto": 0.0,
            })
        df_res_p = pd.concat([df_t_p, pd.DataFrame(nuevas_filas)[COLS_TASKS]], ignore_index=True)
        _reemplazar_tareas_partida_in_place(proy_name, nombre_partida_completo, df_res_p)
    elif nueva_cant_tareas < actual_cant:
        df_keep = df_t_p.iloc[:nueva_cant_tareas].copy()
        for i_k in range(len(df_keep)):
            cod_k = f"{idx_partida_1based}.{i_k + 1:02d}"
            desc_k = extraer_descripcion_tarea(df_keep.at[i_k, "Tarea"], safe_str(df_keep.at[i_k, "Codigo_Item"]))
            df_keep.at[i_k, "Orden"] = i_k + 1
            df_keep.at[i_k, "Codigo_Item"] = cod_k
            df_keep.at[i_k, "Tarea"] = f"{cod_k} {desc_k}".strip() if desc_k else cod_k
        _reemplazar_tareas_partida_in_place(proy_name, nombre_partida_completo, df_keep)

    clear_base_dfs()
    recalcular_totales_presupuesto_proyecto(proy_name)


def mover_tarea_presupuesto(
    proy_name: str,
    idx_partida_1based: int,
    nombre_partida_completo: str,
    pos_tarea_0based: int,
    direccion: str,
) -> bool:
    """
    Mueve verticalmente una tarea ('subir' o 'bajar') dentro de su Partida,
    renumerando automáticamente Codigo_Item (1.01, 1.02...) y Orden sin perder datos.
    """
    guardar_snapshot_presupuesto(proy_name)
    _sanitizar_dtypes_tasks_y_proyectos()
    mask_t = (st.session_state.tasks["Proyecto"] == proy_name) & (st.session_state.tasks["Partida"] == nombre_partida_completo)
    df_t_p = st.session_state.tasks[mask_t].copy().reset_index(drop=True)
    n = len(df_t_p)
    if n <= 1:
        return False

    target_pos = pos_tarea_0based - 1 if direccion == "subir" else pos_tarea_0based + 1
    if not (0 <= pos_tarea_0based < n and 0 <= target_pos < n):
        return False

    # Extraer descripciones limpias antes de intercambiar filas
    descs = [
        extraer_descripcion_tarea(df_t_p.at[i, "Tarea"], f"{idx_partida_1based}.{i + 1:02d}")
        for i in range(n)
    ]

    # Intercambiar posiciones
    orden_idx = list(range(n))
    orden_idx[pos_tarea_0based], orden_idx[target_pos] = orden_idx[target_pos], orden_idx[pos_tarea_0based]
    descs[pos_tarea_0based], descs[target_pos] = descs[target_pos], descs[pos_tarea_0based]

    df_reord = df_t_p.iloc[orden_idx].copy().reset_index(drop=True)
    for i_new in range(n):
        cod_new = f"{idx_partida_1based}.{i_new + 1:02d}"
        d_limp = descs[i_new]
        df_reord.at[i_new, "Orden"] = i_new + 1
        df_reord.at[i_new, "Codigo_Item"] = cod_new
        df_reord.at[i_new, "Tarea"] = f"{cod_new} {d_limp}".strip() if d_limp else cod_new

    _reemplazar_tareas_partida_in_place(proy_name, nombre_partida_completo, df_reord)
    clear_base_dfs()
    recalcular_totales_presupuesto_proyecto(proy_name)
    return True


def reordenar_tareas_partida_por_indices(
    proy_name: str,
    idx_partida_1based: int,
    nombre_partida_completo: str,
    nuevos_indices_0based: list[int],
) -> bool:
    """
    Reordena las tareas de una Partida según la permutación `nuevos_indices_0based`
    (proveniente de arrastrar y soltar / drag-and-drop), renumerando automáticamente
    Orden (1, 2, 3...) y Codigo_Item (1.01, 1.02, 1.03...).
    """
    guardar_snapshot_presupuesto(proy_name)
    _sanitizar_dtypes_tasks_y_proyectos()
    mask_t = (st.session_state.tasks["Proyecto"] == proy_name) & (st.session_state.tasks["Partida"] == nombre_partida_completo)
    df_t_p = st.session_state.tasks[mask_t].copy().reset_index(drop=True)
    n = len(df_t_p)
    if n <= 1:
        return False

    valid_idx = [int(i) for i in nuevos_indices_0based if 0 <= int(i) < n]
    # Completar cualquier índice faltante por seguridad
    for i_f in range(n):
        if i_f not in valid_idx:
            valid_idx.append(i_f)
    if valid_idx == list(range(n)):
        return False

    descs = [
        extraer_descripcion_tarea(df_t_p.at[i, "Tarea"], f"{idx_partida_1based}.{i + 1:02d}")
        for i in range(n)
    ]

    df_reord = df_t_p.iloc[valid_idx].copy().reset_index(drop=True)
    descs_reord = [descs[i] for i in valid_idx]
    for i_new in range(n):
        cod_new = f"{idx_partida_1based}.{i_new + 1:02d}"
        d_limp = descs_reord[i_new]
        df_reord.at[i_new, "Orden"] = i_new + 1
        df_reord.at[i_new, "Codigo_Item"] = cod_new
        df_reord.at[i_new, "Tarea"] = f"{cod_new} {d_limp}".strip() if d_limp else cod_new

    _reemplazar_tareas_partida_in_place(proy_name, nombre_partida_completo, df_reord)
    clear_base_dfs()
    recalcular_totales_presupuesto_proyecto(proy_name)
    return True


def mover_partida_presupuesto(
    proy_name: str,
    idx_partida_1based: int,
    direccion: str,
) -> bool:
    """
    Intercambia la posición de una Partida completa con la anterior ('subir') o posterior ('bajar'),
    renumerando automáticamente los nombres (1. ..., 2. ...) y los códigos de tarea (1.01, 2.01...).
    """
    guardar_snapshot_presupuesto(proy_name)
    _sanitizar_dtypes_tasks_y_proyectos()

    if "partidas" not in st.session_state or st.session_state.partidas.empty:
        return False

    parts_proy = st.session_state.partidas[st.session_state.partidas["Proyecto"] == proy_name].reset_index(drop=True)
    n = len(parts_proy)
    curr_pos = int(idx_partida_1based) - 1
    target_pos = curr_pos - 1 if direccion == "subir" else curr_pos + 1

    if not (0 <= curr_pos < n and 0 <= target_pos < n):
        return False

    p_actuales = parts_proy["Nombre_Partida"].tolist()
    p_actuales[curr_pos], p_actuales[target_pos] = p_actuales[target_pos], p_actuales[curr_pos]

    map_renombre = {}
    for i_new, old_p_name in enumerate(p_actuales, start=1):
        limp = limpiar_nombre_partida(old_p_name, i_new)
        nuevo_nom = formatear_nombre_partida(i_new, limp)
        map_renombre[old_p_name] = (i_new, nuevo_nom)

    nuevas_partidas_proy = []
    for old_p_name in p_actuales:
        _, nuevo_nom = map_renombre[old_p_name]
        nuevas_partidas_proy.append({"Proyecto": proy_name, "Nombre_Partida": nuevo_nom})

    st.session_state.partidas = st.session_state.partidas[st.session_state.partidas["Proyecto"] != proy_name].reset_index(drop=True)
    st.session_state.partidas = pd.concat([st.session_state.partidas, pd.DataFrame(nuevas_partidas_proy)], ignore_index=True)

    if "p_info" in st.session_state and not st.session_state.p_info.empty:
        mask_pi = st.session_state.p_info["Proyecto"] == proy_name
        for old_p, (i_new, nuevo_nom) in map_renombre.items():
            st.session_state.p_info.loc[mask_pi & (st.session_state.p_info["Nombre_Partida"] == old_p), "Nombre_Partida"] = nuevo_nom

    if "tasks" in st.session_state and not st.session_state.tasks.empty:
        mask_t = st.session_state.tasks["Proyecto"] == proy_name
        df_proy_t = st.session_state.tasks[mask_t].copy()
        for old_p, (i_new, nuevo_nom) in map_renombre.items():
            m_tp = df_proy_t["Partida"] == old_p
            df_proy_t.loc[m_tp, "Partida"] = nuevo_nom
            indices_sub = df_proy_t[m_tp].index.tolist()
            for pos_sub, idx_sub in enumerate(indices_sub, start=1):
                cod_sub = f"{i_new}.{pos_sub:02d}"
                desc_sub = extraer_descripcion_tarea(df_proy_t.at[idx_sub, "Tarea"], cod_sub)
                df_proy_t.at[idx_sub, "Orden"] = pos_sub
                df_proy_t.at[idx_sub, "Codigo_Item"] = cod_sub
                df_proy_t.at[idx_sub, "Tarea"] = f"{cod_sub} {desc_sub}".strip() if desc_sub else cod_sub

        st.session_state.tasks = st.session_state.tasks[~mask_t].reset_index(drop=True)
        st.session_state.tasks = pd.concat([st.session_state.tasks, df_proy_t], ignore_index=True)

    if "pagos" in st.session_state and not st.session_state.pagos.empty:
        mask_pag = st.session_state.pagos["Proyecto"] == proy_name
        for old_p, (i_new, nuevo_nom) in map_renombre.items():
            st.session_state.pagos.loc[mask_pag & (st.session_state.pagos["Partida"] == old_p), "Partida"] = nuevo_nom

    clear_base_dfs()
    recalcular_totales_presupuesto_proyecto(proy_name)
    return True


def aplicar_edicion_tabla_presupuesto_partida(
    base_key: str,
    widget_key: str,
    proy_name: str,
    idx_partida_1based: int,
    nombre_partida_completo: str,
):
    """
    Callback on_change para el st.data_editor de cada Partida en '03. PRESUPUESTO'.
    Actualiza Descripción, Especialidad, Contratista, Moneda, Unidad, Metrado, Valor Unitario y Presupuesto
    en st.session_state.tasks sincronizado con '05. PARTIDAS Y TAREAS' y '06. COBROS & FINANZAS'.
    """
    if widget_key not in st.session_state or base_key not in st.session_state:
        return

    edits = st.session_state[widget_key]
    edited_rows = edits.get("edited_rows", {})
    if not edited_rows:
        return

    guardar_snapshot_presupuesto(proy_name)

    df_base = st.session_state[base_key].copy()
    for c_num_b in ["Metrado", "Valor_Unitario", "Valor_Parcial"]:
        if c_num_b in df_base.columns:
            df_base[c_num_b] = pd.to_numeric(df_base[c_num_b], errors="coerce").fillna(0.0).astype(float)
    for c_txt_b in ["⠿", "Item", "Descripcion", "Especialidad", "Contratista", "Moneda", "Unidad"]:
        if c_txt_b in df_base.columns and df_base[c_txt_b].dtype != object:
            df_base[c_txt_b] = df_base[c_txt_b].fillna("").astype(str).astype(object)

    for r_str, changes in edited_rows.items():
        r_i = int(r_str)
        if r_i in df_base.index:
            for col_k, val in changes.items():
                if col_k in ("Metrado", "Valor_Unitario", "Valor_Parcial"):
                    df_base.at[r_i, col_k] = safe_float(val, 0.0)
                else:
                    if col_k in df_base.columns and df_base[col_k].dtype != object:
                        df_base[col_k] = df_base[col_k].astype(object)
                    df_base.at[r_i, col_k] = safe_str(val)

    _sanitizar_dtypes_tasks_y_proyectos()
    params = obtener_parametros_presupuesto_proyecto(proy_name)
    tc = params["tc"] if params["tc"] > 0 else 3.42

    mask_t = (st.session_state.tasks["Proyecto"] == proy_name) & (st.session_state.tasks["Partida"] == nombre_partida_completo)
    indices_tasks = st.session_state.tasks[mask_t].index.tolist()

    for pos_i, (_, row_ed) in enumerate(df_base.iterrows()):
        if pos_i >= len(indices_tasks):
            break
        t_idx_global = indices_tasks[pos_i]
        cod_t = f"{idx_partida_1based}.{pos_i + 1:02d}"
        desc_limpia = safe_str(row_ed.get("Descripcion", "")).strip()
        tarea_completa = f"{cod_t} {desc_limpia}".strip() if desc_limpia else cod_t
        moneda = safe_str(row_ed.get("Moneda", "Soles")) or "Soles"
        unidad = safe_str(row_ed.get("Unidad", "glb")) or "glb"
        metrado = safe_float(row_ed.get("Metrado", 1.0), 1.0)
        if metrado <= 0:
            metrado = 1.0
        v_unit = safe_float(row_ed.get("Valor_Unitario", 0.0), 0.0)
        es_usd = any(d in moneda.lower() for d in ["dolar", "dólar", "usd", "$"])
        ppto_soles = round(metrado * v_unit * (tc if es_usd else 1.0), 2)

        st.session_state.tasks.at[t_idx_global, "Orden"] = int(pos_i + 1)
        st.session_state.tasks.at[t_idx_global, "Codigo_Item"] = cod_t
        st.session_state.tasks.at[t_idx_global, "Tarea"] = tarea_completa
        st.session_state.tasks.at[t_idx_global, "Moneda"] = moneda
        st.session_state.tasks.at[t_idx_global, "Unidad"] = unidad
        st.session_state.tasks.at[t_idx_global, "Metrado"] = float(metrado)
        st.session_state.tasks.at[t_idx_global, "Valor_Unitario"] = float(v_unit)
        st.session_state.tasks.at[t_idx_global, "Presupuesto"] = float(ppto_soles)
        if "Especialidad" in row_ed:
            esp_v = safe_str(row_ed.get("Especialidad", "Ninguna")).strip() or "Ninguna"
            st.session_state.tasks.at[t_idx_global, "Especialidad"] = esp_v
        if "Contratista" in row_ed:
            con_v = safe_str(row_ed.get("Contratista", "")).strip()
            st.session_state.tasks.at[t_idx_global, "Contratista"] = con_v

    if base_key in st.session_state:
        del st.session_state[base_key]
    c_key = f"c_{base_key}"
    st.session_state[c_key] = st.session_state.get(c_key, 0) + 1

    recalcular_totales_presupuesto_proyecto(proy_name)


# ==========================================
# 2. GENERADOR Y EXPORTADOR A EXCEL (.XLSX) PARA CLIENTES
# ==========================================

def generar_plantilla_cotizacion_venti(
    proyecto: str = "",
    cliente: str = "",
    tc: float = 3.42,
    fecha: str = None,
    partidas_precargadas: list = None,
    gg_pct: float = 0.02,
    gestion_pct: float = 0.15,
    incluir_igv: bool = False,
) -> bytes:
    """
    Construye un archivo Excel (.xlsx) con el Formato Oficial de Cotización COTA.
    Incluye:
    - Hoja 1 ('Presupuesto'): Partidas (1, 2, 3...), Tareas (1.01, 1.02...), fórmulas de Valor Parcial,
      Sub Total (S/.), Mix (%), SUB TOTAL (Casco), Gastos generales, Gestión y supervisión de obras,
      IGV (18% si incluir_igv=True o 0% si False) y TOTAL integrado en Soles y Dólares.
    - Hoja 2 ('Control de Pagos'): Cronograma de desembolsos por semanas integrado simétricamente.
    """
    if not fecha:
        fecha = datetime.date.today().strftime("%Y-%m-%d")

    gg_dec = float(gg_pct) / 100.0 if float(gg_pct) > 1.0 else float(gg_pct)
    ges_dec = float(gestion_pct) / 100.0 if float(gestion_pct) > 1.0 else float(gestion_pct)

    wb = openpyxl.Workbook()

    font_family = "Calibri"
    f_title = Font(name=font_family, size=11, bold=True, color="262320")
    f_header = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    f_partida = Font(name=font_family, size=10, bold=True, color="262320")
    f_task = Font(name=font_family, size=10, bold=False, color="262320")
    f_footer = Font(name=font_family, size=10, bold=True, color="262320")
    f_total = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    f_usd = Font(name=font_family, size=10, bold=True, color="262320")

    fill_header = PatternFill(start_color="262320", end_color="262320", fill_type="solid")
    fill_partida = PatternFill(start_color="EFECE6", end_color="EFECE6", fill_type="solid")
    fill_subtotal = PatternFill(start_color="DFD9CE", end_color="DFD9CE", fill_type="solid")
    fill_total = PatternFill(start_color="C85A32", end_color="C85A32", fill_type="solid")
    fill_usd = PatternFill(start_color="EFECE6", end_color="EFECE6", fill_type="solid")
    fill_input = PatternFill(start_color="FAF8F5", end_color="FAF8F5", fill_type="solid")

    thin_side = Side(border_style="thin", color="D0C9BD")
    med_bottom = Side(border_style="medium", color="262320")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_partida = Border(left=thin_side, right=thin_side, top=thin_side, bottom=med_bottom)
    border_closing = Border(left=thin_side, right=thin_side, top=thin_side, bottom=med_bottom)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # ==========================================
    # HOJA 1: PRESUPUESTO
    # ==========================================
    ws1 = wb.active
    ws1.title = "Presupuesto"
    ws1.views.sheetView[0].showGridLines = True

    ws1["B2"] = "Proyecto"
    ws1["B2"].font = f_title
    ws1["C2"] = proyecto
    ws1["C2"].font = f_title
    ws1["I2"] = fecha
    ws1["I2"].font = f_task
    ws1["I2"].alignment = align_center

    ws1["B4"] = "Cliente:"
    ws1["B4"].font = f_title
    ws1["C4"] = cliente
    ws1["C4"].font = f_title
    ws1["I4"] = "TC:"
    ws1["I4"].font = f_title
    ws1["I4"].alignment = align_right
    ws1["J4"] = float(tc)
    ws1["J4"].font = f_title
    ws1["J4"].number_format = "0.00"
    ws1["J4"].alignment = align_center

    headers_h1 = [
        ("B", "Ítem"),
        ("C", "Descripción"),
        ("D", "Moneda"),
        ("E", "Unidad"),
        ("F", "Metrado"),
        ("G", "Valor\nUnitario"),
        ("H", "Valor\nParcial"),
        ("I", "Sub Total (S/.)"),
        ("J", "Mix"),
    ]
    ws1.row_dimensions[5].height = 28
    for col_let, text in headers_h1:
        cell = ws1[f"{col_let}5"]
        cell.value = text
        cell.font = f_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    estructura = partidas_precargadas if partidas_precargadas else ESTRUCTURA_COTIZACION_DEFAULT

    curr_row = 6
    partida_rows = []

    for part in estructura:
        p_row = curr_row
        partida_rows.append(p_row)
        p_cod = part.get("codigo", str(len(partida_rows)))
        p_nom = part.get("nombre", "")
        subitems = part.get("subitems", [])

        first_task_row = p_row + 1
        last_task_row = p_row + len(subitems)

        ws1[f"B{p_row}"] = int(p_cod) if str(p_cod).isdigit() else p_cod
        ws1[f"B{p_row}"].alignment = align_center
        ws1[f"C{p_row}"] = p_nom
        ws1[f"C{p_row}"].alignment = align_left

        if last_task_row >= first_task_row:
            ws1[f"I{p_row}"] = f"=SUM(H{first_task_row}:H{last_task_row})"
        else:
            ws1[f"I{p_row}"] = float(part.get("monto_pactado", 0.0))

        ws1[f"I{p_row}"].number_format = "#,##0.00"
        ws1[f"I{p_row}"].alignment = align_right
        ws1[f"J{p_row}"].number_format = "0.0%"
        ws1[f"J{p_row}"].alignment = align_center

        for c_idx in range(2, 11):
            c_let = get_column_letter(c_idx)
            ws1[f"{c_let}{p_row}"].font = f_partida
            ws1[f"{c_let}{p_row}"].fill = fill_partida
            ws1[f"{c_let}{p_row}"].border = border_partida

        curr_row += 1

        for sub in subitems:
            t_row = curr_row
            if isinstance(sub, dict):
                t_cod = sub.get("codigo", "")
                t_desc = sub.get("descripcion", "")
                t_mon = sub.get("moneda", "Soles")
                t_und = sub.get("unidad", "glb")
                t_met = safe_float(sub.get("metrado", 1), 1.0)
                t_vu = sub.get("valor_unitario", None)
            else:
                t_cod = str(sub)
                t_desc = ""
                t_mon = "Soles"
                t_und = "glb"
                t_met = 1
                t_vu = None

            try:
                ws1[f"B{t_row}"] = float(t_cod)
                ws1[f"B{t_row}"].number_format = "0.00"
            except ValueError:
                ws1[f"B{t_row}"] = t_cod
            ws1[f"B{t_row}"].alignment = align_center

            ws1[f"C{t_row}"] = t_desc
            ws1[f"C{t_row}"].alignment = align_left
            ws1[f"C{t_row}"].fill = fill_input

            ws1[f"D{t_row}"] = t_mon
            ws1[f"D{t_row}"].alignment = align_center

            ws1[f"E{t_row}"] = t_und
            ws1[f"E{t_row}"].alignment = align_center

            ws1[f"F{t_row}"] = t_met
            ws1[f"F{t_row}"].alignment = align_center
            ws1[f"F{t_row}"].number_format = "#,##0.00"
            ws1[f"F{t_row}"].fill = fill_input

            if t_vu is not None and safe_float(t_vu, 0.0) > 0:
                ws1[f"G{t_row}"] = safe_float(t_vu, 0.0)
            ws1[f"G{t_row}"].alignment = align_right
            ws1[f"G{t_row}"].number_format = "#,##0.00"
            ws1[f"G{t_row}"].fill = fill_input

            ws1[f"H{t_row}"] = f'=IF(OR(UPPER(TRIM(D{t_row}))="DÓLARES",UPPER(TRIM(D{t_row}))="DOLARES",UPPER(TRIM(D{t_row}))="USD"),F{t_row}*G{t_row}*$J$4,F{t_row}*G{t_row})'
            ws1[f"H{t_row}"].alignment = align_right
            ws1[f"H{t_row}"].number_format = "#,##0.00"

            for c_idx in range(2, 11):
                c_let = get_column_letter(c_idx)
                ws1[f"{c_let}{t_row}"].font = f_task
                ws1[f"{c_let}{t_row}"].border = border_cell

            curr_row += 1

    # Fila de SUB TOTAL (Casco)
    row_subtotal_casco = curr_row
    for p_r in partida_rows:
        ws1[f"J{p_r}"] = f"=IFERROR(I{p_r}/$I${row_subtotal_casco},0)"

    sum_parts_expr = "+".join([f"I{r}" for r in partida_rows]) if partida_rows else "0"
    ws1[f"C{row_subtotal_casco}"] = "SUB TOTAL (Casco)"
    ws1[f"D{row_subtotal_casco}"] = "Soles"
    ws1[f"D{row_subtotal_casco}"].alignment = align_center
    ws1[f"I{row_subtotal_casco}"] = f"={sum_parts_expr}"
    ws1[f"I{row_subtotal_casco}"].number_format = "#,##0.00"
    ws1[f"I{row_subtotal_casco}"].alignment = align_right
    ws1[f"J{row_subtotal_casco}"] = f"=IFERROR(I{row_subtotal_casco}/I{row_subtotal_casco},0)"
    ws1[f"J{row_subtotal_casco}"].number_format = "0.0%"
    ws1[f"J{row_subtotal_casco}"].alignment = align_center

    for c_idx in range(2, 11):
        c_let = get_column_letter(c_idx)
        ws1[f"{c_let}{row_subtotal_casco}"].font = f_footer
        ws1[f"{c_let}{row_subtotal_casco}"].fill = fill_subtotal
        ws1[f"{c_let}{row_subtotal_casco}"].border = border_cell

    # Fila de Gastos generales
    row_gg = row_subtotal_casco + 1
    ws1[f"C{row_gg}"] = "Gastos generales"
    ws1[f"D{row_gg}"] = "Soles"
    ws1[f"D{row_gg}"].alignment = align_center
    ws1[f"F{row_gg}"] = gg_dec
    ws1[f"F{row_gg}"].number_format = "0.0%"
    ws1[f"F{row_gg}"].alignment = align_center
    ws1[f"H{row_gg}"] = f"=I{row_subtotal_casco}*F{row_gg}"
    ws1[f"H{row_gg}"].number_format = "#,##0.00"
    ws1[f"H{row_gg}"].alignment = align_right
    for c_idx in range(2, 11):
        c_let = get_column_letter(c_idx)
        ws1[f"{c_let}{row_gg}"].font = f_footer
        ws1[f"{c_let}{row_gg}"].border = border_cell

    # Fila de Gestión y supervisión de obras (Utilidades)
    row_gestion = row_gg + 1
    ws1[f"C{row_gestion}"] = "Gestión y supervisión de obras"
    ws1[f"D{row_gestion}"] = "Soles"
    ws1[f"D{row_gestion}"].alignment = align_center
    ws1[f"F{row_gestion}"] = ges_dec
    ws1[f"F{row_gestion}"].number_format = "0.0%"
    ws1[f"F{row_gestion}"].alignment = align_center
    ws1[f"H{row_gestion}"] = f"=I{row_subtotal_casco}*F{row_gestion}"
    ws1[f"H{row_gestion}"].number_format = "#,##0.00"
    ws1[f"H{row_gestion}"].alignment = align_right
    for c_idx in range(2, 11):
        c_let = get_column_letter(c_idx)
        ws1[f"{c_let}{row_gestion}"].font = f_footer
        ws1[f"{c_let}{row_gestion}"].border = border_cell

    # Fila de IGV
    row_igv = row_gestion + 1
    ws1[f"C{row_igv}"] = "IGV"
    ws1[f"D{row_igv}"] = "Soles"
    ws1[f"D{row_igv}"].alignment = align_center
    ws1[f"F{row_igv}"] = 0.18 if incluir_igv else 0.0
    ws1[f"F{row_igv}"].number_format = "0%"
    ws1[f"F{row_igv}"].alignment = align_center
    ws1[f"H{row_igv}"] = f"=IF(ISNUMBER(F{row_igv}),(I{row_subtotal_casco}+H{row_gg}+H{row_gestion})*F{row_igv},0)"
    ws1[f"H{row_igv}"].number_format = "#,##0.00"
    ws1[f"H{row_igv}"].alignment = align_right
    for c_idx in range(2, 11):
        c_let = get_column_letter(c_idx)
        ws1[f"{c_let}{row_igv}"].font = f_footer
        ws1[f"{c_let}{row_igv}"].border = border_cell

    # Fila TOTAL Soles
    row_total = row_igv + 1
    ws1[f"C{row_total}"] = "TOTAL"
    ws1[f"D{row_total}"] = "Soles"
    ws1[f"D{row_total}"].alignment = align_center
    ws1[f"I{row_total}"] = f"=I{row_subtotal_casco}+H{row_gg}+H{row_gestion}+H{row_igv}"
    ws1[f"I{row_total}"].number_format = "#,##0.00"
    ws1[f"I{row_total}"].alignment = align_right
    ws1[f"J{row_total}"] = f"=IFERROR(I{row_total}/I{row_subtotal_casco},0)"
    ws1[f"J{row_total}"].number_format = "0.0%"
    ws1[f"J{row_total}"].alignment = align_center
    for c_idx in range(2, 11):
        c_let = get_column_letter(c_idx)
        ws1[f"{c_let}{row_total}"].font = f_total
        ws1[f"{c_let}{row_total}"].fill = fill_total
        ws1[f"{c_let}{row_total}"].border = border_cell

    # Fila TOTAL Dólares integrada simétricamente al cuadro
    row_usd = row_total + 1
    ws1[f"C{row_usd}"] = "TOTAL (USD)"
    ws1[f"D{row_usd}"] = "Dólares"
    ws1[f"D{row_usd}"].alignment = align_center
    ws1[f"I{row_usd}"] = f"=IFERROR(I{row_total}/J4,0)"
    ws1[f"I{row_usd}"].number_format = "#,##0.00"
    ws1[f"I{row_usd}"].alignment = align_right
    for c_idx in range(2, 11):
        c_let = get_column_letter(c_idx)
        ws1[f"{c_let}{row_usd}"].font = f_usd
        ws1[f"{c_let}{row_usd}"].fill = fill_usd
        ws1[f"{c_let}{row_usd}"].border = border_closing

    col_widths_h1 = {
        "A": 3,
        "B": 9,
        "C": 44,
        "D": 12,
        "E": 10,
        "F": 12,
        "G": 15,
        "H": 15,
        "I": 16,
        "J": 12,
    }
    for col_l, w in col_widths_h1.items():
        ws1.column_dimensions[col_l].width = w

    # ==========================================
    # HOJA 2: CONTROL DE PAGOS
    # ==========================================
    ws2 = wb.create_sheet(title="Control de Pagos")
    ws2.views.sheetView[0].showGridLines = True

    ws2["B1"] = "Proyecto"
    ws2["B1"].font = f_title
    ws2["C1"] = "=Presupuesto!C2"
    ws2["C1"].font = f_title

    ws2["B2"] = "Cliente:"
    ws2["B2"].font = f_title
    ws2["C2"] = "=Presupuesto!C4"
    ws2["C2"].font = f_title

    headers_h2 = [
        ("B", "Descripción"),
        ("C", "Moneda"),
        ("D", "TOTAL"),
        ("E", "ADELANTO 1"),
        ("F", "SEMANA 1"),
        ("G", "SEMANA 2"),
        ("H", "SEMANA 3"),
        ("I", "SEMANA 4"),
        ("J", "SEMANA 5"),
        ("K", "PAGADO"),
        ("L", "SALDO"),
    ]
    ws2.row_dimensions[4].height = 24
    for col_let, text in headers_h2:
        cell = ws2[f"{col_let}4"]
        cell.value = text
        cell.font = f_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    ws2["B6"] = "TRABAJOS CONTRATISTAS"
    ws2["B6"].font = f_partida

    r_c = 7
    first_c_row = r_c
    for idx_p, p_r in enumerate(partida_rows, start=1):
        ws2[f"B{r_c}"] = f'=IF(Presupuesto!C{p_r}="","Partida {idx_p}",Presupuesto!C{p_r})'
        ws2[f"C{r_c}"] = "Soles"
        ws2[f"C{r_c}"].alignment = align_center
        ws2[f"D{r_c}"] = f"=Presupuesto!I{p_r}"
        ws2[f"D{r_c}"].number_format = "#,##0.00"
        ws2[f"K{r_c}"] = f"=SUM(E{r_c}:J{r_c})"
        ws2[f"K{r_c}"].number_format = "#,##0.00"
        ws2[f"L{r_c}"] = f"=D{r_c}-K{r_c}"
        ws2[f"L{r_c}"].number_format = "#,##0.00"
        for c_i in range(2, 13):
            c_l = get_column_letter(c_i)
            ws2[f"{c_l}{r_c}"].font = f_task
            ws2[f"{c_l}{r_c}"].border = border_cell
        r_c += 1
    last_c_row = r_c - 1

    row_sub_h2 = r_c
    ws2[f"B{row_sub_h2}"] = "SUB TOTAL"
    ws2[f"C{row_sub_h2}"] = "Soles"
    ws2[f"C{row_sub_h2}"].alignment = align_center
    for c_i in range(4, 13):
        c_l = get_column_letter(c_i)
        ws2[f"{c_l}{row_sub_h2}"] = f"=SUM({c_l}{first_c_row}:{c_l}{last_c_row})"
        ws2[f"{c_l}{row_sub_h2}"].number_format = "#,##0.00"
    for c_i in range(2, 13):
        c_l = get_column_letter(c_i)
        ws2[f"{c_l}{row_sub_h2}"].font = f_footer
        ws2[f"{c_l}{row_sub_h2}"].fill = fill_subtotal
        ws2[f"{c_l}{row_sub_h2}"].border = border_cell

    row_ges_h2 = row_sub_h2 + 1
    ws2[f"B{row_ges_h2}"] = "Gestión y supervisión de obras"
    ws2[f"C{row_ges_h2}"] = "Soles"
    ws2[f"C{row_ges_h2}"].alignment = align_center
    ws2[f"D{row_ges_h2}"] = f"=Presupuesto!H{row_gestion}"
    ws2[f"D{row_ges_h2}"].number_format = "#,##0.00"
    ws2[f"K{row_ges_h2}"] = f"=SUM(E{row_ges_h2}:J{row_ges_h2})"
    ws2[f"K{row_ges_h2}"].number_format = "#,##0.00"
    ws2[f"L{row_ges_h2}"] = f"=D{row_ges_h2}-K{row_ges_h2}"
    ws2[f"L{row_ges_h2}"].number_format = "#,##0.00"
    for c_i in range(2, 13):
        c_l = get_column_letter(c_i)
        ws2[f"{c_l}{row_ges_h2}"].font = f_task
        ws2[f"{c_l}{row_ges_h2}"].border = border_cell

    row_gg_h2 = row_ges_h2 + 1
    ws2[f"B{row_gg_h2}"] = "Gastos Generales"
    ws2[f"C{row_gg_h2}"] = "Soles"
    ws2[f"C{row_gg_h2}"].alignment = align_center
    ws2[f"D{row_gg_h2}"] = f"=Presupuesto!H{row_gg}"
    ws2[f"D{row_gg_h2}"].number_format = "#,##0.00"
    ws2[f"K{row_gg_h2}"] = f"=SUM(E{row_gg_h2}:J{row_gg_h2})"
    ws2[f"K{row_gg_h2}"].number_format = "#,##0.00"
    ws2[f"L{row_gg_h2}"] = f"=D{row_gg_h2}-K{row_gg_h2}"
    ws2[f"L{row_gg_h2}"].number_format = "#,##0.00"
    for c_i in range(2, 13):
        c_l = get_column_letter(c_i)
        ws2[f"{c_l}{row_gg_h2}"].font = f_task
        ws2[f"{c_l}{row_gg_h2}"].border = border_cell

    row_tot_h2 = row_gg_h2 + 1
    ws2[f"B{row_tot_h2}"] = "TOTAL"
    ws2[f"C{row_tot_h2}"] = "Soles"
    ws2[f"C{row_tot_h2}"].alignment = align_center
    for c_i in range(4, 13):
        c_l = get_column_letter(c_i)
        ws2[f"{c_l}{row_tot_h2}"] = f"={c_l}{row_sub_h2}+{c_l}{row_ges_h2}+{c_l}{row_gg_h2}"
        ws2[f"{c_l}{row_tot_h2}"].number_format = "#,##0.00"
    for c_i in range(2, 13):
        c_l = get_column_letter(c_i)
        ws2[f"{c_l}{row_tot_h2}"].font = f_total
        ws2[f"{c_l}{row_tot_h2}"].fill = fill_total
        ws2[f"{c_l}{row_tot_h2}"].border = border_cell

    row_usd_h2 = row_tot_h2 + 1
    ws2[f"B{row_usd_h2}"] = "TOTAL (USD)"
    ws2[f"C{row_usd_h2}"] = "Dólares"
    ws2[f"C{row_usd_h2}"].alignment = align_center
    for c_i in range(4, 13):
        c_l = get_column_letter(c_i)
        ws2[f"{c_l}{row_usd_h2}"] = f"=IFERROR({c_l}{row_tot_h2}/Presupuesto!$J$4,0)"
        ws2[f"{c_l}{row_usd_h2}"].number_format = "#,##0.00"
    for c_i in range(2, 13):
        c_l = get_column_letter(c_i)
        ws2[f"{c_l}{row_usd_h2}"].font = f_usd
        ws2[f"{c_l}{row_usd_h2}"].fill = fill_usd
        ws2[f"{c_l}{row_usd_h2}"].border = border_closing

    col_widths_h2 = {
        "A": 3,
        "B": 32,
        "C": 14,
        "D": 15,
        "E": 15,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 15,
        "L": 15,
    }
    for col_l, w in col_widths_h2.items():
        ws2.column_dimensions[col_l].width = w

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def exportar_proyecto_a_formato_venti(
    proy_name: str,
    tc: float = None,
    gg_pct_override: float = None,
    gestion_pct_override: float = None,
    incluir_igv: bool = None,
) -> bytes:
    """
    Exporta el presupuesto actual del proyecto en COTA a un archivo Excel (.xlsx)
    listo para compartir con el cliente.
    """
    params = obtener_parametros_presupuesto_proyecto(proy_name)
    tc_final = tc if tc is not None else params["tc"]
    gg_final = gg_pct_override if gg_pct_override is not None else (params["gg_pct"] / 100.0)
    ges_final = gestion_pct_override if gestion_pct_override is not None else (params["utilidades_pct"] / 100.0)
    igv_final = incluir_igv if incluir_igv is not None else params["incluir_igv"]

    df_p = st.session_state.get("proyectos", pd.DataFrame())
    proy_row = df_p[df_p["Nombre"] == proy_name].iloc[0] if not df_p.empty and proy_name in df_p["Nombre"].values else {}

    cliente = safe_str(proy_row.get("Cliente", ""))
    fecha = safe_str(proy_row.get("Inicio", "")) or datetime.date.today().strftime("%Y-%m-%d")

    df_parts = st.session_state.get("partidas", pd.DataFrame())
    parts_proy = df_parts[df_parts["Proyecto"] == proy_name] if not df_parts.empty and "Proyecto" in df_parts.columns else pd.DataFrame()

    df_pinfo = st.session_state.get("p_info", pd.DataFrame())
    df_tasks = st.session_state.get("tasks", pd.DataFrame())

    partidas_precargadas = []
    if not parts_proy.empty:
        for idx_p, (_, p_r) in enumerate(parts_proy.iterrows(), start=1):
            p_nom_raw = safe_str(p_r.get("Nombre_Partida", f"Partida {idx_p}"))
            p_nom_clean = limpiar_nombre_partida(p_nom_raw, idx_p)

            pinfo_m = df_pinfo[(df_pinfo["Proyecto"] == proy_name) & (df_pinfo["Nombre_Partida"] == p_nom_raw)] if not df_pinfo.empty else pd.DataFrame()
            monto_pact = safe_float(pinfo_m["Monto_Pactado"].iloc[0], 0.0) if not pinfo_m.empty else 0.0

            tasks_m = df_tasks[(df_tasks["Proyecto"] == proy_name) & (df_tasks["Partida"] == p_nom_raw)] if not df_tasks.empty else pd.DataFrame()
            subitems = []
            if not tasks_m.empty:
                for idx_t, (_, t_r) in enumerate(tasks_m.iterrows(), start=1):
                    cod_item = f"{idx_p}.{idx_t:02d}"
                    t_nom_raw = safe_str(t_r.get("Tarea", ""))
                    t_desc_clean = extraer_descripcion_tarea(t_nom_raw, cod_item)
                    t_mon = safe_str(t_r.get("Moneda", "Soles")) or "Soles"
                    t_und = safe_str(t_r.get("Unidad", "glb")) or "glb"
                    t_met = safe_float(t_r.get("Metrado", 1.0), 1.0)
                    if t_met <= 0:
                        t_met = 1.0
                    t_ppto = safe_float(t_r.get("Presupuesto", 0.0), 0.0)
                    t_vu = safe_float(t_r.get("Valor_Unitario", 0.0), 0.0)
                    if t_vu <= 0 and t_ppto > 0:
                        t_vu = round(t_ppto / t_met, 2)

                    subitems.append({
                        "codigo": cod_item,
                        "descripcion": t_desc_clean,
                        "moneda": t_mon,
                        "unidad": t_und,
                        "metrado": t_met,
                        "valor_unitario": t_vu,
                    })
            else:
                subitems.append({
                    "codigo": f"{idx_p}.01",
                    "descripcion": p_nom_clean if monto_pact > 0 else "",
                    "moneda": "Soles",
                    "unidad": "glb",
                    "metrado": 1.0,
                    "valor_unitario": monto_pact if monto_pact > 0 else None,
                })

            partidas_precargadas.append({
                "codigo": str(idx_p),
                "nombre": p_nom_clean,
                "monto_pactado": monto_pact,
                "subitems": subitems,
            })

    return generar_plantilla_cotizacion_venti(
        proyecto=proy_name,
        cliente=cliente,
        tc=tc_final,
        fecha=fecha,
        partidas_precargadas=partidas_precargadas if partidas_precargadas else None,
        gg_pct=gg_final,
        gestion_pct=ges_final,
        incluir_igv=igv_final,
    )


# ==========================================
# 3. LECTOR E INTERPRETADOR INTELIGENTE DE PRESUPUESTO EXTERNO
# ==========================================

def _clean_cell_str(val) -> str:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    s = str(val).strip()
    if s.lower() in ["nan", "none", "null", "<na>"]:
        return ""
    if s.startswith("#") and any(err in s.upper() for err in ["#DIV/0", "#VALUE", "#REF", "#NAME", "#N/A", "#NUM"]):
        return ""
    return s


def _parse_number_or_pct(val, default=0.0):
    s = _clean_cell_str(val)
    if not s or s in ["-%", "-", "—"]:
        return default, False

    is_pct = "%" in s
    s_clean = re.sub(r"[S/\$\s%USDsolesdólares]", "", s, flags=re.IGNORECASE).strip()
    if not s_clean or s_clean in ["-", "."]:
        return default, is_pct

    num = safe_float(s_clean, default=default)
    if is_pct and abs(num) > 1.0:
        num = num / 100.0
    return num, is_pct


def _classify_item_code(raw_val) -> tuple:
    if raw_val is None or (isinstance(raw_val, float) and pd.isna(raw_val)):
        return None, None, None

    if isinstance(raw_val, int):
        if 1 <= raw_val <= 999:
            return "partida", str(raw_val), None
        return None, None, None

    if isinstance(raw_val, float):
        if raw_val <= 0 or raw_val >= 1000:
            return None, None, None
        if abs(raw_val - round(raw_val)) < 1e-6:
            return "partida", str(int(round(raw_val))), None
        else:
            p_code = str(int(raw_val))
            t_code = f"{raw_val:.2f}"
            return "tarea", p_code, t_code

    s = _clean_cell_str(raw_val)
    if not s:
        return None, None, None

    if re.match(r"^\d+,\d+$", s):
        s = s.replace(",", ".")

    m_part = re.match(r"^0*(\d+)\.?$", s)
    if m_part:
        return "partida", m_part.group(1), None

    m_task = re.match(r"^0*(\d+)\.(\d+(?:\.\d+)?)$", s)
    if m_task:
        p_code = m_task.group(1)
        sub_part = m_task.group(2)
        if len(sub_part) == 1:
            sub_part = f"0{sub_part}"
        return "tarea", p_code, f"{p_code}.{sub_part}"

    return None, None, None


def _rows_from_input(file_or_df) -> list:
    if isinstance(file_or_df, pd.DataFrame):
        return file_or_df.values.tolist()

    raw_bytes = None
    filename = ""

    if hasattr(file_or_df, "getvalue"):
        raw_bytes = file_or_df.getvalue()
        filename = getattr(file_or_df, "name", "").lower()
    elif hasattr(file_or_df, "read"):
        raw_bytes = file_or_df.read()
        if hasattr(file_or_df, "seek"):
            file_or_df.seek(0)
        filename = getattr(file_or_df, "name", "").lower()
    elif isinstance(file_or_df, (bytes, bytearray)):
        raw_bytes = bytes(file_or_df)
    elif isinstance(file_or_df, str):
        s_in = file_or_df.strip()
        if "\n" in s_in or "|" in s_in or "\t" in s_in or ("," in s_in and not s_in.lower().endswith((".xlsx", ".xls", ".csv"))):
            # Soportar tablas Markdown de Gemini Spark (| col1 | col2 | ...)
            lines_md = [ln.strip() for ln in s_in.splitlines() if ln.strip()]
            pipe_lines = [ln for ln in lines_md if "|" in ln]
            if len(pipe_lines) >= 2:
                md_rows = []
                for ln in pipe_lines:
                    # Ignorar separadores markdown como |---|---|
                    if re.match(r"^\|?[\s\-\:\|]+\|?$", ln):
                        continue
                    parts = [p.strip() for p in ln.strip("|").split("|")]
                    md_rows.append([""] + parts)  # Offset col 0 vacía para compatibilidad con col 1..N
                if md_rows:
                    return md_rows
            if "\t" in s_in:
                df_tsv = pd.read_csv(io.StringIO(s_in), sep="\t", header=None, dtype=object)
                return df_tsv.values.tolist()
            df_csv = pd.read_csv(io.StringIO(s_in), header=None, dtype=object)
            return df_csv.values.tolist()
        else:
            filename = s_in.lower()
            with open(file_or_df, "rb") as f:
                raw_bytes = f.read()

    if raw_bytes is None:
        return []

    if filename.endswith((".xlsx", ".xlsm")) or raw_bytes[:4] == b"PK\x03\x04":
        wb_val = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        wb_f = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=False)
        sheet_name = wb_val.sheetnames[0]
        for sn in wb_val.sheetnames:
            if "presupuesto" in sn.lower() or "cotiz" in sn.lower():
                sheet_name = sn
                break
        ws_v = wb_val[sheet_name]
        ws_f = wb_f[sheet_name]

        rows = []
        for r_idx in range(1, ws_v.max_row + 1):
            row_vals = []
            for c_idx in range(1, ws_v.max_column + 1):
                v = ws_v.cell(r_idx, c_idx).value
                f_val = ws_f.cell(r_idx, c_idx).value
                if v is None and f_val is not None and not str(f_val).startswith("="):
                    v = f_val
                row_vals.append(v)
            rows.append(row_vals)
        return rows

    for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
        try:
            text = raw_bytes.decode(enc)
            df_csv = pd.read_csv(io.StringIO(text), header=None, dtype=object)
            return df_csv.values.tolist()
        except Exception:
            continue

    return []


def parsear_cotizacion_venti(file_or_df) -> dict:
    """
    Lee e interpreta un archivo Excel (.xlsx) o CSV con el formato de presupuesto de COTA.
    """
    rows = _rows_from_input(file_or_df)

    resultado = {
        "proyecto": "",
        "cliente": "",
        "fecha": "",
        "tc": 3.42,
        "partidas": [],
        "subtotal_casco": 0.0,
        "gg_pct": 0.0,
        "gg_monto": 0.0,
        "utilidades_pct": 0.0,
        "utilidades_monto": 0.0,
        "igv_pct": 0.0,
        "igv_monto": 0.0,
        "presupuesto_total": 0.0,
        "presupuesto_dolares": 0.0,
    }

    if not rows:
        return resultado

    header_row_idx = None
    col_map = {
        "codigo": 1,
        "descripcion": 2,
        "moneda": 3,
        "unidad": 4,
        "metrado": 5,
        "valor_unitario": 6,
        "valor_parcial": 7,
        "subtotal": 8,
        "mix": 9,
    }

    for r_i, row in enumerate(rows[:25]):
        row_strs = [_clean_cell_str(c).lower() for c in row]
        has_desc = any("descrip" in s for s in row_strs)
        has_mon_or_und = any(("moneda" in s or "unidad" in s or "metrado" in s) for s in row_strs)
        if has_desc and has_mon_or_und:
            header_row_idx = r_i
            for c_i, s in enumerate(row_strs):
                if "descrip" in s:
                    col_map["descripcion"] = c_i
                    if c_i > 0:
                        col_map["codigo"] = c_i - 1
                elif "moneda" in s:
                    col_map["moneda"] = c_i
                elif "unidad" in s:
                    col_map["unidad"] = c_i
                elif "metrado" in s or "cant" in s:
                    col_map["metrado"] = c_i
                elif "unitario" in s:
                    col_map["valor_unitario"] = c_i
                elif "parcial" in s:
                    col_map["valor_parcial"] = c_i
                elif "sub total" in s or "subtotal" in s:
                    col_map["subtotal"] = c_i
                elif "mix" in s:
                    col_map["mix"] = c_i
            break

    if header_row_idx is None:
        header_row_idx = 4

    for r_i in range(0, min(header_row_idx + 1, len(rows))):
        row = rows[r_i]
        for c_i, cell in enumerate(row):
            s = _clean_cell_str(cell)
            if not s:
                continue
            s_low = s.lower()

            if s_low.startswith("proyecto"):
                inline_val = re.sub(r"^proyecto\s*:?\s*", "", s, flags=re.IGNORECASE).strip()
                if inline_val:
                    resultado["proyecto"] = inline_val
                else:
                    for next_c in range(c_i + 1, len(row)):
                        cand = _clean_cell_str(row[next_c])
                        if not cand:
                            continue
                        if re.match(r"^\d{4}-\d{2}-\d{2}", cand) or re.match(r"^\d{2}/\d{2}/\d{4}", cand):
                            resultado["fecha"] = cand.split(" ")[0]
                        elif not cand.lower().startswith("tc") and not resultado["proyecto"]:
                            resultado["proyecto"] = cand

            if re.match(r"^\d{4}-\d{2}-\d{2}", s) or re.match(r"^\d{2}/\d{2}/\d{4}", s):
                if not resultado["fecha"]:
                    resultado["fecha"] = s.split(" ")[0]

            if s_low.startswith("cliente"):
                inline_cli = re.sub(r"^cliente\s*:?\s*", "", s, flags=re.IGNORECASE).strip()
                if inline_cli:
                    resultado["cliente"] = inline_cli
                else:
                    for next_c in range(c_i + 1, len(row)):
                        cand = _clean_cell_str(row[next_c])
                        if not cand:
                            continue
                        if cand.lower().startswith("tc"):
                            break
                        if not resultado["cliente"]:
                            resultado["cliente"] = cand
                            break

            if s_low in ["tc", "tc:", "t.c.", "t.c:", "tipo de cambio", "tipo de cambio:"]:
                for next_c in range(c_i + 1, len(row)):
                    tc_val, _ = _parse_number_or_pct(row[next_c], default=0.0)
                    if tc_val > 0:
                        resultado["tc"] = tc_val
                        break

    tc = resultado["tc"] if resultado["tc"] > 0 else 3.42

    partidas_dict = {}
    partidas_order = []
    in_footer = False

    raw_subtotal_casco = 0.0
    raw_gg_factor_or_val = None
    raw_gg_is_pct = False
    raw_gg_amount = 0.0

    raw_util_factor_or_val = None
    raw_util_is_pct = False
    raw_util_amount = 0.0

    raw_igv_factor_or_val = None
    raw_igv_is_pct = False
    raw_igv_amount = 0.0

    raw_total_soles = 0.0

    def _get_col(row_list, key):
        idx = col_map.get(key)
        if idx is not None and 0 <= idx < len(row_list):
            return row_list[idx]
        return None

    for r_i in range(header_row_idx + 1, len(rows)):
        row = rows[r_i]
        if not any(_clean_cell_str(x) for x in row):
            continue

        row_joined = " ".join(_clean_cell_str(x) for x in row).lower()
        if "numeros de cuenta" in row_joined or "números de cuenta" in row_joined:
            break
        if "adelanto 1" in row_joined and "semana 1" in row_joined:
            break
        if "trabajos contratistas" in row_joined:
            break

        raw_code = _get_col(row, "codigo")
        desc = _clean_cell_str(_get_col(row, "descripcion"))
        moneda = _clean_cell_str(_get_col(row, "moneda")) or "Soles"
        unidad = _clean_cell_str(_get_col(row, "unidad")) or "glb"

        desc_low = desc.lower()

        if any(
            kw in desc_low
            for kw in [
                "sub total",
                "subtotal",
                "gastos generales",
                "gestión y supervisión",
                "gestion y supervision",
                "supervisión de obra",
                "supervision de obra",
                "utilidad",
                "honorario",
                "igv",
                "total (usd)",
            ]
        ) or desc_low == "total":
            in_footer = True
            val_met, met_is_pct = _parse_number_or_pct(_get_col(row, "metrado"), 0.0)
            val_vu, vu_is_pct = _parse_number_or_pct(_get_col(row, "valor_unitario"), 0.0)
            val_vp, _ = _parse_number_or_pct(_get_col(row, "valor_parcial"), 0.0)
            val_st, _ = _parse_number_or_pct(_get_col(row, "subtotal"), 0.0)

            best_amount = val_st if val_st > 0 else val_vp
            factor_cand = val_met if val_met > 0 else val_vu
            is_pct_cand = met_is_pct or vu_is_pct or (0 < factor_cand <= 1.0)

            if "sub total" in desc_low or "subtotal" in desc_low:
                if raw_subtotal_casco == 0.0 and best_amount > 0:
                    raw_subtotal_casco = best_amount
            elif "gastos generales" in desc_low:
                if best_amount > 0:
                    raw_gg_amount = best_amount
                if factor_cand > 0:
                    raw_gg_factor_or_val = factor_cand
                    raw_gg_is_pct = is_pct_cand
            elif any(k in desc_low for k in ["gestión", "gestion", "supervisión", "supervision", "utilidad", "honorario"]):
                if best_amount > 0:
                    raw_util_amount = best_amount
                if factor_cand > 0:
                    raw_util_factor_or_val = factor_cand
                    raw_util_is_pct = is_pct_cand
            elif desc_low == "igv" or desc_low.startswith("igv "):
                if best_amount > 0:
                    raw_igv_amount = best_amount
                if factor_cand > 0:
                    raw_igv_factor_or_val = factor_cand
                    raw_igv_is_pct = is_pct_cand
            elif desc_low == "total" and not any(d in moneda.lower() for d in ["dolar", "dólar", "usd"]):
                if raw_total_soles == 0.0 and best_amount > 0:
                    raw_total_soles = best_amount
            continue

        if in_footer:
            continue

        item_type, p_code, t_code = _classify_item_code(raw_code)

        if item_type == "partida":
            st_val, _ = _parse_number_or_pct(_get_col(row, "subtotal"), 0.0)
            vp_val, _ = _parse_number_or_pct(_get_col(row, "valor_parcial"), 0.0)
            sub_directo = st_val if st_val > 0 else vp_val

            if p_code not in partidas_dict:
                partidas_dict[p_code] = {
                    "codigo": p_code,
                    "nombre": desc,
                    "subtotal_celda": sub_directo,
                    "tareas": [],
                }
                partidas_order.append(p_code)
            else:
                if desc and not partidas_dict[p_code]["nombre"]:
                    partidas_dict[p_code]["nombre"] = desc
                if sub_directo > 0:
                    partidas_dict[p_code]["subtotal_celda"] = sub_directo

        elif item_type == "tarea":
            if p_code not in partidas_dict:
                partidas_dict[p_code] = {
                    "codigo": p_code,
                    "nombre": "",
                    "subtotal_celda": 0.0,
                    "tareas": [],
                }
                partidas_order.append(p_code)

            met_val, _ = _parse_number_or_pct(_get_col(row, "metrado"), 0.0)
            vu_val, _ = _parse_number_or_pct(_get_col(row, "valor_unitario"), 0.0)
            vp_val, _ = _parse_number_or_pct(_get_col(row, "valor_parcial"), 0.0)
            st_val, _ = _parse_number_or_pct(_get_col(row, "subtotal"), 0.0)

            has_content = bool(desc) or vu_val > 0 or vp_val > 0 or st_val > 0
            if not has_content:
                continue

            if met_val <= 0:
                met_val = 1.0

            if vp_val > 0:
                parcial_calc = vp_val
            elif vu_val > 0:
                parcial_calc = round(met_val * vu_val, 2)
            elif st_val > 0:
                parcial_calc = st_val
            else:
                parcial_calc = 0.0

            if vu_val <= 0 and parcial_calc > 0 and met_val > 0:
                vu_val = round(parcial_calc / met_val, 2)

            es_dolares = any(d in moneda.lower() for d in ["dolar", "dólar", "usd", "$"])
            factor_moneda = tc if es_dolares else 1.0
            presupuesto_soles = round(parcial_calc * factor_moneda, 2)
            vu_soles = round(vu_val * factor_moneda, 2)

            partidas_dict[p_code]["tareas"].append({
                "codigo": t_code,
                "descripcion": desc if desc else f"Tarea {t_code}",
                "moneda": moneda,
                "unidad": unidad,
                "metrado": met_val,
                "valor_unitario": vu_val,
                "valor_unitario_soles": vu_soles,
                "valor_parcial": parcial_calc,
                "presupuesto_soles": presupuesto_soles,
            })

    partidas_final = []
    suma_partidas = 0.0

    for p_code in partidas_order:
        p_info = partidas_dict[p_code]
        suma_tareas = round(sum(t["presupuesto_soles"] for t in p_info["tareas"]), 2)
        monto_pactado = suma_tareas if suma_tareas > 0 else round(p_info["subtotal_celda"], 2)

        nombre_base = p_info["nombre"].strip()
        if not nombre_base:
            nombre_base = f"Partida {p_code}"

        p_num_fmt = f"{int(p_code):02d}" if p_code.isdigit() else p_code
        if re.match(r"^\d+[\.\-\s]", nombre_base):
            nombre_completo = nombre_base
        else:
            nombre_completo = f"{p_num_fmt}. {nombre_base}"

        partidas_final.append({
            "codigo": p_code,
            "nombre": nombre_base,
            "nombre_completo": nombre_completo,
            "monto_pactado": monto_pactado,
            "mix_pct": 0.0,
            "tareas": p_info["tareas"],
        })
        suma_partidas += monto_pactado

    subtotal_casco = round(suma_partidas if suma_partidas > 0 else raw_subtotal_casco, 2)
    resultado["subtotal_casco"] = subtotal_casco

    for p_item in partidas_final:
        if subtotal_casco > 0:
            p_item["mix_pct"] = round((p_item["monto_pactado"] / subtotal_casco) * 100.0, 2)
    resultado["partidas"] = partidas_final

    if raw_gg_amount > 0:
        gg_monto = redondear_arriba_10(raw_gg_amount)
        gg_pct = round((raw_gg_factor_or_val * 100.0) if (raw_gg_factor_or_val and raw_gg_is_pct) else ((gg_monto / subtotal_casco * 100.0) if subtotal_casco > 0 else 0.0), 2)
    elif raw_gg_factor_or_val is not None and raw_gg_factor_or_val > 0:
        if raw_gg_is_pct:
            gg_pct = round(raw_gg_factor_or_val * 100.0, 2)
            gg_monto = redondear_arriba_10(subtotal_casco * raw_gg_factor_or_val)
        else:
            gg_monto = redondear_arriba_10(raw_gg_factor_or_val)
            gg_pct = round((gg_monto / subtotal_casco * 100.0) if subtotal_casco > 0 else 0.0, 2)
    else:
        gg_pct = 0.0
        gg_monto = 0.0

    resultado["gg_pct"] = gg_pct
    resultado["gg_monto"] = gg_monto

    if raw_util_amount > 0:
        util_monto = redondear_arriba_100(raw_util_amount)
        util_pct = round((raw_util_factor_or_val * 100.0) if (raw_util_factor_or_val and raw_util_is_pct) else ((util_monto / subtotal_casco * 100.0) if subtotal_casco > 0 else 0.0), 2)
    elif raw_util_factor_or_val is not None and raw_util_factor_or_val > 0:
        if raw_util_is_pct:
            util_pct = round(raw_util_factor_or_val * 100.0, 2)
            util_monto = redondear_arriba_100(subtotal_casco * raw_util_factor_or_val)
        else:
            util_monto = redondear_arriba_100(raw_util_factor_or_val)
            util_pct = round((util_monto / subtotal_casco * 100.0) if subtotal_casco > 0 else 0.0, 2)
    else:
        util_pct = 0.0
        util_monto = 0.0

    resultado["utilidades_pct"] = util_pct
    resultado["utilidades_monto"] = util_monto

    base_imponible = subtotal_casco + gg_monto + util_monto
    if raw_igv_amount > 0:
        igv_monto = round(raw_igv_amount, 2)
        igv_pct = round((raw_igv_factor_or_val * 100.0) if (raw_igv_factor_or_val and raw_igv_is_pct) else ((igv_monto / base_imponible * 100.0) if base_imponible > 0 else 0.0), 2)
    elif raw_igv_factor_or_val is not None and raw_igv_factor_or_val > 0:
        if raw_igv_is_pct:
            igv_pct = round(raw_igv_factor_or_val * 100.0, 2)
            igv_monto = round(base_imponible * raw_igv_factor_or_val, 2)
        else:
            igv_monto = round(raw_igv_factor_or_val, 2)
            igv_pct = round((igv_monto / base_imponible * 100.0) if base_imponible > 0 else 0.0, 2)
    else:
        igv_pct = 0.0
        igv_monto = 0.0

    resultado["igv_pct"] = igv_pct
    resultado["igv_monto"] = igv_monto

    total_calc = round(subtotal_casco + gg_monto + util_monto + igv_monto, 2)
    presupuesto_total = total_calc if total_calc > 0 else round(raw_total_soles, 2)
    resultado["presupuesto_total"] = presupuesto_total
    resultado["presupuesto_dolares"] = round(presupuesto_total / tc, 2) if tc > 0 else 0.0

    # Si el parser estándar no encontró tareas (ej. tabla flexible de Gemini Spark / Excel con columnas Partida, Tarea, Metraje, P.U.), usar parser flexible
    total_tareas_std = sum(len(p.get("tareas", [])) for p in resultado.get("partidas", []))
    if total_tareas_std == 0 and subtotal_casco == 0.0:
        flex_res = parsear_presupuesto_flexible(rows, tc_default=tc)
        if flex_res and (flex_res.get("partidas") or flex_res.get("subtotal_casco", 0.0) > 0):
            return flex_res

    return resultado


def parsear_presupuesto_flexible(rows_or_input, tc_default: float = 3.42) -> dict:
    """
    Parser flexible para importar presupuestos generados en Gemini Spark, Excel o CSV
    que contengan columnas equivalentes a Partida, Ítem, Tarea/Sub-partida/Descripción,
    Unidad, Metrado/Metraje/Cantidad, Valor Unitario/P.U., Valor Parcial/Subtotal,
    Moneda, Especialidad y Contratista.
    """
    rows = rows_or_input if isinstance(rows_or_input, list) else _rows_from_input(rows_or_input)
    resultado = {
        "proyecto": "",
        "cliente": "",
        "fecha": "",
        "tc": tc_default,
        "partidas": [],
        "subtotal_casco": 0.0,
        "gg_pct": 2.0,
        "gg_monto": 0.0,
        "utilidades_pct": 15.0,
        "utilidades_monto": 0.0,
        "igv_pct": 0.0,
        "igv_monto": 0.0,
        "presupuesto_total": 0.0,
        "presupuesto_dolares": 0.0,
    }
    if not rows:
        return resultado

    header_idx = None
    cmap = {}
    for r_i, row in enumerate(rows[:25]):
        r_low = [_clean_cell_str(c).lower() for c in row]
        score = sum(
            1
            for s in r_low
            if any(
                k in s
                for k in [
                    "partida", "categor", "capítulo", "capitulo", "item", "ítem", "cod", "cód",
                    "tarea", "subpartida", "sub-partida", "descrip", "actividad",
                    "unidad", "und", "u.m", "metrado", "metraje", "cant",
                    "unitario", "p.u", "precio", "costo", "parcial", "total", "importe",
                ]
            )
        )
        if score >= 2:
            header_idx = r_i
            for c_i, s in enumerate(r_low):
                if not s:
                    continue
                if any(k in s for k in ["subpartida", "sub-partida", "tarea", "descrip", "actividad", "concepto"]) and "descripcion" not in cmap:
                    cmap["descripcion"] = c_i
                elif any(k in s for k in ["partida", "categor", "capítulo", "capitulo", "fase"]) and "partida" not in cmap:
                    cmap["partida"] = c_i
                elif any(k in s for k in ["item", "ítem", "cod", "cód", "n°", "nro"]) and "codigo" not in cmap:
                    cmap["codigo"] = c_i
                elif any(k in s for k in ["moneda", "divisa"]) and "moneda" not in cmap:
                    cmap["moneda"] = c_i
                elif any(k in s for k in ["unidad", "und", "u.m", "medida"]) and "unidad" not in cmap:
                    cmap["unidad"] = c_i
                elif any(k in s for k in ["metrado", "metraje", "cantidad", "cant"]) and "metrado" not in cmap:
                    cmap["metrado"] = c_i
                elif any(k in s for k in ["unitario", "p.u", "pu", "precio unit", "v. unit", "v unit"]) and "valor_unitario" not in cmap:
                    cmap["valor_unitario"] = c_i
                elif any(k in s for k in ["parcial", "subtotal", "sub total", "importe", "total", "monto"]) and "valor_parcial" not in cmap:
                    cmap["valor_parcial"] = c_i
                elif "especialidad" in s and "especialidad" not in cmap:
                    cmap["especialidad"] = c_i
                elif any(k in s for k in ["contratista", "proveedor", "responsable"]) and "contratista" not in cmap:
                    cmap["contratista"] = c_i
            break

    if header_idx is None:
        return resultado

    # Si no hay columna "descripcion" pero sí "partida", usar "partida" como descripción jerárquica
    if "descripcion" not in cmap and "partida" in cmap:
        cmap["descripcion"] = cmap.pop("partida")

    def _cell(r, k):
        idx = cmap.get(k)
        return r[idx] if (idx is not None and 0 <= idx < len(r)) else None

    partidas_map = {}
    partidas_order = []
    current_p_code = "1"
    current_p_name = "Partida General"

    def _ensure_partida(p_key: str, p_nom: str):
        if p_key not in partidas_map:
            partidas_map[p_key] = {
                "codigo": str(len(partidas_order) + 1),
                "nombre": limpiar_nombre_partida(p_nom, len(partidas_order) + 1),
                "tareas": [],
            }
            partidas_order.append(p_key)

    for r_i in range(header_idx + 1, len(rows)):
        row = rows[r_i]
        if not any(_clean_cell_str(x) for x in row):
            continue

        raw_cod = _clean_cell_str(_cell(row, "codigo"))
        raw_part = _clean_cell_str(_cell(row, "partida"))
        raw_desc = _clean_cell_str(_cell(row, "descripcion"))
        raw_mon = _clean_cell_str(_cell(row, "moneda")) or "Soles"
        raw_und = _clean_cell_str(_cell(row, "unidad")) or "glb"
        raw_esp = _clean_cell_str(_cell(row, "especialidad")) or "Ninguna"
        raw_con = _clean_cell_str(_cell(row, "contratista"))

        met_v, _ = _parse_number_or_pct(_cell(row, "metrado"), 0.0)
        vu_v, _ = _parse_number_or_pct(_cell(row, "valor_unitario"), 0.0)
        vp_v, _ = _parse_number_or_pct(_cell(row, "valor_parcial"), 0.0)

        texto_comb = f"{raw_part} {raw_desc}".strip().lower()
        if any(
            kw in texto_comb
            for kw in [
                "sub total", "subtotal", "gastos generales", "gestión y supervisión",
                "gestion y supervision", "utilidad", "igv", "total general", "total presupuesto",
            ]
        ) or raw_desc.lower() == "total":
            if "gastos generales" in texto_comb and vp_v > 0:
                resultado["gg_monto"] = vp_v
            elif any(k in texto_comb for k in ["gestión", "gestion", "supervisión", "supervision", "utilidad"]) and vp_v > 0:
                resultado["utilidades_monto"] = vp_v
            elif "igv" in texto_comb and vp_v > 0:
                resultado["igv_monto"] = vp_v
                resultado["igv_pct"] = 18.0
            continue

        # Caso 1: Existe columna explícita de Partida en cada fila
        if raw_part and raw_desc:
            p_key = raw_part.upper()
            _ensure_partida(p_key, raw_part)
            current_p_code = p_key
        else:
            # Caso 2: Tabla jerárquica donde una fila sin metrado ni precio unitario es cabecera de Partida
            it_type, p_c, _ = _classify_item_code(raw_cod)
            es_cabecera_partida = (
                it_type == "partida"
                or (bool(raw_desc) and met_v == 0.0 and vu_v == 0.0 and (vp_v == 0.0 or not _clean_cell_str(_cell(row, "unidad"))))
            )
            if es_cabecera_partida and raw_desc:
                p_key = raw_desc.upper()
                _ensure_partida(p_key, raw_desc)
                current_p_code = p_key
                continue

        if not raw_desc and vu_v == 0.0 and vp_v == 0.0:
            continue

        _ensure_partida(current_p_code, current_p_name)
        p_obj = partidas_map[current_p_code]
        p_num_str = p_obj["codigo"]
        t_idx = len(p_obj["tareas"]) + 1
        t_code = f"{p_num_str}.{t_idx:02d}"

        if met_v <= 0:
            met_v = 1.0
        if vp_v > 0 and vu_v <= 0:
            vu_v = round(vp_v / met_v, 2)
        parcial_calc = round(met_v * vu_v, 2) if vu_v > 0 else round(vp_v, 2)

        es_usd = any(d in raw_mon.lower() for d in ["dolar", "dólar", "usd", "$"])
        factor_tc = tc_default if es_usd else 1.0
        ppto_soles = round(parcial_calc * factor_tc, 2)

        p_obj["tareas"].append({
            "codigo": t_code,
            "descripcion": extraer_descripcion_tarea(raw_desc, t_code) or f"Tarea {t_code}",
            "moneda": "Dólares" if es_usd else "Soles",
            "unidad": raw_und,
            "metrado": met_v,
            "valor_unitario": vu_v,
            "valor_unitario_soles": round(vu_v * factor_tc, 2),
            "valor_parcial": parcial_calc,
            "presupuesto_soles": ppto_soles,
            "especialidad": raw_esp,
            "contratista": raw_con,
        })

    partidas_final = []
    subtotal_casco = 0.0
    for idx_p, p_key in enumerate(partidas_order, start=1):
        p_obj = partidas_map[p_key]
        if not p_obj["tareas"]:
            continue
        p_code_str = str(idx_p)
        for idx_t, t_it in enumerate(p_obj["tareas"], start=1):
            t_it["codigo"] = f"{p_code_str}.{idx_t:02d}"
        monto_p = round(sum(t["presupuesto_soles"] for t in p_obj["tareas"]), 2)
        nom_limpio = limpiar_nombre_partida(p_obj["nombre"], idx_p)
        nom_comp = formatear_nombre_partida(idx_p, nom_limpio)
        partidas_final.append({
            "codigo": p_code_str,
            "nombre": nom_limpio,
            "nombre_completo": nom_comp,
            "monto_pactado": monto_p,
            "mix_pct": 0.0,
            "tareas": p_obj["tareas"],
        })
        subtotal_casco += monto_p

    subtotal_casco = round(subtotal_casco, 2)
    for p_it in partidas_final:
        p_it["mix_pct"] = round((p_it["monto_pactado"] / subtotal_casco) * 100.0, 2) if subtotal_casco > 0 else 0.0

    resultado["partidas"] = partidas_final
    resultado["subtotal_casco"] = subtotal_casco
    if resultado["gg_monto"] <= 0:
        calc_gg_flex = subtotal_casco * (resultado["gg_pct"] / 100.0)
        resultado["gg_monto"] = redondear_arriba_10(calc_gg_flex)
        resultado["gg_pct"] = round((resultado["gg_monto"] / subtotal_casco * 100.0) if subtotal_casco > 0 else resultado["gg_pct"], 2)
    else:
        resultado["gg_monto"] = redondear_arriba_10(resultado["gg_monto"])
        resultado["gg_pct"] = round((resultado["gg_monto"] / subtotal_casco * 100.0) if subtotal_casco > 0 else 2.0, 2)

    if resultado["utilidades_monto"] <= 0:
        calc_util_flex = subtotal_casco * (resultado["utilidades_pct"] / 100.0)
        resultado["utilidades_monto"] = redondear_arriba_100(calc_util_flex)
        resultado["utilidades_pct"] = round((resultado["utilidades_monto"] / subtotal_casco * 100.0) if subtotal_casco > 0 else resultado["utilidades_pct"], 2)
    else:
        resultado["utilidades_monto"] = redondear_arriba_100(resultado["utilidades_monto"])
        resultado["utilidades_pct"] = round((resultado["utilidades_monto"] / subtotal_casco * 100.0) if subtotal_casco > 0 else 15.0, 2)

    tot_soles = round(subtotal_casco + resultado["gg_monto"] + resultado["utilidades_monto"] + resultado["igv_monto"], 2)
    resultado["presupuesto_total"] = tot_soles
    resultado["presupuesto_dolares"] = round(tot_soles / tc_default, 2) if tc_default > 0 else 0.0
    return resultado


def aplicar_cotizacion_a_proyecto(
    proy_name: str,
    parsed_data: dict,
    modo: str = "reemplazar",
    actualizar_ficha: bool = True,
    nombres_partidas_override: dict = None,
    incluir_partidas_vacias: bool = False,
) -> dict:
    """
    Importa los datos leídos por `parsear_cotizacion_venti` en `st.session_state` para el proyecto `proy_name`.
    """
    guardar_snapshot_presupuesto(proy_name)
    if "proyectos" not in st.session_state:
        st.session_state.proyectos = pd.DataFrame(columns=COLS_PROY)
    if "p_info" not in st.session_state:
        st.session_state.p_info = pd.DataFrame(columns=COLS_P_INFO)
    if "tasks" not in st.session_state:
        st.session_state.tasks = pd.DataFrame(columns=COLS_TASKS)
    if "partidas" not in st.session_state:
        st.session_state.partidas = pd.DataFrame(columns=["Proyecto", "Nombre_Partida"])

    for col in COLS_PROY:
        if col not in st.session_state.proyectos.columns:
            st.session_state.proyectos[col] = ""
    for col in COLS_TASKS:
        if col not in st.session_state.tasks.columns:
            st.session_state.tasks[col] = ""
    for col in COLS_P_INFO:
        if col not in st.session_state.p_info.columns:
            st.session_state.p_info[col] = ""

    mask_proy = st.session_state.proyectos["Nombre"] == proy_name
    if actualizar_ficha and mask_proy.any():
        idx_p = st.session_state.proyectos[mask_proy].index[0]
        if parsed_data.get("presupuesto_total", 0.0) > 0:
            st.session_state.proyectos.at[idx_p, "Presupuesto"] = float(parsed_data["presupuesto_total"])
        if parsed_data.get("utilidades_monto", 0.0) >= 0:
            st.session_state.proyectos.at[idx_p, "Utilidades"] = float(parsed_data["utilidades_monto"])
        if parsed_data.get("gg_monto", 0.0) >= 0:
            st.session_state.proyectos.at[idx_p, "GG_Presupuesto"] = float(parsed_data["gg_monto"])
        if parsed_data.get("tc", 0.0) > 0:
            st.session_state.proyectos.at[idx_p, "TC_Presupuesto"] = float(parsed_data["tc"])
        if parsed_data.get("gg_pct", 0.0) >= 0:
            st.session_state.proyectos.at[idx_p, "GG_Pct"] = float(parsed_data["gg_pct"])
        if parsed_data.get("utilidades_pct", 0.0) >= 0:
            st.session_state.proyectos.at[idx_p, "Utilidades_Pct"] = float(parsed_data["utilidades_pct"])
        if "modo_gg" in parsed_data:
            st.session_state.proyectos.at[idx_p, "Modo_GG"] = str(parsed_data["modo_gg"])
        if "monto_fijo_gg" in parsed_data:
            st.session_state.proyectos.at[idx_p, "Monto_Fijo_GG"] = float(parsed_data["monto_fijo_gg"])
        if "modo_util" in parsed_data:
            st.session_state.proyectos.at[idx_p, "Modo_Utilidades"] = str(parsed_data["modo_util"])
        if "monto_fijo_util" in parsed_data:
            st.session_state.proyectos.at[idx_p, "Monto_Fijo_Utilidades"] = float(parsed_data["monto_fijo_util"])
        st.session_state.proyectos.at[idx_p, "Incluir_IGV"] = bool(parsed_data.get("igv_pct", 0.0) > 0)

        cli_parsed = safe_str(parsed_data.get("cliente", ""))
        if cli_parsed and not safe_str(st.session_state.proyectos.at[idx_p, "Cliente"]):
            st.session_state.proyectos.at[idx_p, "Cliente"] = cli_parsed

    if modo == "reemplazar":
        st.session_state.partidas = st.session_state.partidas[st.session_state.partidas["Proyecto"] != proy_name].reset_index(drop=True)
        st.session_state.p_info = st.session_state.p_info[st.session_state.p_info["Proyecto"] != proy_name].reset_index(drop=True)
        st.session_state.tasks = st.session_state.tasks[st.session_state.tasks["Proyecto"] != proy_name].reset_index(drop=True)

    nuevas_partidas_rows = []
    nuevas_pinfo_rows = []
    nuevas_tasks_rows = []

    partidas_importadas = 0
    tareas_importadas = 0

    for p in parsed_data.get("partidas", []):
        p_code = p["codigo"]
        tiene_contenido = bool(p.get("tareas")) or p.get("monto_pactado", 0.0) > 0 or bool(p.get("nombre", "").strip())
        if not incluir_partidas_vacias and not tiene_contenido:
            continue

        p_nom_final = (
            nombres_partidas_override.get(p_code, p["nombre_completo"]).strip()
            if nombres_partidas_override and p_code in nombres_partidas_override
            else p["nombre_completo"]
        )
        monto_pactado = float(p.get("monto_pactado", 0.0))

        mask_pi_exist = (st.session_state.p_info["Proyecto"] == proy_name) & (st.session_state.p_info["Nombre_Partida"] == p_nom_final)
        if mask_pi_exist.any():
            st.session_state.p_info.loc[mask_pi_exist, "Monto_Pactado"] = monto_pactado
        else:
            nuevas_pinfo_rows.append({
                "Proyecto": proy_name,
                "Nombre_Partida": p_nom_final,
                "Monto_Pactado": monto_pactado,
                "Nivel_Actual": "",
                "Contratista": "",
                "Inicio": "",
                "Fin": "",
                "Estado": "🟡 Planificado",
                "Etapa": "Construcción",
                "Nivel_Inicio": "",
                "Nivel_Fin": "",
            })

        mask_pa_exist = (st.session_state.partidas["Proyecto"] == proy_name) & (st.session_state.partidas["Nombre_Partida"] == p_nom_final)
        if not mask_pa_exist.any():
            nuevas_partidas_rows.append({
                "Proyecto": proy_name,
                "Nombre_Partida": p_nom_final,
            })

        partidas_importadas += 1

        for idx_t, t in enumerate(p.get("tareas", []), start=1):
            t_cod = safe_str(t.get("codigo", "")) or f"{p_code}.{idx_t:02d}"
            t_desc = safe_str(t.get("descripcion", ""))
            if t_cod and not t_desc.startswith(t_cod):
                nombre_tarea = f"{t_cod} {t_desc}".strip()
            else:
                nombre_tarea = t_desc or f"Tarea {t_cod}"

            moneda_t = safe_str(t.get("moneda", "Soles")) or "Soles"
            unidad_t = safe_str(t.get("unidad", "glb")) or "glb"
            metrado_t = float(safe_float(t.get("metrado", 1.0), 1.0))
            vu_orig_t = float(safe_float(t.get("valor_unitario", 0.0), 0.0))
            ppto_soles_t = float(safe_float(t.get("presupuesto_soles", 0.0), 0.0))

            notas_t = f"Ppto: S/ {ppto_soles_t:,.2f} ({metrado_t:g} {unidad_t} x {vu_orig_t:,.2f} {moneda_t})"

            nuevas_tasks_rows.append({
                "Proyecto": proy_name,
                "Partida": p_nom_final,
                "Orden": idx_t,
                "Codigo_Item": t_cod,
                "Tarea": nombre_tarea,
                "Especialidad": safe_str(t.get("especialidad", "Ninguna")).strip() or "Ninguna",
                "Contratista": safe_str(t.get("contratista", "")).strip(),
                "Estado": "🟡 Planificado",
                "Supervisado": "No",
                "Inicio": parsed_data.get("fecha") or datetime.date.today().strftime("%Y-%m-%d"),
                "Fin": "",
                "Causa_Cumplimiento": "",
                "Culpa_Contratista": False,
                "Permitir_Cruce": False,
                "Dependencia": "",
                "Notas": notas_t,
                "Moneda": moneda_t,
                "Unidad": unidad_t,
                "Metrado": metrado_t,
                "Valor_Unitario": vu_orig_t,
                "Presupuesto": ppto_soles_t,
            })
            tareas_importadas += 1

    if nuevas_partidas_rows:
        st.session_state.partidas = pd.concat(
            [st.session_state.partidas, pd.DataFrame(nuevas_partidas_rows)],
            ignore_index=True,
        ).drop_duplicates().reset_index(drop=True)

    if nuevas_pinfo_rows:
        df_new_pi = pd.DataFrame(nuevas_pinfo_rows)
        st.session_state.p_info = pd.concat(
            [st.session_state.p_info, df_new_pi[COLS_P_INFO]],
            ignore_index=True,
        )

    if nuevas_tasks_rows:
        df_new_t = pd.DataFrame(nuevas_tasks_rows)
        st.session_state.tasks = pd.concat(
            [st.session_state.tasks, df_new_t[COLS_TASKS]],
            ignore_index=True,
        )

    clear_base_dfs()

    return {
        "partidas_importadas": partidas_importadas,
        "tareas_importadas": tareas_importadas,
        "presupuesto_total": float(parsed_data.get("presupuesto_total", 0.0)),
        "utilidades_monto": float(parsed_data.get("utilidades_monto", 0.0)),
        "gg_monto": float(parsed_data.get("gg_monto", 0.0)),
    }


# ==========================================
# 4. CUADRO DE PAGOS SEMANALES POR PARTIDA / CONTRATISTA (06. COBROS & FINANZAS)
# ==========================================

PARTIDAS_PAGOS_DEFAULT = [
    "OBRAS CIVILES",
    "INSTALACIONES E&S",
    "CARPINTERÍA Y MOBILIARIO",
    "TABLERO DE PIEDRA",
    "PINTURA",
    "TRABAJOS PRELIMINARES",
    "VIDRIOS",
]


def obtener_columnas_pago_semanal(num_semanas: int = 5) -> list[str]:
    """Retorna la lista de columnas de desembolso: ['ADELANTO 1', 'SEMANA 1', ..., 'SEMANA N']."""
    n = max(1, min(30, int(num_semanas or 5)))
    return ["ADELANTO 1"] + [f"SEMANA {i}" for i in range(1, n + 1)]


def _normalizar_fila_pago(fila: dict, cols_pago: list[str], desc_default: str = "", partida_default: str = "") -> dict:
    """Garantiza que una fila del cuadro de pagos tenga todas las columnas numéricas, de texto y vínculo de Partida."""
    desc_val = safe_str(fila.get("Descripcion", desc_default)).strip() or desc_default
    part_val = safe_str(fila.get("Partida", partida_default)).strip() or partida_default or desc_val
    res = {
        "Partida": part_val,
        "Descripcion": desc_val,
        "Moneda": safe_str(fila.get("Moneda", "Soles")).strip() or "Soles",
        "TOTAL": round(safe_float(fila.get("TOTAL", 0.0), 0.0), 2),
    }
    for c in cols_pago:
        res[c] = round(safe_float(fila.get(c, 0.0), 0.0), 2)
    res["Notas"] = safe_str(fila.get("Notas", "")).strip()
    return res


def _fila_pertenece_a_partida(fila: dict, p_raw: str, p_clean: str, idx_p: int) -> bool:
    """Verifica si una fila del cuadro de pagos corresponde a la partida dada."""
    p_raw_u = safe_str(p_raw).strip().upper()
    p_clean_u = safe_str(p_clean).strip().upper()
    f_part_u = safe_str(fila.get("Partida", "")).strip().upper()
    f_desc_u = safe_str(fila.get("Descripcion", "")).strip().upper()

    if f_part_u and (
        f_part_u == p_raw_u
        or f_part_u == p_clean_u
        or limpiar_nombre_partida(f_part_u, idx_p).strip().upper() == p_clean_u
    ):
        return True
    if not f_part_u or f_part_u == f_desc_u:
        if (
            f_desc_u == p_raw_u
            or f_desc_u == p_clean_u
            or limpiar_nombre_partida(f_desc_u, idx_p).strip().upper() == p_clean_u
            or (p_clean_u and (f_desc_u.startswith(p_clean_u + " -") or f_desc_u.startswith(p_clean_u + " (") or f_desc_u.startswith(p_clean_u + " /")))
        ):
            return True
    return False


def obtener_cuadro_pagos_semanales_proyecto(proy_name: str, sincronizar_partidas: bool = False) -> dict:
    """
    Obtiene o inicializa la estructura del Cuadro de Pagos Semanales para el proyecto `proy_name`.
    Sincroniza automáticamente las partidas, contratistas y montos desde 03. PRESUPUESTO
    conservando todos los pagos semanales ya registrados y permitiendo múltiples filas por partida.
    """
    df_parts = st.session_state.get("partidas", pd.DataFrame())
    if not df_parts.empty and "Proyecto" in df_parts.columns and not df_parts[df_parts["Proyecto"] == proy_name].empty:
        recalcular_totales_presupuesto_proyecto(proy_name, sincronizar_pactado=False)

    df_p = st.session_state.get("proyectos", pd.DataFrame())
    proy_row = {}
    if not df_p.empty and "Nombre" in df_p.columns:
        m = df_p[df_p["Nombre"] == proy_name]
        if not m.empty:
            proy_row = m.iloc[0].to_dict()

    tc_proy = safe_float(proy_row.get("TC_Presupuesto", 3.60), 3.60)
    if tc_proy <= 0:
        tc_proy = 3.60
    util_proy = round(safe_float(proy_row.get("Utilidades", 0.0), 0.0), 2)
    gg_proy = round(safe_float(proy_row.get("GG_Presupuesto", 0.0), 0.0), 2)

    raw_json = safe_str(proy_row.get("Pagos_Semanales_JSON", "")).strip()
    data_guardada = None
    if raw_json:
        try:
            data_guardada = json.loads(raw_json)
        except Exception:
            data_guardada = None

    # Extraer partidas, montos pactados y contratistas actuales del proyecto
    df_pinfo = st.session_state.get("p_info", pd.DataFrame())
    df_tasks = st.session_state.get("tasks", pd.DataFrame())

    partidas_proy = []
    if not df_parts.empty and "Proyecto" in df_parts.columns:
        p_sub = df_parts[df_parts["Proyecto"] == proy_name]
        for idx_p, (_, r_p) in enumerate(p_sub.iterrows(), start=1):
            p_raw = safe_str(r_p.get("Nombre_Partida", f"Partida {idx_p}"))
            p_clean = limpiar_nombre_partida(p_raw, idx_p).upper()
            monto_p = 0.0
            contra_p = ""
            if not df_pinfo.empty and "Proyecto" in df_pinfo.columns:
                pi_m = df_pinfo[(df_pinfo["Proyecto"] == proy_name) & (df_pinfo["Nombre_Partida"] == p_raw)]
                if not pi_m.empty:
                    monto_p = safe_float(pi_m["Monto_Pactado"].iloc[0], 0.0)
                    contra_p = safe_str(pi_m["Contratista"].iloc[0]).strip() if "Contratista" in pi_m.columns else ""
            if not df_tasks.empty and "Proyecto" in df_tasks.columns:
                t_m = df_tasks[(df_tasks["Proyecto"] == proy_name) & (df_tasks["Partida"] == p_raw)]
                if not t_m.empty:
                    if monto_p <= 0 and "Presupuesto" in t_m.columns:
                        monto_p = float(pd.to_numeric(t_m["Presupuesto"], errors="coerce").fillna(0.0).sum())
                    if not contra_p and "Contratista" in t_m.columns:
                        cs = [safe_str(x).strip() for x in t_m["Contratista"].dropna().unique().tolist() if safe_str(x).strip()]
                        contra_p = ", ".join(cs)
            partidas_proy.append((idx_p, p_raw, p_clean, round(monto_p, 2), contra_p))

    tiene_ppto_activo = any(m > 0 for _, _, _, m, _ in partidas_proy)

    if data_guardada and isinstance(data_guardada, dict):
        num_sem = max(1, min(30, int(data_guardada.get("num_semanas", 5) or 5)))
        tc_val = safe_float(data_guardada.get("tc", tc_proy), tc_proy)
        if tc_val <= 0:
            tc_val = tc_proy
        cols_pago = obtener_columnas_pago_semanal(num_sem)
        cols_por_pagar = [c for c in data_guardada.get("cols_por_pagar", []) if c in cols_pago]

        filas_c = [
            _normalizar_fila_pago(f, cols_pago, f"Partida {i+1}")
            for i, f in enumerate(data_guardada.get("filas_contratistas", []))
        ]
        fila_ges = _normalizar_fila_pago(
            data_guardada.get("fila_gestion", {"TOTAL": util_proy}),
            cols_pago,
            "Gestión y supervisión de obras",
        )
        fila_gg = _normalizar_fila_pago(
            data_guardada.get("fila_gg", {"TOTAL": gg_proy}),
            cols_pago,
            "Gastos Generales",
        )

        if (sincronizar_partidas or tiene_ppto_activo) and partidas_proy:
            nuevas_filas_c = []
            indices_usados = set()

            for idx_pos, (idx_p, p_raw, p_nom_u, p_monto, p_contra) in enumerate(partidas_proy):
                match_idxs = [
                    i_f for i_f, f_cand in enumerate(filas_c)
                    if i_f not in indices_usados and _fila_pertenece_a_partida(f_cand, p_raw, p_nom_u, idx_p)
                ]
                if not match_idxs and idx_pos < len(filas_c) and idx_pos not in indices_usados:
                    cand_d = filas_c[idx_pos]["Descripcion"].strip().upper()
                    if cand_d.startswith("PARTIDA "):
                        match_idxs = [idx_pos]

                if len(match_idxs) == 1:
                    i_m = match_idxs[0]
                    indices_usados.add(i_m)
                    f_ex = dict(filas_c[i_m])
                    f_ex["Partida"] = p_nom_u
                    if f_ex["Descripcion"].strip().upper().startswith("PARTIDA "):
                        f_ex["Descripcion"] = p_nom_u
                    if p_monto > 0 and (sincronizar_partidas or safe_float(f_ex.get("TOTAL", 0.0), 0.0) == 0.0):
                        f_ex["TOTAL"] = p_monto
                    if p_contra and not safe_str(f_ex.get("Notas", "")).strip():
                        f_ex["Notas"] = f"Contratista: {p_contra}"
                    nuevas_filas_c.append(_normalizar_fila_pago(f_ex, cols_pago, p_nom_u, p_nom_u))
                elif len(match_idxs) > 1:
                    # Múltiples filas para la misma partida (ej. Eléctricas y Sanitarias dentro de Instalaciones)
                    for i_m in match_idxs:
                        indices_usados.add(i_m)
                        f_ex = dict(filas_c[i_m])
                        f_ex["Partida"] = p_nom_u
                        nuevas_filas_c.append(_normalizar_fila_pago(f_ex, cols_pago, f_ex.get("Descripcion", p_nom_u), p_nom_u))
                else:
                    nota_init = f"Contratista: {p_contra}" if p_contra else ""
                    nuevas_filas_c.append(
                        _normalizar_fila_pago(
                            {"Partida": p_nom_u, "Descripcion": p_nom_u, "TOTAL": p_monto, "Notas": nota_init},
                            cols_pago,
                            p_nom_u,
                            p_nom_u,
                        )
                    )

            # Conservar cualquier fila adicional creada por el usuario
            for i_old, f_old in enumerate(filas_c):
                if i_old not in indices_usados:
                    nuevas_filas_c.append(_normalizar_fila_pago(f_old, cols_pago, f_old.get("Descripcion", ""), f_old.get("Partida", "")))

            filas_c = nuevas_filas_c
            if util_proy > 0 or sincronizar_partidas:
                fila_ges["TOTAL"] = util_proy
            if gg_proy > 0 or sincronizar_partidas:
                fila_gg["TOTAL"] = gg_proy

        res_data = {
            "num_semanas": num_sem,
            "tc": tc_val,
            "cols_por_pagar": cols_por_pagar,
            "filas_contratistas": filas_c,
            "fila_gestion": fila_ges,
            "fila_gg": fila_gg,
        }
        if sincronizar_partidas:
            guardar_cuadro_pagos_semanales_proyecto(proy_name, res_data)
        return res_data

    # Inicialización por primera vez para este proyecto
    num_sem = 5
    cols_pago = obtener_columnas_pago_semanal(num_sem)
    if partidas_proy:
        filas_c = [
            _normalizar_fila_pago(
                {"Partida": p_n, "Descripcion": p_n, "TOTAL": p_m, "Notas": f"Contratista: {p_c}" if p_c else ""},
                cols_pago,
                p_n,
                p_n,
            )
            for _, _, p_n, p_m, p_c in partidas_proy
        ]
    else:
        filas_c = [
            _normalizar_fila_pago({"Partida": p_def, "Descripcion": p_def, "TOTAL": 0.0}, cols_pago, p_def, p_def)
            for p_def in PARTIDAS_PAGOS_DEFAULT
        ]

    res_init = {
        "num_semanas": num_sem,
        "tc": tc_proy,
        "cols_por_pagar": [],
        "filas_contratistas": filas_c,
        "fila_gestion": _normalizar_fila_pago({"TOTAL": util_proy}, cols_pago, "Gestión y supervisión de obras"),
        "fila_gg": _normalizar_fila_pago({"TOTAL": gg_proy}, cols_pago, "Gastos Generales"),
    }
    return res_init


def obtener_pagado_por_partida_proyecto(proy_name: str) -> dict[str, float]:
    """
    Calcula el monto total PAGADO de cada Partida del proyecto sumando:
    1. Los pagos registrados en el Cuadro de Pagos Semanales (06. COBROS & FINANZAS) para esa Partida
       (incluyendo cuando una Partida tiene múltiples filas de contratistas, ej. Eléctricas y Sanitarias).
    2. Los pagos manuales registrados en `st.session_state.pagos` para esa Partida.
    Retorna un diccionario `{Nombre_Partida_Completo: monto_pagado_total}`.
    """
    resultado_pagado = {}
    df_parts = st.session_state.get("partidas", pd.DataFrame())
    if df_parts.empty or "Proyecto" not in df_parts.columns:
        return resultado_pagado

    p_sub = df_parts[df_parts["Proyecto"] == proy_name]
    if p_sub.empty:
        return resultado_pagado

    data_cuadro = obtener_cuadro_pagos_semanales_proyecto(proy_name)
    res_calc = calcular_resumen_cuadro_pagos_semanales(data_cuadro)
    filas_c_calc = res_calc.get("filas_contratistas", [])
    tc = safe_float(res_calc.get("tc", 3.60), 3.60)
    if tc <= 0:
        tc = 3.60

    df_pagos_man = st.session_state.get("pagos", pd.DataFrame())

    for idx_p, (_, r_p) in enumerate(p_sub.iterrows(), start=1):
        p_raw = safe_str(r_p.get("Nombre_Partida", f"Partida {idx_p}"))
        p_clean = limpiar_nombre_partida(p_raw, idx_p).upper()

        pagado_cuadro = 0.0
        for f_c in filas_c_calc:
            if _fila_pertenece_a_partida(f_c, p_raw, p_clean, idx_p):
                monto_f = safe_float(f_c.get("PAGADO", 0.0), 0.0)
                es_usd = "dólar" in safe_str(f_c.get("Moneda", "")).lower() or "dolar" in safe_str(f_c.get("Moneda", "")).lower()
                pagado_cuadro += monto_f * (tc if es_usd else 1.0)

        pagado_manual = 0.0
        if not df_pagos_man.empty and "Proyecto" in df_pagos_man.columns and "Partida" in df_pagos_man.columns:
            m_man = df_pagos_man[(df_pagos_man["Proyecto"] == proy_name) & (df_pagos_man["Partida"] == p_raw)]
            if not m_man.empty:
                pagado_manual = float(pd.to_numeric(m_man.get("Monto", 0.0), errors="coerce").fillna(0.0).sum())

        resultado_pagado[p_raw] = round(pagado_cuadro + pagado_manual, 2)

    return resultado_pagado


def guardar_cuadro_pagos_semanales_proyecto(proy_name: str, data_cuadro: dict):
    """Guarda el estado del Cuadro de Pagos Semanales en `st.session_state.proyectos`."""
    if "proyectos" not in st.session_state or st.session_state.proyectos.empty:
        return
    if "Pagos_Semanales_JSON" not in st.session_state.proyectos.columns:
        st.session_state.proyectos["Pagos_Semanales_JSON"] = ""
    if st.session_state.proyectos["Pagos_Semanales_JSON"].dtype != object:
        st.session_state.proyectos["Pagos_Semanales_JSON"] = (
            st.session_state.proyectos["Pagos_Semanales_JSON"].fillna("").astype(str).replace({"nan": "", "None": ""}).astype(object)
        )

    mask = st.session_state.proyectos["Nombre"] == proy_name
    if mask.any():
        idx_p = st.session_state.proyectos[mask].index[0]
        st.session_state.proyectos.at[idx_p, "Pagos_Semanales_JSON"] = json.dumps(data_cuadro, ensure_ascii=False)


def calcular_resumen_cuadro_pagos_semanales(data_cuadro: dict) -> dict:
    """
    Calcula PAGADO, POR_PAGAR y SALDO de cada fila, así como las filas de cierre:
    SUB TOTAL (Contratistas), Gestión y supervisión de obras, Gastos Generales,
    TOTAL (Soles) y TOTAL (Dólares).
    """
    num_sem = max(1, min(30, int(data_cuadro.get("num_semanas", 5) or 5)))
    tc = safe_float(data_cuadro.get("tc", 3.60), 3.60)
    if tc <= 0:
        tc = 3.60
    cols_pago = obtener_columnas_pago_semanal(num_sem)
    cols_por_pagar = set(data_cuadro.get("cols_por_pagar", []))
    cols_pagadas = [c for c in cols_pago if c not in cols_por_pagar]

    def _enriquecer_fila(f_raw: dict, default_desc: str = "") -> dict:
        f = _normalizar_fila_pago(f_raw, cols_pago, default_desc)
        pagado = round(sum(safe_float(f.get(c, 0.0), 0.0) for c in cols_pagadas), 2)
        por_pagar = round(sum(safe_float(f.get(c, 0.0), 0.0) for c in cols_pago if c in cols_por_pagar), 2)
        saldo = round(f["TOTAL"] - pagado, 2)
        f["PAGADO"] = pagado
        f["POR_PAGAR"] = por_pagar
        f["SALDO"] = saldo
        return f

    filas_c_calc = [
        _enriquecer_fila(f, f"Partida {i+1}")
        for i, f in enumerate(data_cuadro.get("filas_contratistas", []))
    ]

    sub_total = {
        "Descripcion": "SUB TOTAL",
        "Moneda": "Soles",
        "TOTAL": round(sum(f["TOTAL"] for f in filas_c_calc), 2),
    }
    for c in cols_pago:
        sub_total[c] = round(sum(f[c] for f in filas_c_calc), 2)
    sub_total["PAGADO"] = round(sum(f["PAGADO"] for f in filas_c_calc), 2)
    sub_total["POR_PAGAR"] = round(sum(f["POR_PAGAR"] for f in filas_c_calc), 2)
    sub_total["SALDO"] = round(sum(f["SALDO"] for f in filas_c_calc), 2)
    sub_total["Notas"] = ""

    fila_ges_calc = _enriquecer_fila(data_cuadro.get("fila_gestion", {}), "Gestión y supervisión de obras")
    fila_gg_calc = _enriquecer_fila(data_cuadro.get("fila_gg", {}), "Gastos Generales")

    filas_principales = [sub_total, fila_ges_calc, fila_gg_calc]
    total_soles = {
        "Descripcion": "TOTAL",
        "Moneda": "Soles",
        "TOTAL": round(sum(r["TOTAL"] for r in filas_principales), 2),
    }
    for c in cols_pago:
        total_soles[c] = round(sum(r[c] for r in filas_principales), 2)
    total_soles["PAGADO"] = round(sum(r["PAGADO"] for r in filas_principales), 2)
    total_soles["POR_PAGAR"] = round(sum(r["POR_PAGAR"] for r in filas_principales), 2)
    total_soles["SALDO"] = round(sum(r["SALDO"] for r in filas_principales), 2)
    total_soles["Notas"] = ""

    total_dolares = {
        "Descripcion": "",
        "Moneda": "Dólares",
        "TOTAL": round(total_soles["TOTAL"] / tc, 2),
    }
    for c in cols_pago:
        total_dolares[c] = round(total_soles[c] / tc, 2)
    total_dolares["PAGADO"] = round(total_soles["PAGADO"] / tc, 2)
    total_dolares["POR_PAGAR"] = round(total_soles["POR_PAGAR"] / tc, 2)
    total_dolares["SALDO"] = round(total_soles["SALDO"] / tc, 2)
    total_dolares["Notas"] = f"TC: {tc:.2f}"

    return {
        "num_semanas": num_sem,
        "tc": tc,
        "cols_pago": cols_pago,
        "cols_por_pagar": list(cols_por_pagar),
        "filas_contratistas": filas_c_calc,
        "sub_total": sub_total,
        "fila_gestion": fila_ges_calc,
        "fila_gg": fila_gg_calc,
        "total_soles": total_soles,
        "total_dolares": total_dolares,
    }


def exportar_cuadro_pagos_semanales_excel(proy_name: str, data_cuadro: dict = None) -> bytes:
    """
    Exporta el Cuadro de Pagos Semanales por Partida/Contratista a un archivo Excel (.xlsx)
    con el diseño y fórmulas exactas para entregar al cliente.
    """
    if data_cuadro is None:
        data_cuadro = obtener_cuadro_pagos_semanales_proyecto(proy_name)
    res = calcular_resumen_cuadro_pagos_semanales(data_cuadro)

    df_p = st.session_state.get("proyectos", pd.DataFrame())
    cliente = ""
    if not df_p.empty and "Nombre" in df_p.columns:
        m = df_p[df_p["Nombre"] == proy_name]
        if not m.empty:
            cliente = safe_str(m.iloc[0].get("Cliente", ""))

    cols_pago = res["cols_pago"]
    cols_por_pagar = set(res["cols_por_pagar"])
    tc = res["tc"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control de Pagos"
    ws.views.sheetView[0].showGridLines = True

    font_family = "Calibri"
    f_title = Font(name=font_family, size=11, bold=True, color="262320")
    f_header = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    f_por_pagar = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    f_section = Font(name=font_family, size=10, bold=True, color="262320")
    f_row = Font(name=font_family, size=10, bold=False, color="262320")
    f_subtotal = Font(name=font_family, size=10, bold=True, color="262320")
    f_total = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    f_usd = Font(name=font_family, size=10, bold=True, color="262320")

    fill_header = PatternFill(start_color="262320", end_color="262320", fill_type="solid")
    fill_por_pagar = PatternFill(start_color="D98A2B", end_color="D98A2B", fill_type="solid")
    fill_pagado_tag = PatternFill(start_color="4A7C59", end_color="4A7C59", fill_type="solid")
    fill_section = PatternFill(start_color="EFECE6", end_color="EFECE6", fill_type="solid")
    fill_subtotal = PatternFill(start_color="DFD9CE", end_color="DFD9CE", fill_type="solid")
    fill_total = PatternFill(start_color="C85A32", end_color="C85A32", fill_type="solid")
    fill_usd = PatternFill(start_color="EFECE6", end_color="EFECE6", fill_type="solid")

    thin_side = Side(border_style="thin", color="D0C9BD")
    med_bottom = Side(border_style="medium", color="262320")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_closing = Border(left=thin_side, right=thin_side, top=thin_side, bottom=med_bottom)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # Encabezado de Proyecto y Cliente
    ws["B1"] = "Proyecto:"
    ws["B1"].font = f_title
    ws["C1"] = proy_name
    ws["C1"].font = f_title

    ws["B2"] = "Cliente:"
    ws["B2"].font = f_title
    ws["C2"] = cliente
    ws["C2"].font = f_title

    ws["E2"] = "TC:"
    ws["E2"].font = f_title
    ws["E2"].alignment = align_right
    ws["F2"] = float(tc)
    ws["F2"].font = f_title
    ws["F2"].number_format = "0.00"
    ws["F2"].alignment = align_center

    # Columnas:
    # B: Descripción (col 2)
    # C: Moneda (col 3)
    # D: TOTAL (col 4)
    # E .. E+len(cols_pago)-1: ADELANTO 1, SEMANA 1..N (cols 5 .. 4+len(cols_pago))
    # Col PAGADO: 5 + len(cols_pago)
    # Col SALDO: 6 + len(cols_pago)
    # Col Notas: 7 + len(cols_pago)
    col_idx_pagado = 5 + len(cols_pago)
    col_idx_saldo = 6 + len(cols_pago)
    col_idx_notas = 7 + len(cols_pago)

    let_pagado = get_column_letter(col_idx_pagado)
    let_saldo = get_column_letter(col_idx_saldo)
    let_notas = get_column_letter(col_idx_notas)

    # Fila 4: Cabecera de la tabla
    headers = [
        (2, "Descripción"),
        (3, "Moneda"),
        (4, "TOTAL"),
    ]
    for i_c, c_name in enumerate(cols_pago):
        headers.append((5 + i_c, c_name))
    headers.append((col_idx_pagado, "PAGADO"))
    headers.append((col_idx_saldo, "SALDO"))
    headers.append((col_idx_notas, "Notas"))

    ws.row_dimensions[4].height = 24
    for c_i, h_text in headers:
        c_let = get_column_letter(c_i)
        cell = ws[f"{c_let}4"]
        cell.value = h_text
        cell.font = f_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    # Fila 5: Indicadores "Por Pagar" sobre las semanas pendientes
    ws.row_dimensions[5].height = 18
    for i_c, c_name in enumerate(cols_pago):
        c_let = get_column_letter(5 + i_c)
        cell = ws[f"{c_let}5"]
        if c_name in cols_por_pagar:
            cell.value = "Por Pagar"
            cell.font = f_por_pagar
            cell.fill = fill_por_pagar
            cell.alignment = align_center
            cell.border = border_cell

    # Fila 6: TRABAJOS CONTRATISTAS
    ws["B6"] = "TRABAJOS CONTRATISTAS"
    ws["B6"].font = f_section
    for c_i in range(2, col_idx_notas + 1):
        c_let = get_column_letter(c_i)
        ws[f"{c_let}6"].fill = fill_section
        ws[f"{c_let}6"].border = border_cell

    # Construir fórmula de PAGADO para una fila dada (sumando solo columnas que NO están en Por Pagar)
    paid_col_letters = [
        get_column_letter(5 + i_c)
        for i_c, c_name in enumerate(cols_pago)
        if c_name not in cols_por_pagar
    ]

    def _formula_pagado(row_num: int) -> str:
        if not paid_col_letters:
            return "=0"
        if len(paid_col_letters) == len(cols_pago):
            first_l = get_column_letter(5)
            last_l = get_column_letter(4 + len(cols_pago))
            return f"=SUM({first_l}{row_num}:{last_l}{row_num})"
        terms = "+".join(f"{c_l}{row_num}" for c_l in paid_col_letters)
        return f"={terms}"

    curr_r = 7
    first_c_row = curr_r
    for f_c in res["filas_contratistas"]:
        part_s = safe_str(f_c.get("Partida", "")).strip()
        desc_s = safe_str(f_c.get("Descripcion", "")).strip()
        if part_s and desc_s and part_s.upper() not in desc_s.upper():
            ws[f"B{curr_r}"] = f"{part_s} — {desc_s}"
        else:
            ws[f"B{curr_r}"] = desc_s or part_s
        ws[f"B{curr_r}"].alignment = align_left
        ws[f"C{curr_r}"] = f_c["Moneda"]
        ws[f"C{curr_r}"].alignment = align_center
        ws[f"D{curr_r}"] = float(f_c["TOTAL"])
        ws[f"D{curr_r}"].number_format = "#,##0.00"
        ws[f"D{curr_r}"].alignment = align_right

        for i_c, c_name in enumerate(cols_pago):
            c_let = get_column_letter(5 + i_c)
            val_c = float(f_c.get(c_name, 0.0))
            ws[f"{c_let}{curr_r}"] = val_c if val_c != 0.0 else None
            ws[f"{c_let}{curr_r}"].number_format = "#,##0.00"
            ws[f"{c_let}{curr_r}"].alignment = align_right

        ws[f"{let_pagado}{curr_r}"] = _formula_pagado(curr_r)
        ws[f"{let_pagado}{curr_r}"].number_format = "#,##0.00"
        ws[f"{let_pagado}{curr_r}"].alignment = align_right

        ws[f"{let_saldo}{curr_r}"] = f"=D{curr_r}-{let_pagado}{curr_r}"
        ws[f"{let_saldo}{curr_r}"].number_format = "#,##0.00"
        ws[f"{let_saldo}{curr_r}"].alignment = align_right

        ws[f"{let_notas}{curr_r}"] = f_c.get("Notas", "")
        ws[f"{let_notas}{curr_r}"].alignment = align_left

        for c_i in range(2, col_idx_notas + 1):
            c_l = get_column_letter(c_i)
            ws[f"{c_l}{curr_r}"].font = f_row
            ws[f"{c_l}{curr_r}"].border = border_cell

        curr_r += 1

    last_c_row = max(first_c_row, curr_r - 1)

    # Fila SUB TOTAL
    row_sub = curr_r
    ws[f"B{row_sub}"] = "SUB TOTAL"
    ws[f"C{row_sub}"] = ""
    for c_i in range(4, col_idx_saldo + 1):
        c_l = get_column_letter(c_i)
        ws[f"{c_l}{row_sub}"] = f"=SUM({c_l}{first_c_row}:{c_l}{last_c_row})"
        ws[f"{c_l}{row_sub}"].number_format = "#,##0.00"
        ws[f"{c_l}{row_sub}"].alignment = align_right
    for c_i in range(2, col_idx_notas + 1):
        c_l = get_column_letter(c_i)
        ws[f"{c_l}{row_sub}"].font = f_subtotal
        ws[f"{c_l}{row_sub}"].fill = fill_subtotal
        ws[f"{c_l}{row_sub}"].border = border_cell

    # Fila Gestión y supervisión de obras
    row_ges = row_sub + 1
    f_ges = res["fila_gestion"]
    ws[f"B{row_ges}"] = "Gestión y supervisión de obras"
    ws[f"C{row_ges}"] = f_ges["Moneda"]
    ws[f"C{row_ges}"].alignment = align_center
    ws[f"D{row_ges}"] = float(f_ges["TOTAL"])
    ws[f"D{row_ges}"].number_format = "#,##0.00"
    ws[f"D{row_ges}"].alignment = align_right
    for i_c, c_name in enumerate(cols_pago):
        c_let = get_column_letter(5 + i_c)
        v_g = float(f_ges.get(c_name, 0.0))
        ws[f"{c_let}{row_ges}"] = v_g if v_g != 0.0 else None
        ws[f"{c_let}{row_ges}"].number_format = "#,##0.00"
        ws[f"{c_let}{row_ges}"].alignment = align_right
    ws[f"{let_pagado}{row_ges}"] = _formula_pagado(row_ges)
    ws[f"{let_pagado}{row_ges}"].number_format = "#,##0.00"
    ws[f"{let_pagado}{row_ges}"].alignment = align_right
    ws[f"{let_saldo}{row_ges}"] = f"=D{row_ges}-{let_pagado}{row_ges}"
    ws[f"{let_saldo}{row_ges}"].number_format = "#,##0.00"
    ws[f"{let_saldo}{row_ges}"].alignment = align_right
    ws[f"{let_notas}{row_ges}"] = f_ges.get("Notas", "")
    for c_i in range(2, col_idx_notas + 1):
        c_l = get_column_letter(c_i)
        ws[f"{c_l}{row_ges}"].font = f_row
        ws[f"{c_l}{row_ges}"].border = border_cell

    # Fila Gastos Generales
    row_gg = row_ges + 1
    f_gg = res["fila_gg"]
    ws[f"B{row_gg}"] = "Gastos Generales"
    ws[f"C{row_gg}"] = f_gg["Moneda"]
    ws[f"C{row_gg}"].alignment = align_center
    ws[f"D{row_gg}"] = float(f_gg["TOTAL"])
    ws[f"D{row_gg}"].number_format = "#,##0.00"
    ws[f"D{row_gg}"].alignment = align_right
    for i_c, c_name in enumerate(cols_pago):
        c_let = get_column_letter(5 + i_c)
        v_gg = float(f_gg.get(c_name, 0.0))
        ws[f"{c_let}{row_gg}"] = v_gg if v_gg != 0.0 else None
        ws[f"{c_let}{row_gg}"].number_format = "#,##0.00"
        ws[f"{c_let}{row_gg}"].alignment = align_right
    ws[f"{let_pagado}{row_gg}"] = _formula_pagado(row_gg)
    ws[f"{let_pagado}{row_gg}"].number_format = "#,##0.00"
    ws[f"{let_pagado}{row_gg}"].alignment = align_right
    ws[f"{let_saldo}{row_gg}"] = f"=D{row_gg}-{let_pagado}{row_gg}"
    ws[f"{let_saldo}{row_gg}"].number_format = "#,##0.00"
    ws[f"{let_saldo}{row_gg}"].alignment = align_right
    ws[f"{let_notas}{row_gg}"] = f_gg.get("Notas", "")
    for c_i in range(2, col_idx_notas + 1):
        c_l = get_column_letter(c_i)
        ws[f"{c_l}{row_gg}"].font = f_row
        ws[f"{c_l}{row_gg}"].border = border_cell

    # Fila TOTAL Soles
    row_tot = row_gg + 1
    ws[f"B{row_tot}"] = "TOTAL"
    ws[f"C{row_tot}"] = "Soles"
    ws[f"C{row_tot}"].alignment = align_center
    for c_i in range(4, col_idx_saldo + 1):
        c_l = get_column_letter(c_i)
        ws[f"{c_l}{row_tot}"] = f"={c_l}{row_sub}+{c_l}{row_ges}+{c_l}{row_gg}"
        ws[f"{c_l}{row_tot}"].number_format = "#,##0.00"
        ws[f"{c_l}{row_tot}"].alignment = align_right
    for c_i in range(2, col_idx_notas + 1):
        c_l = get_column_letter(c_i)
        ws[f"{c_l}{row_tot}"].font = f_total
        ws[f"{c_l}{row_tot}"].fill = fill_total
        ws[f"{c_l}{row_tot}"].border = border_cell

    # Fila Dólares
    row_usd = row_tot + 1
    ws[f"B{row_usd}"] = ""
    ws[f"C{row_usd}"] = "Dólares"
    ws[f"C{row_usd}"].alignment = align_center
    for c_i in range(4, col_idx_saldo + 1):
        c_l = get_column_letter(c_i)
        ws[f"{c_l}{row_usd}"] = f"=IFERROR({c_l}{row_tot}/$F$2,0)"
        ws[f"{c_l}{row_usd}"].number_format = "#,##0.00"
        ws[f"{c_l}{row_usd}"].alignment = align_right
    for c_i in range(2, col_idx_notas + 1):
        c_l = get_column_letter(c_i)
        ws[f"{c_l}{row_usd}"].font = f_usd
        ws[f"{c_l}{row_usd}"].fill = fill_usd
        ws[f"{c_l}{row_usd}"].border = border_closing

    # Anchos de columna
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 15
    for i_c in range(len(cols_pago)):
        ws.column_dimensions[get_column_letter(5 + i_c)].width = 14
    ws.column_dimensions[let_pagado].width = 15
    ws.column_dimensions[let_saldo].width = 15
    ws.column_dimensions[let_notas].width = 24

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def parsear_cuadro_pagos_semanales(file_or_df, tc_default: float = 3.60) -> dict:
    """
    Lee e interpreta un archivo CSV o Excel (.xlsx) con el formato del Cuadro de Pagos Semanales
    (como el compartido por el usuario) y retorna la estructura lista para guardarse en el proyecto.
    """
    rows = _rows_from_input(file_or_df)
    if not rows:
        return {}

    header_idx = None
    col_desc = 1
    col_moneda = 2
    col_total = 3
    col_pagado = None
    col_saldo = None
    cols_semanas_map = {}  # {col_idx: nombre_col}

    for r_i, row in enumerate(rows[:15]):
        row_strs = [_clean_cell_str(c) for c in row]
        row_low = [s.lower() for s in row_strs]
        if any("descrip" in s for s in row_low) and any("adelanto" in s or "semana" in s for s in row_low):
            header_idx = r_i
            for c_i, s_l in enumerate(row_low):
                if "descrip" in s_l:
                    col_desc = c_i
                elif "moneda" in s_l:
                    col_moneda = c_i
                elif s_l == "total":
                    col_total = c_i
                elif "adelanto" in s_l or "semana" in s_l:
                    cols_semanas_map[c_i] = row_strs[c_i].strip().upper()
                elif "pagado" in s_l:
                    col_pagado = c_i
                elif "saldo" in s_l:
                    col_saldo = c_i
            break

    if header_idx is None:
        return {}

    # Determinar número de semanas
    num_sem = 5
    for c_name in cols_semanas_map.values():
        m_s = re.search(r"SEMANA\s*(\d+)", c_name, re.IGNORECASE)
        if m_s:
            num_sem = max(num_sem, int(m_s.group(1)))

    cols_pago = obtener_columnas_pago_semanal(num_sem)
    cols_por_pagar = []

    # Revisar la fila inmediatamente posterior al header por si dice "Por Pagar"
    start_data_row = header_idx + 1
    if start_data_row < len(rows):
        next_row = rows[start_data_row]
        has_por_pagar = False
        for c_i, cell_v in enumerate(next_row):
            s_v = _clean_cell_str(cell_v).lower()
            if "por pagar" in s_v and c_i in cols_semanas_map:
                cols_por_pagar.append(cols_semanas_map[c_i])
                has_por_pagar = True
        if has_por_pagar:
            start_data_row += 1

    filas_contratistas = []
    fila_gestion = {"Descripcion": "Gestión y supervisión de obras", "Moneda": "Soles", "TOTAL": 0.0}
    fila_gg = {"Descripcion": "Gastos Generales", "Moneda": "Soles", "TOTAL": 0.0}
    tc_detected = tc_default
    total_soles_val = 0.0

    for r_i in range(start_data_row, len(rows)):
        row = rows[r_i]
        desc = _clean_cell_str(row[col_desc]) if col_desc < len(row) else ""
        mon = _clean_cell_str(row[col_moneda]) if col_moneda < len(row) else "Soles"
        tot_v, _ = _parse_number_or_pct(row[col_total] if col_total < len(row) else 0.0, 0.0)

        desc_low = desc.lower()
        mon_low = mon.lower()

        if "trabajos contratistas" in desc_low:
            continue
        if desc_low in ["sub total", "subtotal"]:
            continue

        # Detectar fila de Dólares al final para inferir TC si aplica
        if any(d in mon_low for d in ["dolar", "dólar", "usd"]) and (not desc or "total" in desc_low):
            if tot_v > 0 and total_soles_val > 0:
                tc_detected = round(total_soles_val / tot_v, 2)
            continue

        if desc_low == "total":
            total_soles_val = tot_v
            continue

        if not desc and tot_v == 0.0:
            continue

        # Extraer notas en columnas posteriores a SALDO
        notas_parts = []
        start_notas_col = (col_saldo + 1) if col_saldo is not None else (max(cols_semanas_map.keys(), default=col_total) + 3)
        for c_n in range(start_notas_col, len(row)):
            s_n = _clean_cell_str(row[c_n])
            if s_n:
                notas_parts.append(s_n)
        notas_str = " | ".join(notas_parts)

        fila_dict = {
            "Descripcion": desc,
            "Moneda": mon or "Soles",
            "TOTAL": round(tot_v, 2),
            "Notas": notas_str,
        }
        for c_i, c_name in cols_semanas_map.items():
            v_sem, _ = _parse_number_or_pct(row[c_i] if c_i < len(row) else 0.0, 0.0)
            fila_dict[c_name] = round(v_sem, 2)

        if any(k in desc_low for k in ["gestión", "gestion", "supervisión", "supervision", "utilidad"]):
            fila_gestion = _normalizar_fila_pago(fila_dict, cols_pago, "Gestión y supervisión de obras")
        elif "gastos generales" in desc_low:
            fila_gg = _normalizar_fila_pago(fila_dict, cols_pago, "Gastos Generales")
        else:
            filas_contratistas.append(_normalizar_fila_pago(fila_dict, cols_pago, desc))

    return {
        "num_semanas": num_sem,
        "tc": tc_detected if tc_detected > 0 else tc_default,
        "cols_por_pagar": cols_por_pagar,
        "filas_contratistas": filas_contratistas,
        "fila_gestion": _normalizar_fila_pago(fila_gestion, cols_pago, "Gestión y supervisión de obras"),
        "fila_gg": _normalizar_fila_pago(fila_gg, cols_pago, "Gastos Generales"),
    }

